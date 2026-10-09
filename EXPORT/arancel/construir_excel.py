import pandas as pd, sys
from openpyxl.styles import Font, PatternFill, Alignment
OUT=sys.argv[1]
d=pd.read_csv(f"{OUT}/diccionario_partidas.csv",dtype=str).fillna("")
c=pd.read_csv(f"{OUT}/correlaciones_sunat.csv",dtype=str).fillna("")
v=pd.read_csv(f"{OUT}/versiones_arancel.csv",dtype=str).fillna("")
cols=["PARTIDA","CODIGO_SUNAT","PARTIDA_4","PARTIDA_4_DESC","DESCRIPCION","V2002","V2007","V2012","V2017","V2022",
      "ESTADO","SERIE_ESTABLE","TRAMOS","MISMA_DESCRIPCION","EQUIV_ANTERIORES","EQUIV_SIGUIENTES"]
idx=d.groupby(["CAPITULO","CAPITULO_DESC"]).agg(
    SUBPARTIDAS=("PARTIDA","size"),
    VIGENTES_2022=("VIGENTE",lambda s:(s=="1").sum()),
    ESTABLES_2002_HOY=("ESTADO",lambda s:(s=="ESTABLE 2002-HOY").sum()),
    DESCONTINUADAS=("ESTADO",lambda s:s.str.startswith("DESCONTINUADA").sum())).reset_index()
idx.insert(0,"HOJA",["Cap "+x for x in idx.CAPITULO])
leeme=pd.DataFrame({"CONTENIDO":[
 "Diccionario de subpartidas nacionales SUNAT (10 dígitos, sin puntos: 4412310000) por capítulo.",
 "Fuentes: Aranceles de Aduanas 2002, 2007, 2012, 2017, 2022 y correlaciones teóricas SUNAT 2002-2007, 2007-2012, 2012-2017, 2017-2022.",
 "V2002..V2022 = 1 si el código existe en esa versión del arancel.",
 "SERIE_ESTABLE = tramo más largo de versiones consecutivas en que el código existe y la correlación es 1 a 1 consigo mismo (sin divisiones, fusiones ni recodificaciones): la serie es comparable en ese periodo.",
 "TRAMOS = todos los tramos estables del código (un código puede desaparecer y reutilizarse).",
 "ESTADO: ESTABLE 2002-HOY / ESTABLE DESDE AAAA (vigente, comparable desde esa versión) / NUEVA 2022 / DESCONTINUADA (ÚLTIMA AAAA).",
 "MISMA_DESCRIPCION = SI si el texto de la subpartida es igual (≥85% de similitud) en todo su tramo estable; NO = cambió la redacción aunque la correlación diga 1 a 1.",
 "EQUIV_ANTERIORES / EQUIV_SIGUIENTES = códigos de la versión anterior / siguiente con los que se correlaciona cuando NO hay continuidad (para empalmar series).",
 "Vigencias: 2002 desde 01/01/2002; 2007 desde 01/04/2007 (ene-mar 2007 rige el 2002); 2012 desde 01/01/2012; 2017 desde 01/01/2017; 2022 desde 01/01/2022.",
 "Advertencia: algunas descripciones 2002/2017 provienen de PDF convertidos a texto y pueden tener errores de OCR; las de 2007/2012/2022 provienen de las tablas de correlación."]})
with pd.ExcelWriter(f"{OUT}/diccionario_partidas.xlsx",engine="openpyxl") as w:
    leeme.to_excel(w,sheet_name="LEEME",index=False); v.to_excel(w,sheet_name="VERSIONES",index=False); idx.to_excel(w,sheet_name="INDICE",index=False)
    for (cap,_),g in d.groupby(["CAPITULO","CAPITULO_DESC"]):
        g[cols].to_excel(w,sheet_name=f"Cap {cap}",index=False)
    c.to_excel(w,sheet_name="CORRELACIONES",index=False)
    for ws in w.book.worksheets:
        ws.freeze_panes="A2"
        for cell in ws[1]:
            cell.font=Font(bold=True,color="FFFFFF"); cell.fill=PatternFill("solid",fgColor="1F4E78")
        for col in ws.columns:
            L=max(len(str(x.value or "")) for x in list(col)[:300])
            ws.column_dimensions[col[0].column_letter].width=min(max(8,L+2),70)
    w.book["LEEME"].column_dimensions["A"].width=160
print(len(w.book.sheetnames), w.book.sheetnames[:6])
