"""Corre FISCAL/Fiscal_Procesamiento.ipynb sin Colab (GitHub Actions) y guarda sus Excel en la base del API.

    python FISCAL/automatizacion/procesamiento_servidor.py

El notebook lee los CSV de Scrapeo/output y el IPC.xlsx de Procesamiento/input y genera Fiscal-<Depto>_AUTO.xlsx (los
cuadros de gasto, inversión y PIM de cada región) y «Variaciones alarmantes AAAA-MM.xlsx». Aquí:
  · los CSV se arman con las bases del API (las que carga el scraper del MEF), con los mismos nombres y columnas que
    deja el scraper; el IPC.xlsx, con el IPC de Lima cargado en la página;
  · se corren todos sus bloques; la subida a Drive del bloque 6 se reemplaza por guardar los Excel en la base, de donde
    se descargan en la página (Ejecución del presupuesto → Archivos). Necesita DATABASE_URL.
"""
import json
import os
import re
import sys
import tempfile
import time
import traceback
import types
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook

AQUI = Path(__file__).resolve().parent
RAIZ = AQUI.parent.parent
RUTA_COLAB = "/content/drive/MyDrive/Fiscal"
CARPETA_ARCHIVOS = "fiscal"           # carpeta de los archivos en la base (db.archivos)
sys.path.insert(0, str(AQUI))
sys.path.insert(0, str(RAIZ / "FISCAL" / "api"))
from scraper_servidor import sustituir_colab  # noqa: E402

DEPTOS = ["16", "22", "25"]
NIVELES = {"E": ("Gobierno Nacional", "GobNacional"), "R": ("Gobierno Regional", "GobRegional"), "M": ("Gobierno Municipal", "Municipalidad")}
SUFIJO = "LOR - SMN - UCA"


def sin_drive():
    """google.colab.auth que no autentica: el bloque 6 de los notebooks cae a su copia en Procesamiento/output."""
    sustituir_colab()
    auth = types.ModuleType("google.colab.auth")

    def authenticate_user(*a, **k):
        raise RuntimeError("en GitHub Actions el Excel se guarda en la página, no en Drive")
    auth.authenticate_user = authenticate_user
    sys.modules["google.colab"].auth = auth
    sys.modules["google.colab.auth"] = auth


def celdas(notebook, raiz, colab, fijar=None):
    """Código del notebook listo para correr aquí: sin «!pip»/«%», con la carpeta de trabajo y las variables fijadas."""
    nb = json.loads(Path(notebook).read_text(encoding="utf-8"))
    Path(colab).mkdir(parents=True, exist_ok=True)              # el «/content» de Colab, donde se guarda antes de subir
    fijar, salida, fijadas = fijar or {}, [], set()
    for c in nb["cells"]:
        if c["cell_type"] != "code":
            continue
        lineas = []
        for linea in "".join(c["source"]).split("\n"):
            if linea.lstrip().startswith(("!", "%")):
                continue
            m = re.match(r"^(\w+)\s*=", linea)
            if m and m.group(1) in fijar:
                linea = f"{m.group(1)} = {fijar[m.group(1)]!r}"
                fijadas.add(m.group(1))
            lineas.append(linea)
        codigo = "\n".join(lineas).replace(RUTA_COLAB, str(raiz))
        salida.append(codigo.replace('os.path.join("/content", ', f'os.path.join({str(colab)!r}, '))
    faltan = set(fijar) - fijadas
    if faltan:
        raise RuntimeError(f"{Path(notebook).name} ya no define {', '.join(sorted(faltan))}: revisa {Path(__file__).name}")
    return salida


def correr(notebook, raiz, fijar=None, g=None):
    """Ejecuta el notebook en un diccionario de variables (que devuelve) con Drive = raiz."""
    sin_drive()
    g = g if g is not None else {"__name__": "__main__", "display": lambda *x, **k: print(*x)}
    for i, codigo in enumerate(celdas(notebook, raiz, raiz.parent / "content", fijar), 1):
        print(f"\n── {Path(notebook).stem} · celda {i} ──", flush=True)
        exec(compile(codigo, f"{Path(notebook).name}#celda{i}", "exec"), g)
    return g


def _mes_num(df):
    return df["MES"].astype(str).str.extract(r"^(\d+)")[0].astype(int)


def _por_mes(df, carpeta, nombre):
    """Un CSV por mes, «<nombre> AAAA - Mes MM - LOR - SMN - UCA.csv», como los deja el scraper."""
    if df.empty:
        return 0
    carpeta.mkdir(parents=True, exist_ok=True)
    df = df.assign(_m=_mes_num(df))
    for (a, m), g in df.groupby(["AÑO", "_m"]):
        g.drop(columns="_m").to_csv(carpeta / f"{nombre} {int(a)} - Mes {int(m):02d} - {SUFIJO}.csv", index=False, encoding="utf-8-sig")
    return len(df)


def generica_y_deuda(gen):
    """La base de genéricas trae la 2.8 desglosada (DESGLOSE); el notebook espera DEV Generica sin desglose y DEV Deuda."""
    if gen.empty:
        return gen, gen
    claves = [c for c in gen.columns if c not in ("DESGLOSE", "DEVENGADO")]
    generica = gen.groupby(claves, as_index=False, sort=False)["DEVENGADO"].sum()
    d = gen[gen["DESGLOSE"].isin(["AMORTIZACIÓN", "INTERESES Y COMISIONES"])].copy()
    if d.empty:
        return generica, d
    d["CLASIFICACION"] = d["DESGLOSE"].str.replace("Ó", "O")
    d["RUTA"] = d["GENERICA"].astype(str) + " > " + d["DESGLOSE"]
    deuda = d[["AÑO", "MES", "NIVEL_GOBIERNO", "DEPARTAMENTO", "CLASIFICACION", "RUTA", "DEVENGADO"]]
    return generica, deuda


def preparar(raiz):
    """Scrapeo/output y Procesamiento/input/IPC.xlsx desde la base. Devuelve las filas escritas por archivo."""
    from app import db, ipc
    salida = raiz / "Scrapeo" / "output"
    hechas = {}
    for nivel, (carpeta, corto) in NIVELES.items():
        destino = salida / "DEVENGADO" / carpeta
        if nivel == "M":
            fun = db.leer("inv-dev-municipalidad", DEPTOS, [], "1990-01", "2100-12")
        else:
            fun = db.leer("inv-dev-funcion", DEPTOS, [nivel], "1990-01", "2100-12")
        hechas[f"DEV Funcion.{corto}"] = _por_mes(fun, destino, f"DEV Funcion.{corto}")
        hechas[f"DEV InvFuente.{corto}"] = _por_mes(db.leer("inv-dev-fuente", DEPTOS, [nivel], "1990-01", "2100-12"), destino, f"DEV InvFuente.{corto}")
        hechas[f"DEV Fuente.{corto}"] = _por_mes(db.leer("gasto-dev-fuente", DEPTOS, [nivel], "1990-01", "2100-12"), destino, f"DEV Fuente.{corto}")
        gen, deu = generica_y_deuda(db.leer("gasto-dev-generica", DEPTOS, [nivel], "1990-01", "2100-12"))
        hechas[f"DEV Generica.{corto}"] = _por_mes(gen, destino, f"DEV Generica.{corto}")
        hechas[f"DEV Deuda.{corto}"] = _por_mes(deu, destino, f"DEV Deuda.{corto}")
    pim = db.leer("inv-pim-funcion", DEPTOS, ["E", "R", "M"], "1990-01", "2100-12")
    destino = salida / "PIM" / "PIM X FUNCIÓN" / "DATA"
    destino.mkdir(parents=True, exist_ok=True)
    for a, g in pim.groupby("AÑO"):
        g.to_csv(destino / f"PIM Funcion {int(a)} - {SUFIJO}.csv", index=False, encoding="utf-8-sig")
    hechas["PIM Funcion"] = len(pim)
    hechas["IPC"] = ipc_xlsx(ipc.serie(), raiz / "Procesamiento" / "input" / "IPC.xlsx")
    return hechas


def ipc_xlsx(serie, ruta):
    """El IPC.xlsx que lee el notebook: Fecha en la columna B e IPC en la C desde la fila 4."""
    if not serie:
        raise RuntimeError("No hay IPC de Lima cargado en la página (Las tres regiones → Actualización).")
    ruta.parent.mkdir(parents=True, exist_ok=True)
    wb = Workbook()
    ws = wb.active
    ws["B1"] = "IPC de Lima Metropolitana (índice), cargado en la página"
    ws["B3"], ws["C3"] = "Fecha", "IPC"
    for i, (p, v) in enumerate(sorted(serie.items()), start=4):
        ws.cell(i, 2, datetime(int(p[:4]), int(p[5:]), 1))
        ws.cell(i, 3, float(v))
    wb.save(ruta)
    return len(serie)


def guardar(carpeta_salida):
    """Guarda en la base los .xlsx de Procesamiento/output. Devuelve sus nombres."""
    from app import db
    nombres = []
    for ruta in sorted(Path(carpeta_salida).glob("*.xlsx")):
        db.guardar_archivo(CARPETA_ARCHIVOS, ruta.name, ruta.read_bytes())
        nombres.append(ruta.name)
    return nombres


def main():
    lineas, ok, t0 = [f"## Procesamiento fiscal · {datetime.now():%d/%m/%Y %H:%M}\n"], True, time.time()
    try:
        with tempfile.TemporaryDirectory() as tmp:
            raiz = Path(tmp) / "Fiscal"
            hechas = preparar(raiz)
            print("Inputs desde la base:", ", ".join(f"{k}: {v}" for k, v in hechas.items()), flush=True)
            os.chdir(tmp)
            correr(AQUI.parent / "Fiscal_Procesamiento.ipynb", raiz)
            nombres = guardar(raiz / "Procesamiento" / "output")
        if not nombres:
            raise RuntimeError("el notebook no generó ningún Excel")
        lineas.append(f"**✅ {len(nombres)} archivos** guardados en la página ({(time.time() - t0) / 60:.1f} min):")
        lineas += [f"- {n}" for n in nombres]
        vacios = [k for k, v in hechas.items() if not v]
        if vacios:
            lineas.append(f"- ⚠ sin datos en la base: {', '.join(vacios)}")
    except BaseException as e:
        traceback.print_exc()
        lineas.append(f"**❌ {type(e).__name__}: {str(e)[:500]}**")
        ok = False
    texto = "\n".join(lineas)
    print("\n" + texto)
    if os.environ.get("GITHUB_STEP_SUMMARY"):
        with open(os.environ["GITHUB_STEP_SUMMARY"], "a", encoding="utf-8") as f:
            f.write(texto + "\n")
    sys.exit(0 if ok else 1)


if __name__ == "__main__":
    main()
