"""Empleo: trabajadores asalariados privados formales (MTPE, Planilla Electrónica: PLAME y T-Registro).

Entra el Excel que genera EMPLEO/Scraping_MTPE_Empleo_x_Region_Actividad.ipynb (en Colab o en GitHub Actions):
columnas AÑO, MES, REGIÓN, ACTIVIDAD_ECONOMICA, EMPLEO, una fila por mes, región y actividad.

También arma lo que publica la síntesis de cada región (junio 2026):
  · el cuadro «Trabajadores asalariados privados formales» (Loreto N° 6, San Martín N° 4, Ucayali N° 7): siete ramas y
    el total, el mes contra el mismo mes del año anterior y contra diciembre del año anterior, con su variación %;
  · el gráfico de la variación % interanual del total (25 meses);
  · el anexo por sectores (Loreto, Anexo 5): variación % interanual de cada sector en los últimos 13 meses.
"""
import io
import re
import unicodedata

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font

from .catalogo import DEPARTAMENTOS

COLUMNAS = ["AÑO", "MES", "REGIÓN", "ACTIVIDAD_ECONOMICA", "EMPLEO"]
MESES = ["Enero", "Febrero", "Marzo", "Abril", "Mayo", "Junio", "Julio", "Agosto", "Setiembre", "Octubre", "Noviembre", "Diciembre"]
MESES_CORTOS = ["Ene.", "Feb.", "Mar.", "Abr.", "May.", "Jun.", "Jul.", "Ago.", "Set.", "Oct.", "Nov.", "Dic."]
# rama del cuadro de la síntesis (orden alfabético, como en el informe) y del anexo
RAMAS = {"Agrícola": "Agrícola", "Comercio": "Comercio", "Construcción": "Construcción", "Manufactura": "Manufactura",
         "Minería": "Minería", "No precisa": "Otros", "Servicios": "Servicios"}
ORDEN_CUADRO = ["Agrícola", "Comercio", "Construcción", "Manufactura", "Minería", "Otros", "Servicios"]
ORDEN_ANEXO = [("Agrícola", "Agricultura"), ("Minería", "Minería"), ("Manufactura", "Manufactura"),
               ("Construcción", "Construcción"), ("Comercio", "Comercio"), ("Servicios", "Servicios")]
FUENTE = "Fuente: MTPE, Planilla Electrónica – PLAME y T-Registro. Elaboración: BCRP, Sucursal Iquitos. Departamento de Estudios Económicos."


def _clave(t):
    t = unicodedata.normalize("NFKD", str(t)).encode("ascii", "ignore").decode()
    return re.sub(r"[^a-z]", "", t.lower())


_DEPTO_DE_REGION = {_clave(n): c for c, n in DEPARTAMENTOS.items()}
_DEPTO_DE_REGION.update({"limametropolitana": "15", "limaprovincias": "15", "lima": "15"})


def depto_de(region):
    return _DEPTO_DE_REGION.get(_clave(region))


def reconocer(ruta):
    """True si el Excel tiene las columnas del scraper de empleo."""
    columnas = {str(c).strip().upper() for c in pd.read_excel(ruta, sheet_name=0, nrows=0).columns}
    return {"AÑO", "MES", "REGIÓN", "ACTIVIDAD_ECONOMICA", "EMPLEO"} <= columnas


def leer(archivos):
    """archivos: [(n_archivo, ruta)] -> DataFrame con columnas internas, avisos."""
    partes, avisos = [], []
    for n, ruta in archivos:
        df = pd.read_excel(ruta, sheet_name=0)
        df.columns = [str(c).strip().upper() for c in df.columns]
        df = df[COLUMNAS].dropna(subset=["AÑO", "MES", "REGIÓN"])
        df["_depto"] = df["REGIÓN"].map(depto_de)
        sin = sorted(set(df.loc[df["_depto"].isna(), "REGIÓN"].astype(str)))
        if sin:
            avisos.append(f"Empleo: {', '.join(sin)} no es un departamento; no se cargó (el total del Perú no se guarda).")
        df = df[df["_depto"].notna()].copy()
        df["AÑO"], df["MES"] = df["AÑO"].astype(int), df["MES"].astype(int)
        df["EMPLEO"] = pd.to_numeric(df["EMPLEO"], errors="coerce").fillna(0)
        df["REGIÓN"], df["ACTIVIDAD_ECONOMICA"] = df["REGIÓN"].astype(str), df["ACTIVIDAD_ECONOMICA"].astype(str)
        # Lima Metropolitana y Lima Provincias comparten departamento: se reemplazan por separado
        df["_anio"], df["_mes"], df["_nivel"], df["_archivo"] = df["AÑO"], df["MES"], df["REGIÓN"].map(_clave), n
        partes.append(df)
    return (pd.concat(partes, ignore_index=True) if partes else pd.DataFrame()), avisos


# ── Lo que publica la síntesis ──
def _serie(df):
    """Trabajadores por rama y periodo (AAAA*100+MM) y el total."""
    df = df.assign(RAMA=df["ACTIVIDAD_ECONOMICA"].map(RAMAS).fillna("Otros"), P=df["AÑO"] * 100 + df["MES"])
    t = df.pivot_table(index="P", columns="RAMA", values="EMPLEO", aggfunc="sum").reindex(columns=ORDEN_CUADRO).fillna(0)
    t["Total"] = t.sum(axis=1)
    return t


def _var(a, b):
    return None if not a else round((b / a - 1) * 100, 1)


def etiqueta(p, corta=False):
    return f"{(MESES_CORTOS if corta else MESES)[p % 100 - 1]} {p // 100}"


def cuadros(df, periodo=None):
    """df: filas de empleo de una región. periodo AAAA-MM (por defecto, el último con datos)."""
    if df.empty:
        return None
    t = _serie(df)
    ultimo = int(t.index.max())
    p = int(periodo.replace("-", "")) if periodo else ultimo
    if p not in t.index:
        raise ValueError(f"No hay datos de empleo para {etiqueta(p)} (último mes: {etiqueta(ultimo)}).")
    p_ant, p_dic = p - 100, (p // 100 - 1) * 100 + 12
    fila = lambda rama: {"rama": rama, "mes_anterior": _val(t, p_ant, rama), "mes": float(t.at[p, rama]),
                         "var_mes": _var(_val(t, p_ant, rama), t.at[p, rama]), "diciembre": _val(t, p_dic, rama),
                         "var_dic": _var(_val(t, p_dic, rama), t.at[p, rama])}
    cuadro = [fila(r) for r in ORDEN_CUADRO + ["Total"]]
    periodos = [q for q in t.index if q <= p]
    grafico = [{"periodo": etiqueta(q, True), "var": _var(_val(t, q - 100, "Total"), t.at[q, "Total"])} for q in periodos[-25:]]
    anexo = [{"periodo": etiqueta(q, True), **{n: _var(_val(t, q - 100, r), t.at[q, r]) for r, n in ORDEN_ANEXO},
              "Total": _var(_val(t, q - 100, "Total"), t.at[q, "Total"])} for q in periodos[-13:]]
    return {"periodo": f"{p // 100}-{p % 100:02d}", "ultimo": f"{ultimo // 100}-{ultimo % 100:02d}",
            "titulo_mes": MESES[p % 100 - 1], "anio": p // 100, "cuadro": cuadro, "grafico": grafico, "anexo": anexo}


def _val(t, p, rama):
    return float(t.at[p, rama]) if p in t.index else None


def excel(region, c):
    """Los tres cuadros en un Excel, con títulos, unidades y fuente como en la síntesis."""
    wb = Workbook()
    neg, centro = Font(bold=True), Alignment(horizontal="center")
    ws = wb.active
    ws.title = "Cuadro"
    a, mes = c["anio"], c["titulo_mes"]
    for fila in ([f"{region.upper()}: TRABAJADORES ASALARIADOS PRIVADOS FORMALES"], ["(Número de trabajadores)"], [],
                 ["Rama de actividad", f"{mes} {a - 1}", f"{mes} {a}", "Var. %", f"Dic. {a - 1}", f"{mes[:3]}. {a}", "Var. %"]):
        ws.append(fila)
    for r in c["cuadro"]:
        ws.append([r["rama"], r["mes_anterior"], r["mes"], r["var_mes"], r["diciembre"], r["mes"], r["var_dic"]])
    ws.append([]); ws.append([FUENTE])
    for celda in ws[4]:
        celda.font, celda.alignment = neg, centro
    for fila in ws.iter_rows(min_row=5, max_row=4 + len(c["cuadro"])):
        for i, celda in enumerate(fila):
            celda.number_format = "0.0" if i in (3, 6) else "#,##0"
        if fila[0].value == "Total":
            for celda in fila:
                celda.font = neg
    ws["A1"].font = neg
    for col, ancho in zip("ABCDEFG", [22, 12, 12, 9, 12, 12, 9]):
        ws.column_dimensions[col].width = ancho

    wg = wb.create_sheet("Gráfico")
    wg.append([f"{region.upper()}: TRABAJADORES ASALARIADOS PRIVADOS FORMALES"]); wg.append(["(Variación % respecto a similar mes del año anterior)"]); wg.append([])
    wg.append(["Mes", "Var. %"])
    for g in c["grafico"]:
        wg.append([g["periodo"], g["var"]])
    wg["A1"].font = neg

    wa = wb.create_sheet("Anexo")
    wa.append([f"{region.upper()}: TRABAJADORES ASALARIADOS PRIVADOS FORMALES POR SECTORES ECONÓMICOS"])
    wa.append(["(Variación % con respecto al mismo mes del año anterior)"]); wa.append([])
    wa.append(["Mes"] + [n for _, n in ORDEN_ANEXO] + ["Total"])
    for r in c["anexo"]:
        wa.append([r["periodo"]] + [("n.d." if r[n] is None else r[n]) for _, n in ORDEN_ANEXO] + ["n.d." if r["Total"] is None else r["Total"]])
    wa.append([]); wa.append([FUENTE])
    wa["A1"].font = neg
    for hoja in (wg, wa):
        for celda in hoja[4]:
            celda.font, celda.alignment = neg, centro
        for fila in hoja.iter_rows(min_row=5):
            for celda in fila[1:]:
                if isinstance(celda.value, (int, float)):
                    celda.number_format = "0.0"
    salida = io.BytesIO()
    wb.save(salida)
    return salida.getvalue()
