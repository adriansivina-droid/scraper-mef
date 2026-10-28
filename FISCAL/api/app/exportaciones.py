"""Exportaciones (e importaciones de Loreto): el ESTAD de cada región, la única fuente de los cuadros (EXPORT/info_export.md).

Se reconoce el libro ESTAD por sus hojas, no por el nombre: la hoja cuyo título (A2) dice «LORETO: EXPORTACIONES FOB …»,
«SAN MARTÍN: …» o «UCAYALI: …» (ESTADLO 15.Export, ESTADSMN 14.XLS, ESTADUCA 15.XLS) y, en Loreto, la de importaciones
(16.Import). De cada hoja:
  · fila 4: año en el primer mes de cada bloque, «Total» o textos largos («Var.% 2026/2025») en las columnas resumen;
  · fila 5: fecha del mes en las columnas mensuales;
  · columna A desde la fila 7: rótulos (TOTAL, tipos en mayúsculas, grupos y productos), hasta las notas.
Solo se leen las columnas con fecha en la fila 5 y menos de 5 caracteres en la fila 4 (la «Var.%» también tiene fecha) y
nunca las «Ene.-Mes» del final: los acumulados se calculan aquí por meses.
"""
import json
import re
import unicodedata
from datetime import datetime

import pandas as pd
from openpyxl import load_workbook

from . import db

DEPTOS = {"LORETO": "16", "SAN MARTIN": "22", "SAN MARTÍN": "22", "UCAYALI": "25"}
NOMBRE = {"16": "LORETO", "22": "SAN MARTÍN", "25": "UCAYALI"}
COLUMNAS = ["AÑO", "MES", "DEPARTAMENTO", "ORDEN", "NIVEL", "TIPO", "GRUPO", "RUBRO", "FOB"]
# rótulos que en el ESTAD agrupan productos (el resto de rótulos en minúsculas son productos de su grupo)
GRUPOS = {"petroleo crudo y derivados", "forestales", "maderas y sus manufacturas", "pesqueros", "agropecuarios",
          "animales vivos", "resto", "agricolas (cafe)", "agricolas"}
_RE_TITULO = re.compile(r"^\s*(LORETO|SAN MART[IÍ]N|UCAYALI)\s*:.*\b(EXPORTACIONES|IMPORTACIONES)\b", re.I)


def _clave(s):
    s = unicodedata.normalize("NFKD", str(s)).encode("ascii", "ignore").decode().lower()
    return " ".join(s.split())


def limpiar(rotulo):
    """«Resto2/» -> «Resto», «OTRAS EXPORTACIONES3/» -> «OTRAS EXPORTACIONES», «Pesqueros (2)» -> «Pesqueros»."""
    t = " ".join(str(rotulo).split())
    t = re.sub(r"\s*\d/\s*$", "", t)
    t = re.sub(r"\s*\(\d\)\s*$", "", t)
    return t.strip()


def hojas(ruta):
    """[(hoja, flujo 'expo'|'impo', depto)] de un libro ESTAD; vacío si no lo es."""
    wb = load_workbook(ruta, read_only=True, data_only=True)
    try:
        salida = []
        for ws in wb.worksheets:
            filas = list(ws.iter_rows(min_row=1, max_row=7, max_col=1, values_only=True))
            if len(filas) < 7:
                continue
            m = _RE_TITULO.match(str(filas[1][0] or ""))
            total = _clave(filas[6][0] or "")
            if m and total.startswith("total"):
                flujo = "expo" if m.group(2).upper().startswith("EXP") else "impo"
                salida.append((ws.title, flujo, DEPTOS[m.group(1).upper()]))
        return salida
    finally:
        wb.close()


def reconocer(ruta):
    try:
        return hojas(ruta) or None
    except Exception:
        return None


def leer_hoja(ruta, hoja, depto):
    """Una hoja ESTAD -> (DataFrame en largo con COLUMNAS, notas al pie)."""
    wb = load_workbook(ruta, read_only=True, data_only=True)
    try:
        filas = [list(r) for r in wb[hoja].iter_rows(values_only=True)]
    finally:
        wb.close()
    f4, f5 = filas[3], filas[4]
    meses = []                                        # (columna, año, mes)
    for j in range(1, len(f5)):
        am = (f5[j].year, f5[j].month) if isinstance(f5[j], datetime) else None
        txt4 = "" if j >= len(f4) or f4[j] is None else str(f4[j])
        if am and len(txt4) < 5:
            meses.append((j, *am))
    registros, notas_pie, tipo, grupo, orden = [], [], "", "", 0
    fin = False
    for i in range(6, len(filas)):
        rot = filas[i][0] if filas[i] else None
        if rot is None or str(rot).strip() == "":
            continue
        t = " ".join(str(rot).split())
        if fin or re.match(r"^(\d/|Fuente|Elaboraci|Nota)", t, re.I):
            fin = True
            if re.match(r"^(\d/|Fuente)", t, re.I):
                notas_pie.append(t)
            continue
        r = limpiar(t)
        k = _clave(r)
        if k.startswith("total"):
            nivel, tipo, grupo = "total", "", ""
        elif r == r.upper():
            nivel, tipo, grupo = "tipo", r, ""          # PRODUCTOS TRADICIONALES, OTRAS EXPORTACIONES, BIENES DE CONSUMO…
        elif k in GRUPOS or k.split(" (")[0] in GRUPOS:
            nivel, grupo = "grupo", r
        else:
            nivel = "producto"
        orden += 1
        for j, a, m in meses:
            v = filas[i][j] if j < len(filas[i]) else None
            v = float(v) if isinstance(v, (int, float)) and not isinstance(v, bool) else 0.0
            registros.append((a, m, NOMBRE[depto], orden, nivel, tipo, grupo, r, v))
    return pd.DataFrame(registros, columns=COLUMNAS), notas_pie


def leer(archivos):
    """archivos: [(n_archivo, ruta, [(hoja, flujo, depto)])] -> ({base_id: DataFrame}, avisos)."""
    partes, avisos = {"expo-estad": [], "impo-estad": []}, []
    for n, ruta, encontradas in archivos:
        for hoja, flujo, depto in encontradas:
            df, pie = leer_hoja(ruta, hoja, depto)
            if df.empty:
                avisos.append(f"{ruta.name} ({hoja}): no tiene meses con datos; se ignoró.")
                continue
            # los meses aún vacíos al final del ESTAD (columnas ya preparadas para el año) no se guardan
            con_datos = df.groupby(["AÑO", "MES"])["FOB"].apply(lambda s: s.abs().sum() > 0)
            ultimo = max((a, m) for (a, m), ok in con_datos.items() if ok) if con_datos.any() else None
            if ultimo:
                df = df[(df["AÑO"] * 100 + df["MES"]) <= ultimo[0] * 100 + ultimo[1]]
            df = df.assign(_anio=df["AÑO"], _mes=df["MES"], _depto=depto, _nivel="", _archivo=n)
            base = "expo-estad" if flujo == "expo" else "impo-estad"
            partes[base].append(df)
            guardar_notas(base, depto, pie)
            avisos.append(f"{ruta.name} ({hoja}): {'exportaciones' if flujo == 'expo' else 'importaciones'} de {NOMBRE[depto].title()}, "
                          f"{df['ORDEN'].nunique()} rubros, {df['AÑO'].min()}-{df['MES'].min():02d} a {ultimo[0]}-{ultimo[1]:02d}.")
    return {b: pd.concat(p, ignore_index=True) for b, p in partes.items() if p}, avisos


def guardar_notas(base, depto, notas):
    """Las notas al pie del ESTAD (qué incluye «Resto» y «Otras») van a los cuadros."""
    db.guardar_documento(f"notas:{base}:{depto}", json.dumps(notas, ensure_ascii=False), datetime.now().isoformat(timespec="seconds"))


def notas(base, depto):
    d = db.documento(f"notas:{base}:{depto}")
    return json.loads(d[0]) if d else []
