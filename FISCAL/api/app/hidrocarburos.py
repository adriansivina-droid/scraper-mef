"""Hidrocarburos: producción fiscalizada por lote y canon petrolero (PERUPETRO).

Entran los CSV que el flujo diario de HIDROCARBUROS/ (PP_Produccion_Lotes.ipynb) deja en la carpeta DATA del almacén de
la página (almacen.py; antes, en Drive):
  · PP_produccion_lote_mes_<ini>-<fin>.csv      petróleo, LGN y condensado (BARRILES, BPD) por año, mes y lote;
  · PP_produccion_gas_lote_mes_<ini>-<fin>.csv  gas natural (MPC, MMPCD) por año, mes y lote;
  · PP_canon_departamento_mes.csv               canon y sobrecanon transferido por departamento (MM S/);
  · PP_produccion_diaria.csv                    producción diaria por lote: la suma del mes cubre los meses que la
    estadística mensual aún no publica (va con dos o tres meses de rezago), y el mes en curso sale como parcial.
El API solo los lee (almacen.sincronizar, al terminar el flujo o con el botón, o subiéndolos en la página).

También arma lo de la síntesis de Loreto y Ucayali: producción del mes y acumulada por producto y lote, contra el
mismo periodo del año anterior; la producción diaria promedio de 25 meses; el anexo mensual por lote (13 meses) y el
canon transferido al departamento.
"""
import calendar
import io
import re
import unicodedata

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font

from .catalogo import DEPARTAMENTOS

PRODUCCION = ["AÑO", "MES", "DEPARTAMENTO", "ZONA", "OPERADOR", "LOTE", "LOTE_HISTORICO", "PRODUCTO", "UNIDAD",
              "VOLUMEN", "PROMEDIO_DIARIO"]
CANON = ["AÑO", "MES", "DEPARTAMENTO", "CANON_MM_SOLES"]
SIN_DEPTO = "99"
# si el CSV no trae el departamento (antes de cruzarlo con el mapa de lotes), se usa el de estos lotes
LOTES_DEPTO = {"1-AB": "16", "192": "16", "1-AB/192": "16", "8": "16", "95": "16", "67": "16", "39": "16", "64": "16",
               "102": "16", "31-C": "25", "131": "25"}
# producto: (nombre en el cuadro, unidad del cuadro, divisor del volumen, unidad del promedio diario)
PRODUCTOS = {"PETROLEO": ("Petróleo", "miles de barriles", 1000, "miles de barriles diarios"),
             "LGN": ("Líquidos de gas natural", "miles de barriles", 1000, "miles de barriles diarios"),
             "CONDENSADO": ("Condensado", "miles de barriles", 1000, "miles de barriles diarios"),
             "GAS NATURAL": ("Gas natural", "millones de pies cúbicos", 1000, "millones de pies cúbicos diarios")}
MESES = ["Enero", "Febrero", "Marzo", "Abril", "Mayo", "Junio", "Julio", "Agosto", "Setiembre", "Octubre", "Noviembre", "Diciembre"]
MESES_CORTOS = ["Ene.", "Feb.", "Mar.", "Abr.", "May.", "Jun.", "Jul.", "Ago.", "Set.", "Oct.", "Nov.", "Dic."]
FUENTE = "Fuente: PERUPETRO, producción fiscalizada y canon petrolero. Elaboración: BCRP, Sucursal Iquitos. Departamento de Estudios Económicos."

_RE_LIQ = re.compile(r"^PP_produccion_lote_mes_\d{4}-\d{4}\.csv$", re.I)
_RE_GAS = re.compile(r"^PP_produccion_gas_lote_mes_\d{4}-\d{4}\.csv$", re.I)
_RE_CANON = re.compile(r"^PP_canon_departamento_mes\.csv$", re.I)
_RE_DIARIA = re.compile(r"^PP_produccion_diaria\.csv$", re.I)
MENSUAL = "ESTADÍSTICA MENSUAL"
_RE_PARCIAL = re.compile(r"(\d+) de (\d+) días")


def _clave(t):
    t = unicodedata.normalize("NFKD", str(t)).encode("ascii", "ignore").decode()
    return re.sub(r"[^A-Z]", "", t.upper())


_CODIGO = {_clave(n): c for c, n in DEPARTAMENTOS.items() if c <= "25"}


def tipo(nombre):
    """'liquidos', 'gas', 'canon', 'diaria' o None, según el nombre que pone el notebook de producción."""
    for t, r in (("liquidos", _RE_LIQ), ("gas", _RE_GAS), ("canon", _RE_CANON), ("diaria", _RE_DIARIA)):
        if r.match(nombre):
            return t
    return None


def _diaria(df):
    """Suma del mes por lote y producto (barriles y Mcf = MPC). FUENTE dice cuántos días trae el mes."""
    df = df[_texto(df["UNIDAD"]).str.upper().isin(["BBL", "MCF"])].copy()
    df["PRODUCTO"] = _texto(df["PRODUCTO"]).str.upper().replace({"PETROLEO Y CONDENSADO": "PETROLEO"})
    df["VALOR"] = pd.to_numeric(df["VALOR"], errors="coerce").fillna(0)
    for c in ("DEPARTAMENTO", "OPERADOR"):
        df[c] = _texto(df[c]) if c in df.columns else ""
    df["LOTE"] = _texto(df["LOTE"])
    dias = df.groupby(["AÑO", "MES"])["DIA"].nunique()
    g = df.groupby(["AÑO", "MES", "DEPARTAMENTO", "OPERADOR", "LOTE", "PRODUCTO", "UNIDAD"], as_index=False)["VALOR"].sum()
    g["DIAS"] = [int(dias[(a, m)]) for a, m in zip(g["AÑO"], g["MES"])]
    g["FUENTE"] = [f"PRODUCCIÓN DIARIA · {d} de {calendar.monthrange(int(a), int(m))[1]} días" for a, m, d in zip(g["AÑO"], g["MES"], g["DIAS"])]
    gas = g["UNIDAD"].str.upper() == "MCF"
    return pd.DataFrame({"AÑO": g["AÑO"], "MES": g["MES"], "DEPARTAMENTO": g["DEPARTAMENTO"], "ZONA": "", "OPERADOR": g["OPERADOR"],
                         "LOTE": g["LOTE"], "LOTE_HISTORICO": g["LOTE"], "PRODUCTO": g["PRODUCTO"],
                         "MPC": g["VALOR"].where(gas), "BARRILES": g["VALOR"].where(~gas),
                         "DIARIO": g["VALOR"] / g["DIAS"], "FUENTE": g["FUENTE"], "_gas": gas})


def _depto(texto, lote=""):
    """Código del departamento: el primero de «UCAYALI / HUANUCO»; si falta, el del lote; si no, 99."""
    primero = str(texto or "").split("/")[0] if pd.notna(texto) else ""
    return _CODIGO.get(_clave(primero)) or LOTES_DEPTO.get(str(lote)) or SIN_DEPTO


def _texto(serie):
    return serie.fillna("").astype(str).str.strip()


def _produccion(df, gas, n, nivel):
    """Columnas de la base a partir del CSV mensual o de la suma de la diaria. Cada nivel (líquidos, gas y sus versiones
    diarias) se reemplaza por separado."""
    volumen, diario = ("MPC", "MMPCD") if gas else ("BARRILES", "BPD")
    df = df.copy()
    for c in ("DEPARTAMENTO", "ZONA", "OPERADOR", "LOTE_HISTORICO", diario):
        if c not in df.columns:
            df[c] = None
    df["LOTE"] = _texto(df["LOTE"])
    df["LOTE_HISTORICO"] = _texto(df["LOTE_HISTORICO"]).where(_texto(df["LOTE_HISTORICO"]) != "", df["LOTE"])
    df["_depto"] = [_depto(d, l) for d, l in zip(df["DEPARTAMENTO"], df["LOTE_HISTORICO"])]
    df["DEPARTAMENTO"] = df["_depto"].map(lambda c: DEPARTAMENTOS.get(c, "SIN DEPARTAMENTO"))
    df["VOLUMEN"] = pd.to_numeric(df[volumen], errors="coerce").fillna(0)
    df["PROMEDIO_DIARIO"] = pd.to_numeric(df[diario], errors="coerce")
    df["UNIDAD"] = "MPC" if gas else "BARRILES"
    df["PRODUCTO"] = _texto(df["PRODUCTO"]).str.upper()
    df["ZONA"], df["OPERADOR"] = _texto(df["ZONA"]), _texto(df["OPERADOR"])
    df["_nivel"] = nivel
    return df.assign(_archivo=n)


def leer(archivos):
    """archivos: [(n_archivo, ruta, tipo)] -> ({base_id: DataFrame con columnas internas}, avisos)."""
    produccion, canon, avisos = [], [], []
    for n, ruta, t in archivos:
        df = pd.read_csv(ruta, encoding="utf-8-sig", dtype={"LOTE": str, "LOTE_HISTORICO": str})
        if t == "canon":
            df = df[_texto(df["DEPARTAMENTO"]).str.upper() != "TOTAL"].copy()
            df["_depto"] = df["DEPARTAMENTO"].map(_depto)
            df["CANON_MM_SOLES"] = pd.to_numeric(df["CANON_MM_SOLES"], errors="coerce").fillna(0)
            df["_nivel"] = "CANON"
            canon.append(df.assign(_archivo=n))
            continue
        if t == "diaria":
            if not {"AÑO", "MES", "DIA", "LOTE", "PRODUCTO", "UNIDAD", "VALOR"} <= set(df.columns):
                avisos.append(f"{ruta.name}: no tiene las columnas de la producción diaria; no se cargó.")
                continue
            d = _diaria(df)
            for es_gas, parte in d.groupby("_gas"):
                parte = parte.drop(columns="_gas").rename(columns={"DIARIO": "MMPCD" if es_gas else "BPD"})
                produccion.append(_produccion(parte, bool(es_gas), n, "DIARIA_GAS" if es_gas else "DIARIA_LIQUIDOS"))
            continue
        gas = t == "gas"
        volumen = "MPC" if gas else "BARRILES"
        faltan = [c for c in ("AÑO", "MES", "LOTE", "PRODUCTO", volumen) if c not in df.columns]
        if faltan:
            avisos.append(f"{ruta.name}: faltan las columnas {', '.join(faltan)}; no se cargó.")
            continue
        produccion.append(_produccion(df.assign(FUENTE=MENSUAL), gas, n, "GAS" if gas else "LIQUIDOS"))
    bases = {}
    for base_id, partes in (("hidro-produccion", produccion), ("hidro-canon", canon)):
        if partes:
            df = pd.concat(partes, ignore_index=True)
            df["AÑO"], df["MES"] = df["AÑO"].astype(int), df["MES"].astype(int)
            df["_anio"], df["_mes"] = df["AÑO"], df["MES"]
            bases[base_id] = df
    sin = sorted(set(bases["hidro-produccion"].query("_depto == @SIN_DEPTO")["LOTE"])) if "hidro-produccion" in bases else []
    if sin:
        avisos.append(f"Hidrocarburos: {len(sin)} lotes sin departamento (zócalo u otros; se guardan como «SIN DEPARTAMENTO»): "
                      f"{', '.join(sin[:12])}{'…' if len(sin) > 12 else ''}.")
    return bases, avisos


# ── Lo que publica la síntesis ──
def etiqueta(p, corta=False):
    return f"{(MESES_CORTOS if corta else MESES)[p % 100 - 1]} {p // 100}"


def _var(a, b):
    return None if not a else round((b / a - 1) * 100, 1)


def _lote(l):
    return l if l.upper().startswith("LOTE") else f"Lote {l}"


def cuadros(prod, canon, periodo=None):
    """prod / canon: filas de un departamento. periodo AAAA-MM (por defecto, el último mes con producción)."""
    if prod.empty:
        return None
    prod, parcial, de_diaria = _unir(prod)
    ultimo = int(prod["P"].max())
    p = int(periodo.replace("-", "")) if periodo else ultimo
    if p not in set(prod["P"]):
        raise ValueError(f"No hay producción de hidrocarburos para {etiqueta(p)} (último mes completo: {etiqueta(ultimo)}).")
    a, m = p // 100, p % 100
    ventanas = {"mes_anterior": [p - 100], "mes": [p],
                "acum_anterior": list(range((a - 1) * 100 + 1, p - 100 + 1)), "acum": list(range(a * 100 + 1, p + 1))}
    productos = [k for k in PRODUCTOS if k in set(prod["PRODUCTO"])]
    bloques, grafico, anexo = [], None, []
    for k in productos:
        nombre, unidad, div, unidad_diaria = PRODUCTOS[k]
        d = prod[prod["PRODUCTO"] == k]
        t = d.pivot_table(index="P", columns="LOTE_C", values="VOLUMEN", aggfunc="sum").fillna(0) / div
        suma = lambda col, ps: float(t.loc[t.index.isin(ps), col].sum()) if col in t.columns else 0.0
        tot = lambda ps: float(t.loc[t.index.isin(ps)].sum().sum())
        # lotes con producción en alguno de los periodos del cuadro, el de mayor producción del año primero
        lotes = sorted([l for l in t.columns if any(suma(l, ps) for ps in ventanas.values())],
                       key=lambda l: (-suma(l, ventanas["acum"]), -suma(l, ventanas["acum_anterior"]), l))
        if not lotes:
            continue
        filas = []
        for l in lotes + ["Total"]:
            v = {w: (tot(ps) if l == "Total" else suma(l, ps)) for w, ps in ventanas.items()}
            filas.append({"lote": l, **{w: round(x, 3) for w, x in v.items()},
                          "var_mes": _var(v["mes_anterior"], v["mes"]), "var_acum": _var(v["acum_anterior"], v["acum"])})
        bloques.append({"producto": nombre, "unidad": unidad, "filas": filas})
        periodos = [q for q in t.index if q <= p]
        if grafico is None:                                   # el producto principal: petróleo, si hay
            grafico = {"producto": nombre, "unidad": unidad_diaria,
                       "serie": [{"periodo": etiqueta(q, True), "valor": round(float(t.loc[q].sum()) / calendar.monthrange(q // 100, q % 100)[1], 2),
                                  "var": _var(float(t.loc[q - 100].sum()) if q - 100 in t.index else None, float(t.loc[q].sum()))}
                                 for q in periodos[-25:]]}
            # el mes en curso (días que van de la producción diaria): promedio de los días reportados
            if parcial and p == ultimo:
                x = parcial["filas"][parcial["filas"]["PRODUCTO"] == k]
                if not x.empty:
                    grafico["serie"].append({"periodo": etiqueta(parcial["P"], True), "valor": round(float(x["VOLUMEN"].sum()) / div / parcial["dias"], 2),
                                             "var": None, "parcial": True})
                    grafico["serie"] = grafico["serie"][-25:]
        anexo.append({"producto": nombre, "unidad": unidad, "lotes": lotes,
                      "filas": [{"periodo": etiqueta(q, True), **{l: round(float(t.at[q, l]), 3) for l in lotes},
                                 "Total": round(float(t.loc[q].sum()), 3)} for q in periodos[-13:]]})
    return {"periodo": f"{a}-{m:02d}", "ultimo": f"{ultimo // 100}-{ultimo % 100:02d}", "titulo_mes": MESES[m - 1],
            "anio": a, "bloques": bloques, "grafico": grafico, "anexo": anexo, "canon": _canon(canon),
            "de_diaria": [etiqueta(q, True) for q in de_diaria if q <= p],
            "parcial": None if not parcial else {"periodo": etiqueta(parcial["P"]), "dias": parcial["dias"], "de": parcial["de"]}}


def _unir(prod):
    """Estadística mensual primero; la suma de la producción diaria solo para los meses y productos que aquella aún no
    trae. El mes en curso (incompleto) se separa: no entra en el cuadro ni en el anexo, solo en el gráfico."""
    prod = prod.assign(P=prod["AÑO"] * 100 + prod["MES"])
    fuente = prod["FUENTE"].fillna(MENSUAL).astype(str) if "FUENTE" in prod.columns else pd.Series(MENSUAL, index=prod.index)
    es_diaria = fuente.str.startswith("PRODUCCIÓN DIARIA")
    mensual, diaria = prod[~es_diaria], prod[es_diaria].copy()
    # en la diaria el 192 sale como «192»; en la mensual, como «1-AB/192»
    historico = dict(zip(mensual["LOTE"], mensual["LOTE_HISTORICO"]))
    diaria["LOTE_HISTORICO"] = diaria["LOTE"].map(historico).fillna(diaria["LOTE_HISTORICO"])
    cubiertos = set(zip(mensual["P"], mensual["PRODUCTO"]))
    diaria = diaria.loc[pd.Series([(q, k) not in cubiertos for q, k in zip(diaria["P"], diaria["PRODUCTO"])], index=diaria.index, dtype=bool)]
    dias = diaria["FUENTE"].astype(str).str.extract(_RE_PARCIAL).astype(float)
    incompleta = (dias[0] < dias[1]).fillna(False)
    parcial = None
    if incompleta.any():
        q = int(diaria.loc[incompleta, "P"].max())
        filas = diaria[incompleta & (diaria["P"] == q)]
        d = dias[incompleta & (diaria["P"] == q)].iloc[0]
        parcial = {"P": q, "dias": int(d[0]), "de": int(d[1]), "filas": filas}
    unida = pd.concat([mensual, diaria[~incompleta]], ignore_index=True)
    unida["LOTE_C"] = unida["LOTE_HISTORICO"].map(_lote)
    return unida, parcial, sorted(set(diaria.loc[~incompleta, "P"]))


def _canon(canon):
    """Canon y sobrecanon transferido al departamento: los últimos 13 meses con su variación y el acumulado del año."""
    if canon is None or canon.empty:
        return None
    t = canon.assign(P=canon["AÑO"] * 100 + canon["MES"]).groupby("P")["CANON_MM_SOLES"].sum()
    q = int(t.index.max())
    a = q // 100
    acum = lambda anio: float(t[(t.index >= anio * 100 + 1) & (t.index <= anio * 100 + q % 100)].sum())
    return {"ultimo": f"{a}-{q % 100:02d}", "titulo_mes": MESES[q % 100 - 1], "anio": a,
            "acum_anterior": round(acum(a - 1), 2), "acum": round(acum(a), 2), "var_acum": _var(acum(a - 1), acum(a)),
            "meses": [{"periodo": etiqueta(x, True), "valor": round(float(t[x]), 2),
                       "var": _var(float(t[x - 100]) if x - 100 in t.index else None, float(t[x]))} for x in list(t.index)[-13:]]}


def excel(region, c):
    """Cuadro de producción, gráfico, anexo por lote y canon en un Excel, con títulos y fuente como en la síntesis."""
    wb = Workbook()
    neg, centro = Font(bold=True), Alignment(horizontal="center")
    a, mes = c["anio"], c["titulo_mes"]
    ws = wb.active
    ws.title = "Cuadro"
    ws.append([f"{region.upper()}: PRODUCCIÓN DE HIDROCARBUROS"]); ws.append([])
    ws.append(["Producto / lote", f"{mes} {a - 1}", f"{mes} {a}", "Var. %", f"Ene.-{mes[:3]}. {a - 1}", f"Ene.-{mes[:3]}. {a}", "Var. %"])
    for celda in ws[3]:
        celda.font, celda.alignment = neg, centro
    for b in c["bloques"]:
        ws.append([f"{b['producto']} ({b['unidad']})"]); ws.cell(ws.max_row, 1).font = neg
        for f in b["filas"]:
            ws.append([f["lote"], f["mes_anterior"], f["mes"], f["var_mes"], f["acum_anterior"], f["acum"], f["var_acum"]])
            for i, celda in enumerate(ws[ws.max_row]):
                celda.number_format = "0.0" if i in (3, 6) else "#,##0.0"
                if f["lote"] == "Total":
                    celda.font = neg
    ws.append([]); ws.append([FUENTE])
    ws["A1"].font = neg
    for col, ancho in zip("ABCDEFG", [30, 13, 13, 9, 15, 15, 9]):
        ws.column_dimensions[col].width = ancho

    if c["grafico"]:
        g = c["grafico"]
        wg = wb.create_sheet("Gráfico")
        wg.append([f"{region.upper()}: PRODUCCIÓN DE {g['producto'].upper()}"]); wg.append([f"({g['unidad'].capitalize()})"]); wg.append([])
        wg.append(["Mes", "Promedio diario", "Var. % interanual"])
        for x in g["serie"]:
            wg.append([x["periodo"], x["valor"], x["var"]])
        wg["A1"].font = neg

    wa = wb.create_sheet("Anexo")
    wa.append([f"{region.upper()}: PRODUCCIÓN DE HIDROCARBUROS POR LOTE"]); wa.append([])
    for b in c["anexo"]:
        wa.append([f"{b['producto']} ({b['unidad']})"]); wa.cell(wa.max_row, 1).font = neg
        wa.append(["Mes", *b["lotes"], "Total"])
        for celda in wa[wa.max_row]:
            celda.font, celda.alignment = neg, centro
        for f in b["filas"]:
            wa.append([f["periodo"], *[f[l] for l in b["lotes"]], f["Total"]])
            for celda in wa[wa.max_row][1:]:
                celda.number_format = "#,##0.0"
        wa.append([])
    wa["A1"].font = neg

    if c["canon"]:
        k = c["canon"]
        wc = wb.create_sheet("Canon")
        wc.append([f"{region.upper()}: CANON Y SOBRECANON PETROLERO TRANSFERIDO"]); wc.append(["(Millones de soles)"]); wc.append([])
        wc.append(["Mes", "Millones de S/", "Var. %"])
        for x in k["meses"]:
            wc.append([x["periodo"], x["valor"], x["var"]])
        wc.append([f"Ene.-{k['titulo_mes'][:3]}. {k['anio'] - 1}", k["acum_anterior"]])
        wc.append([f"Ene.-{k['titulo_mes'][:3]}. {k['anio']}", k["acum"], k["var_acum"]])
        wc["A1"].font = neg
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
