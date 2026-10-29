"""Exportaciones desde la base mensual de SUNAT: el flujo de ANUAL_PRO (hoja ORIGEN) en la página.

  1. Se sube la base de SUNAT (X_AAAAMM.xlsx: aduana, ubigeo, ruc11, ano_mes, fecembar, pais, partorig, dcom, fob, …).
  2. Cada fila se completa con los diccionarios (los de la hoja UNICOS y la hoja PARTIDA de ANUAL_PRO, guardados en la
     página): tipo, categoría y producto por partida; producto corregido por descripción (PROD_UNICOS); país y
     continente; aduana; empresa y BORRAR PARA por RUC.
  3. Estado de cada fila:
       · ok         BORRAR PARA = OK (solo estas entran a los cuadros);
       · verificar  empresa nueva o con BORRAR PARA vacío (falta revisarla);
       · borrar     BORRAR PARA con otra cosa que OK (p. ej. «LOR SMN UCA») o una regla de limpieza (p. ej. café en Loreto);
       · clasificar la partida no tiene tipo ni producto en el diccionario.
  4. Cada fila ok va a una fila del ESTAD de su región (el último ESTAD cargado): por categoría al grupo (maderas,
     pesqueros, animales vivos, agropecuarios, resto), dentro del grupo al producto con el mismo nombre (o el de la tabla
     de equivalencias); lo que no tiene fila propia va a «Otros acabados» (maderas) o a «Resto». Los tipo 3 van a
     «OTRAS EXPORTACIONES».
  5. El «ESTAD propuesto» de un mes es esa tabla dinámica; se compara con el ESTAD cargado y se puede pasar al ESTAD.
Subir ANUAL_PRO_AAAA.xlsx carga los diccionarios y además su hoja ORIGEN como base de ese año.
"""
import io
import re
import unicodedata
from decimal import ROUND_HALF_UP, Decimal

import pandas as pd
from openpyxl import load_workbook
from sqlalchemy import text

from . import db

CORTO = {"16": "LOR", "22": "SMN", "25": "UCA"}
NOMBRE = {"16": "LORETO", "22": "SAN MARTÍN", "25": "UCAYALI"}
CRUDAS = ["aduana", "ubigeo", "ruc11", "ano_mes", "fecembar", "pais", "partorig", "dcom", "fob", "kilos", "qunifis", "tunifis", "ubigeo_dep"]
COLUMNAS = ["AÑO", "MES", "DEPARTAMENTO", "RUC", "EMPRESA", "PARTIDA", "DESCRIPCION", "PRODUCTO", "TIPO", "CATEGORIA",
            "TIPO_ESTAD", "RUBRO_ESTAD", "COD_PAIS", "PAIS", "CONTINENTE", "COD_ADUANA", "ADUANA", "FECHA_EMBARQUE", "KILOS",
            "CANTIDAD", "UNIDAD", "BORRAR_PARA", "ESTADO", "MOTIVO", "FOB_USD", "FOB"]
# grupo del ESTAD de cada categoría de PARTIDA (el primero que exista en la región)
GRUPO_DE = {"2.1": ["forestales", "maderas y sus manufacturas"], "2.2": ["pesqueros"], "2.3": ["animales vivos"],
            "2.4": ["agropecuarios"], "2.5": ["resto"], "1.2": ["petroleo crudo y derivados"], "1.1": ["agricolas (cafe)", "agricolas"]}
# reglas obligatorias de limpieza (se pueden cambiar en la página)
REGLAS_INICIALES = [
    {"departamento": "LOR", "campo": "partida", "valor": "0901", "motivo": "Loreto no exporta café: lo embarca una empresa de otra región."},
    {"departamento": "LOR", "campo": "partida", "valor": "1404902000", "motivo": "Tara en polvo: no es de Loreto."},
]
# productos con otro nombre en el ESTAD (además del parecido por nombre)
EQUIVALENCIAS_INICIALES = [
    {"departamento": "SMN", "producto": p, "rubro": "Cacao (derivados)"} for p in ("Cacao en manteca", "Cacao en pasta", "Cacao en polvo", "Cacao residuos")
] + [{"departamento": "SMN", "producto": "Aceite vegetal de sacha inchi", "rubro": "Aceite vegetal"},
     {"departamento": "SMN", "producto": "Puros", "rubro": "Puros, cigarros y pack"},
     {"departamento": "SMN", "producto": "Tabaco", "rubro": "Tabaco negro y rubio"}]


def _clave(s):
    s = unicodedata.normalize("NFKD", str(s or "")).encode("ascii", "ignore").decode().lower()
    return " ".join(s.split())


def _r0(v):
    """ROUND(fob/1000, 0) de la hoja ORIGEN (la mitad hacia arriba)."""
    return float(Decimal(repr(float(v))).quantize(Decimal(1), rounding=ROUND_HALF_UP)) + 0.0


# ── Diccionarios y tablas de la página ──
TABLAS = {"partidas": ["PARTIDA", "TIPO", "CATEGORIA", "PRODUCTO", "DESCRIPCION", "NOTA"],
          "unicos": ["PRODUCTO_ORIGINAL", "PRODUCTO_CORREGIDO"],
          "empresas": ["RUC", "RAZON_SOCIAL", "DOMICILIO_FISCAL", "BORRAR_PARA", "NOTA", "REVISADO", "ESTADO_SUNAT", "CONDICION"],
          "paises": ["COD", "PAIS", "CONTINENTE"],
          "aduanas": ["COD", "ADUANA"],
          "reglas": ["departamento", "campo", "valor", "motivo"],
          "equivalencias": ["departamento", "producto", "rubro"]}


def _tabla(nombre):
    return f"expo_dic_{nombre}"


def leer_tabla(nombre):
    t = _tabla(nombre)
    if not db.existe(t):
        if nombre == "reglas":
            return pd.DataFrame(REGLAS_INICIALES, columns=TABLAS[nombre])
        if nombre == "equivalencias":
            return pd.DataFrame(EQUIVALENCIAS_INICIALES, columns=TABLAS[nombre])
        return pd.DataFrame(columns=TABLAS[nombre])
    with db.motor().connect() as cx:
        df = pd.read_sql(text(f'SELECT * FROM "{t}"'), cx)
    for c in TABLAS[nombre]:
        if c not in df.columns:
            df[c] = None
    return df[TABLAS[nombre]]


def guardar_tabla(nombre, df):
    df = df.reindex(columns=TABLAS[nombre]).astype(object).where(lambda x: x.notna(), None)
    with db.motor().begin() as cx:
        cx.execute(text(f'DROP TABLE IF EXISTS "{_tabla(nombre)}"'))
        df.to_sql(_tabla(nombre), cx, index=False, chunksize=5000)
    return len(df)


def actualizar_tabla(nombre, filas, clave):
    """Agrega o reemplaza filas (dicts) por su clave; una fila con «borrar»: True se quita."""
    df = leer_tabla(nombre)
    for f in filas:
        k = str(f.get(clave) or "").strip()
        if not k:
            continue
        df = df[df[clave].astype(str) != k]
        if not f.get("borrar"):
            df = pd.concat([df, pd.DataFrame([{c: f.get(c) for c in TABLAS[nombre]}])], ignore_index=True)
    guardar_tabla(nombre, df)
    return df


def _columnas_unicos(hoja):
    """La hoja UNICOS tiene varias tablas lado a lado (PROD_UNICOS, PAISES, EMPRESAS, ADUANA): se ubican por encabezado."""
    cab = [str(c).strip() if c is not None else "" for c in hoja.columns]
    def bloque(primero, siguiente, n):
        for i in range(len(cab) - 1):
            if cab[i].split(".")[0] == primero and cab[i + 1].split(".")[0] == siguiente:
                d = hoja.iloc[:, i:i + n].copy()
                d.columns = [c.split(".")[0] for c in cab[i:i + n]]
                return d
        return None
    return bloque("PRODUCTO_ORIGINAL", "PRODUCTO_CORREGIDO", 2), bloque("COD", "COD_TEXTO", 4), bloque("RUC", "RUC_TEXTO", 9), bloque("COD", "ADUANA_JURISDICCION", 2)


def es_anual_pro(ruta):
    try:
        wb = load_workbook(ruta, read_only=True)
        nombres = set(wb.sheetnames)
        wb.close()
        return {"UNICOS", "PARTIDA"} <= nombres
    except Exception:
        return False


def es_sunat(ruta):
    try:
        cab = pd.read_excel(ruta, sheet_name=0, nrows=0, dtype=str)
        return {"ruc11", "partorig", "dcom", "fob", "ubigeo_dep"} <= {str(c).strip().lower() for c in cab.columns}
    except Exception:
        return False


def cargar_diccionarios(ruta):
    """ANUAL_PRO: los diccionarios de UNICOS y PARTIDA. Empresas y productos únicos se suman a lo que ya tiene la página
    (lo del archivo gana); partidas, países y aduanas se reemplazan. Devuelve avisos."""
    u = pd.read_excel(ruta, sheet_name="UNICOS", dtype=str)
    prod, paises, emp, adu = _columnas_unicos(u)
    avisos = []
    if prod is not None:
        prod = prod.dropna(subset=["PRODUCTO_ORIGINAL"])
        viejo = leer_tabla("unicos")
        n = guardar_tabla("unicos", pd.concat([viejo[~viejo["PRODUCTO_ORIGINAL"].isin(prod["PRODUCTO_ORIGINAL"])], prod], ignore_index=True))
        avisos.append(f"Productos únicos (PROD_UNICOS): {n} descripciones.")
    if paises is not None:
        paises = paises.dropna(subset=["COD_TEXTO"]).rename(columns={"COD_TEXTO": "COD_"})
        guardar_tabla("paises", pd.DataFrame({"COD": paises["COD_"].str.zfill(3), "PAIS": paises["PAIS"], "CONTINENTE": paises["CONTINENTE"]}))
        avisos.append(f"Países: {len(paises)}.")
    if adu is not None:
        adu = adu.dropna(subset=["COD"])
        guardar_tabla("aduanas", pd.DataFrame({"COD": adu["COD"].astype(str).str.replace(r"\.0$", "", regex=True).str.lstrip("0").replace("", "0"),
                                               "ADUANA": adu["ADUANA_JURISDICCION"]}))
        avisos.append(f"Aduanas: {len(adu)}.")
    if emp is not None:
        emp = emp.dropna(subset=["RUC_TEXTO"])
        nuevo = pd.DataFrame({"RUC": emp["RUC_TEXTO"].str.strip(), "RAZON_SOCIAL": emp["RAZON_SOCIAL"], "DOMICILIO_FISCAL": emp["DOMICILIO_FISCAL"],
                              "BORRAR_PARA": emp["BORRAR_PARA"], "NOTA": emp["NOTA"], "REVISADO": emp["REVISADO"]})
        viejo = leer_tabla("empresas")
        guardar_tabla("empresas", pd.concat([viejo[~viejo["RUC"].astype(str).isin(nuevo["RUC"])], nuevo], ignore_index=True))
        avisos.append(f"Empresas (EMPRESAS): {len(nuevo)} RUC.")
    p = pd.read_excel(ruta, sheet_name="PARTIDA", dtype=str)
    p.columns = [str(c).strip() for c in p.columns]
    if "PARTIDA_TEXTO" in p.columns:
        prod_col = next((c for c in p.columns if c.startswith("PRODUCTO") and "✔" in c), None)
        tabla = pd.DataFrame({"PARTIDA": p["PARTIDA_TEXTO"].str.strip().str.zfill(10), "TIPO": p.get("|"), "CATEGORIA": p.get("CATEGORIA"),
                              "PRODUCTO": p[prod_col] if prod_col else None, "DESCRIPCION": p.get("PRODUCTO"), "NOTA": p.get("NOTA")})
        guardar_tabla("partidas", tabla.dropna(subset=["PARTIDA"]))
        avisos.append(f"Partidas (PARTIDA): {len(tabla)}, {int(tabla['TIPO'].notna().sum())} clasificadas.")
    return avisos


def leer_origen(ruta):
    """La hoja ORIGEN de ANUAL_PRO (encabezado en la fila 7): sus columnas crudas, como una base de SUNAT."""
    d = pd.read_excel(ruta, sheet_name="ORIGEN", header=6, dtype=str)
    d.columns = [str(c).strip().lower() for c in d.columns]
    return d[[c for c in CRUDAS if c in d.columns]]


def leer_crudo(ruta):
    d = pd.read_excel(ruta, sheet_name=0, dtype=str)
    d.columns = [str(c).strip().lower() for c in d.columns]
    return d[[c for c in CRUDAS if c in d.columns]]


# ── Del crudo a la base ──
def estructura_estad(depto):
    """Filas del último ESTAD cargado de la región: [{tipo, grupo, rubro, nivel, orden}]."""
    from .catalogo import BASES
    t = BASES["expo-estad"].tabla
    if not db.existe(t):
        return []
    with db.motor().connect() as cx:
        p = cx.execute(text(f'SELECT MAX(_anio * 100 + _mes) FROM "{t}" WHERE _depto = :d'), {"d": depto}).scalar()
        if p is None:
            return []
        df = pd.read_sql(text(f'SELECT "ORDEN", "NIVEL", "TIPO", "GRUPO", "RUBRO" FROM "{t}" WHERE _depto = :d AND (_anio * 100 + _mes) = :p'),
                         cx, params={"d": depto, "p": int(p)})
    df = df.drop_duplicates("ORDEN").sort_values("ORDEN")
    return [{"orden": int(r.ORDEN), "nivel": r.NIVEL, "tipo": r.TIPO or "", "grupo": r.GRUPO or "", "rubro": r.RUBRO} for r in df.itertuples()]


def _parecido(producto, rubro):
    """«Tabaco» ~ «Tabaco negro y rubio», «Manteca vegetal de palma» ~ «Manteca vegetal», «Pescado congelado» ~
    «Pescado congelado (paiche)»; pero «Cacao en grano» no es «Cacao (derivados)»."""
    a, b, sin_par = _clave(producto), _clave(rubro), _clave(re.sub(r"\s*\(.*\)$", "", rubro))
    return a in (b, sin_par) or a.startswith(b + " ") or b.startswith(a + " ")


def ubicar(est, corto, tipo, categoria, producto, equivalencias):
    """(tipo del ESTAD, rubro del ESTAD) de una fila ok; (None, None) si la región no tiene ESTAD."""
    if not est:
        return None, None
    tipos = [f for f in est if f["nivel"] == "tipo"]
    otras = next((f for f in tipos if _clave(f["rubro"]).startswith("otras")), None)
    nt = next((f for f in tipos if "no tradicionales" in _clave(f["rubro"])), None)
    tr = next((f for f in tipos if _clave(f["rubro"]) == "productos tradicionales"), None)
    eq = {(_clave(e["producto"])): e["rubro"] for e in equivalencias if str(e.get("departamento") or "").upper() in ("", corto)}
    # equivalencia explícita: el producto va a esa fila, esté donde esté (también los tipo 3)
    if _clave(producto) in eq:
        f = next((f for f in est if _clave(f["rubro"]) == _clave(eq[_clave(producto)]) and f["nivel"] != "total"), None)
        if f:
            return (f["rubro"], f["rubro"]) if f["nivel"] == "tipo" else (f["tipo"], f["rubro"])
    t = str(tipo or "")[:1]
    if t == "3" or not t:
        return (otras["rubro"], otras["rubro"]) if otras else (None, None)
    cat = str(categoria or "")[:3]
    prod_filas = [f for f in est if f["nivel"] == "producto"]
    # tradicionales: petróleo y café a su grupo; el café sin grupo propio va a «Café en grano» de agropecuarios
    candidatos = []
    for nombre in GRUPO_DE.get(cat, []):
        candidatos += [f for f in est if f["nivel"] == "grupo" and _clave(f["rubro"]) == nombre and
                       (t != "2" or (nt and f["tipo"] == nt["rubro"]))]
    if cat == "1.1" and not candidatos:
        f = next((f for f in prod_filas if _parecido(producto or "Café en grano", f["rubro"])), None)
        if f:
            return f["tipo"], f["rubro"]
    resto = next((f for f in est if f["nivel"] == "grupo" and _clave(f["rubro"]) == "resto"), None)
    if not candidatos:
        return (resto["tipo"], resto["rubro"]) if resto else ((nt or tr or otras)["rubro"], (nt or tr or otras)["rubro"])
    g = candidatos[0]
    hijos = [f for f in prod_filas if f["tipo"] == g["tipo"] and f["grupo"] == g["rubro"]]
    if not hijos:
        return g["tipo"], g["rubro"]
    f = next((f for f in hijos if _clave(producto) == _clave(f["rubro"])), None) or next((f for f in hijos if _parecido(producto, f["rubro"])), None)
    if f:
        return f["tipo"], f["rubro"]
    if cat == "2.1":
        f = next((f for f in hijos if _clave(f["rubro"]) == "otros acabados"), None)
        if f:
            return f["tipo"], f["rubro"]
    return (resto["tipo"], resto["rubro"]) if resto else (g["tipo"], g["rubro"])


def procesar(crudo):
    """crudo: columnas de SUNAT (texto) o la base ya guardada (se vuelve a procesar con los diccionarios del momento)."""
    if "ruc11" not in crudo.columns:                 # filas de la base: a columnas crudas
        crudo = pd.DataFrame({"aduana": crudo["COD_ADUANA"], "ruc11": crudo["RUC"], "ano_mes": None, "anio": crudo["AÑO"], "mes": crudo["MES"],
                              "fecembar": crudo["FECHA_EMBARQUE"], "pais": crudo["COD_PAIS"], "partorig": crudo["PARTIDA"], "dcom": crudo["DESCRIPCION"],
                              "fob": crudo["FOB_USD"], "kilos": crudo["KILOS"], "qunifis": crudo["CANTIDAD"], "tunifis": crudo["UNIDAD"],
                              "ubigeo_dep": crudo["DEPARTAMENTO"].map({v: k for k, v in NOMBRE.items()})})
    d = crudo.copy()
    d = d[d["ruc11"].notna() & d["fob"].notna()]
    d["ubigeo_dep"] = d["ubigeo_dep"].astype(str).str.replace(r"\.0$", "", regex=True).str.zfill(2)
    fuera = d[~d["ubigeo_dep"].isin(CORTO)]
    d = d[d["ubigeo_dep"].isin(CORTO)].copy()
    if "anio" not in d.columns or d["anio"].isna().all():
        am = d["ano_mes"].astype(str).str.replace(r"\.0$", "", regex=True).str.zfill(4)
        d["anio"], d["mes"] = 2000 + am.str[:2].astype(int), am.str[2:4].astype(int)
    d["anio"], d["mes"] = d["anio"].astype(int), d["mes"].astype(int)
    d["ruc11"] = d["ruc11"].astype(str).str.replace(r"\.0$", "", regex=True).str.strip()
    d["partorig"] = d["partorig"].astype(str).str.replace(r"\.0$", "", regex=True).str.strip().str.zfill(10)
    d["pais"] = d["pais"].astype(str).str.replace(r"\.0$", "", regex=True).str.zfill(3)
    d["aduana"] = d["aduana"].astype(str).str.replace(r"\.0$", "", regex=True).str.lstrip("0").replace("", "0")
    d["fob"] = pd.to_numeric(d["fob"], errors="coerce").fillna(0.0)

    part = leer_tabla("partidas").drop_duplicates("PARTIDA").set_index("PARTIDA")
    unicos = leer_tabla("unicos").drop_duplicates("PRODUCTO_ORIGINAL").set_index("PRODUCTO_ORIGINAL")["PRODUCTO_CORREGIDO"]
    emp = leer_tabla("empresas").drop_duplicates("RUC", keep="last")
    emp["RUC"] = emp["RUC"].astype(str)
    emp = emp.set_index("RUC")
    paises = leer_tabla("paises").drop_duplicates("COD").set_index("COD")
    aduanas = leer_tabla("aduanas").drop_duplicates("COD").set_index("COD")["ADUANA"]
    reglas = leer_tabla("reglas").to_dict("records")
    equivalencias = leer_tabla("equivalencias").to_dict("records")
    estructuras = {dep: estructura_estad(dep) for dep in CORTO}

    out = pd.DataFrame({"AÑO": d["anio"], "MES": d["mes"], "DEPARTAMENTO": d["ubigeo_dep"].map(NOMBRE), "RUC": d["ruc11"],
                        "PARTIDA": d["partorig"], "DESCRIPCION": d["dcom"].fillna(""), "COD_PAIS": d["pais"], "COD_ADUANA": d["aduana"],
                        "FECHA_EMBARQUE": d["fecembar"].astype(str).str.replace(r"\.0$", "", regex=True), "KILOS": pd.to_numeric(d["kilos"], errors="coerce"),
                        "CANTIDAD": pd.to_numeric(d["qunifis"], errors="coerce"), "UNIDAD": d["tunifis"], "FOB_USD": d["fob"]})
    out["_depto"] = d["ubigeo_dep"].values
    out["TIPO"] = out["PARTIDA"].map(part["TIPO"]) if len(part) else None
    out["CATEGORIA"] = out["PARTIDA"].map(part["CATEGORIA"]) if len(part) else None
    corregido = out["DESCRIPCION"].map(unicos) if len(unicos) else pd.Series(None, index=out.index)
    out["PRODUCTO"] = corregido.where(corregido.notna() & (corregido.astype(str).str.strip() != ""), out["PARTIDA"].map(part["PRODUCTO"]) if len(part) else None)
    out["PAIS"] = out["COD_PAIS"].map(paises["PAIS"]) if len(paises) else None
    out["CONTINENTE"] = out["COD_PAIS"].map(paises["CONTINENTE"]) if len(paises) else None
    out["ADUANA"] = out["COD_ADUANA"].map(aduanas) if len(aduanas) else None
    razon = out["RUC"].map(emp["RAZON_SOCIAL"]) if len(emp) else pd.Series(None, index=out.index)
    out["EMPRESA"] = out["RUC"] + " - " + razon.fillna("(sin razón social)").astype(str)
    out["BORRAR_PARA"] = out["RUC"].map(emp["BORRAR_PARA"]) if len(emp) else None
    out["FOB"] = [_r0(v / 1000) for v in out["FOB_USD"]]

    conocidas = set(emp.index)
    estado, motivo = [], []
    for ruc, bp, part_, desc, prod, dep, tipo in zip(out["RUC"], out["BORRAR_PARA"], out["PARTIDA"], out["DESCRIPCION"], out["PRODUCTO"], out["_depto"], out["TIPO"]):
        corto = CORTO[dep]
        regla = next((g for g in reglas if str(g.get("departamento") or "").upper() in ("", corto) and _cumple(g, ruc, part_, desc, prod)), None)
        bp_txt = "" if bp is None or (isinstance(bp, float) and pd.isna(bp)) else str(bp).strip()
        if regla:
            estado.append("borrar"); motivo.append(f"Regla de limpieza: {regla.get('motivo') or regla.get('valor')}")
        elif ruc not in conocidas:
            estado.append("verificar"); motivo.append("Empresa nueva: falta revisarla (BORRAR PARA)")
        elif bp_txt == "":
            estado.append("verificar"); motivo.append("BORRAR PARA vacío: falta verificar")
        elif bp_txt.upper() != "OK":
            estado.append("borrar"); motivo.append(f"BORRAR PARA: {bp_txt}")
        elif tipo is None or (isinstance(tipo, float) and pd.isna(tipo)) or str(tipo).strip() == "":
            estado.append("clasificar"); motivo.append("Partida sin tipo ni producto en el diccionario")
        else:
            estado.append("ok"); motivo.append("")
    out["ESTADO"], out["MOTIVO"] = estado, motivo
    ubic = [ubicar(estructuras[dep], CORTO[dep], t, c, p, equivalencias) if e == "ok" else (None, None)
            for dep, t, c, p, e in zip(out["_depto"], out["TIPO"], out["CATEGORIA"], out["PRODUCTO"], out["ESTADO"])]
    out["TIPO_ESTAD"], out["RUBRO_ESTAD"] = [u[0] for u in ubic], [u[1] for u in ubic]
    out["_anio"], out["_mes"], out["_nivel"] = out["AÑO"], out["MES"], ""
    return out[COLUMNAS + ["_anio", "_mes", "_depto", "_nivel"]].reset_index(drop=True), len(fuera)


def _cumple(regla, ruc, partida, descripcion, producto):
    campo, valor = str(regla.get("campo") or "").lower(), str(regla.get("valor") or "").strip()
    if not valor:
        return False
    if campo == "partida":
        return str(partida).startswith(valor)
    if campo == "ruc":
        return str(ruc) == valor
    if campo == "producto":
        return _clave(producto) == _clave(valor)
    if campo == "descripcion":
        return _clave(valor) in _clave(descripcion)
    return False


def leer(archivos):
    """archivos: [(n, ruta, 'sunat'|'anual')] -> ({"expo-sunat": df}, avisos)."""
    partes, avisos = [], []
    for n, ruta, tipo in archivos:
        if tipo == "anual":
            avisos += [f"{ruta.name}: {a}" for a in cargar_diccionarios(ruta)]
            crudo = leer_origen(ruta)
        else:
            crudo = leer_crudo(ruta)
        df, fuera = procesar(crudo)
        if fuera:
            avisos.append(f"{ruta.name}: {fuera} filas de otros departamentos; no se guardaron.")
        if df.empty:
            continue
        cuenta = df.groupby("ESTADO")["FOB"].agg(["count", "sum"])
        avisos.append(f"{ruta.name}: base mensual {df['AÑO'].min()}-{df['MES'].min():02d} a {df['AÑO'].max()}-{df['MES'].max():02d}, " +
                      ", ".join(f"{k} {int(v['count'])} filas ({v['sum']:,.0f} mil US$)".replace(",", " ") for k, v in cuenta.iterrows()) + ".")
        partes.append(df.assign(_archivo=n))
    return ({"expo-sunat": pd.concat(partes, ignore_index=True)} if partes else {}), avisos


def reprocesar():
    """Vuelve a procesar toda la base con los diccionarios y reglas del momento (tras una corrección en la página)."""
    base = db.leer_todo("expo-sunat")
    if base is None or base.empty:
        return 0
    df, _ = procesar(base)
    from .catalogo import BASES
    with db.motor().begin() as cx:
        cx.execute(text(f'DELETE FROM "{BASES["expo-sunat"].tabla}"'))
    return db.guardar("expo-sunat", df)


# ── Revisión, ESTAD propuesto y cuadros de texto ──
def _base(depto, desde="1990-01", hasta="2100-12"):
    return db.leer("expo-sunat", [depto], [], desde, hasta)


def propuesta(depto, anio, mes, df=None):
    """El ESTAD que sale de las filas ok del mes, fila por fila del último ESTAD de la región, y el ESTAD cargado."""
    est = estructura_estad(depto)
    if not est:
        return None
    df = _base(depto, f"{anio}-{mes:02d}", f"{anio}-{mes:02d}") if df is None else df
    ok = df[df["ESTADO"] == "ok"]
    directo = ok.groupby(["TIPO_ESTAD", "RUBRO_ESTAD"])["FOB"].sum().to_dict()
    valores = {}
    for f in est:                                    # productos, grupos sin productos y tipos sin grupos
        valores[f["orden"]] = directo.get((f["tipo"] if f["nivel"] != "tipo" else f["rubro"], f["rubro"]), 0.0)
    for f in est:
        if f["nivel"] == "grupo":
            hijos = [h for h in est if h["nivel"] == "producto" and h["tipo"] == f["tipo"] and h["grupo"] == f["rubro"]]
            if hijos:
                valores[f["orden"]] = sum(valores[h["orden"]] for h in hijos)
    for f in est:
        if f["nivel"] == "tipo":
            grupos = [g for g in est if g["nivel"] == "grupo" and g["tipo"] == f["rubro"]]
            if grupos:
                valores[f["orden"]] = sum(valores[g["orden"]] for g in grupos) + valores[f["orden"]]
    total = next((f for f in est if f["nivel"] == "total"), None)
    if total:
        valores[total["orden"]] = sum(valores[f["orden"]] for f in est if f["nivel"] == "tipo")
    cargado = db.leer("expo-estad", [depto], [], f"{anio}-{mes:02d}", f"{anio}-{mes:02d}")
    car = dict(zip(cargado["RUBRO"] + "|" + cargado["TIPO"].fillna("") + "|" + cargado["NIVEL"], cargado["FOB"])) if not cargado.empty else {}
    filas = []
    for f in est:
        v = valores[f["orden"]]
        c = car.get(f"{f['rubro']}|{f['tipo']}|{f['nivel']}")
        filas.append({**f, "sunat": v, "estad": None if c is None else float(c), "dif": None if c is None else round(v - float(c), 1)})
    return {"filas": filas, "hay_estad": bool(car)}


def pasar_al_estad(depto, anio, mes):
    p = propuesta(depto, anio, mes)
    if p is None:
        raise ValueError("Primero sube el ESTAD de la región: la propuesta usa sus filas.")
    df = pd.DataFrame([{"AÑO": anio, "MES": mes, "DEPARTAMENTO": NOMBRE[depto], "ORDEN": f["orden"], "NIVEL": f["nivel"], "TIPO": f["tipo"],
                        "GRUPO": f["grupo"], "RUBRO": f["rubro"], "FOB": float(f["sunat"])} for f in p["filas"]])
    df = df.assign(_anio=anio, _mes=mes, _depto=depto, _nivel="")
    return db.guardar("expo-estad", df)


def _top(ok, campo, n=5):
    tot = ok["FOB_USD"].sum()
    g = ok.assign(**{campo: ok[campo].fillna("(sin dato)")}).groupby(campo)["FOB_USD"].sum().sort_values(ascending=False)
    salida = []
    for k, v in g.head(n).items():
        if v <= 0:
            continue
        prods = ok[ok[campo].fillna("(sin dato)") == k].groupby("PRODUCTO")["FOB_USD"].sum().sort_values(ascending=False).head(3)
        salida.append({"nombre": k, "fob": float(v) / 1000, "pct": float(_r1(v / tot * 100)) if tot else None,
                       "productos": [p for p, x in prods.items() if x > 0]})
    return salida


def _r1(v):
    return float(Decimal(repr(float(v))).quantize(Decimal("0.1"), rounding=ROUND_HALF_UP)) + 0.0


# ── Resumen para la síntesis: destinos, aduanas, empresas y el borrador de sus párrafos ──
MESES = ["enero", "febrero", "marzo", "abril", "mayo", "junio", "julio", "agosto", "setiembre", "octubre", "noviembre", "diciembre"]
ORDINAL = {"m": ["el primero", "el segundo", "el tercero", "el cuarto", "el quinto"], "f": ["la primera", "la segunda", "la tercera", "la cuarta", "la quinta"]}


def _pct(v):
    return f"{v:.1f}".replace(".", ",")


def _y(lista):
    lista = [x for x in lista if x]
    return lista[0] if len(lista) == 1 else ", ".join(lista[:-1]) + " y " + lista[-1] if lista else ""


def _minus(p):
    return p[:1].lower() + p[1:] if p and not p[:2].isupper() else p


def _detalle(top, genero, verbo):
    """«El primero adquirió cacao; el segundo, madera aserrada; y el cuarto y el quinto, madera aserrada»: los seguidos
    con los mismos productos se juntan."""
    grupos = []
    for i, t in enumerate(top):
        prods = _y([_minus(p) for p in t["productos"][:2]])
        if grupos and grupos[-1][1] == prods:
            grupos[-1][0].append(i)
        else:
            grupos.append([[i], prods])
    partes = []
    for k, (idx, prods) in enumerate(grupos):
        quien = _y([ORDINAL[genero][i] for i in idx])
        partes.append(f"{quien} {verbo} {prods}" if k == 0 else f"{quien}, {prods}")
    frase = partes[0] if len(partes) == 1 else "; ".join(partes[:-1]) + "; y " + partes[-1]
    return frase[0].upper() + frase[1:]


def _texto_destinos(top, cuando):
    if not top:
        return ""
    partes = _y([f"{t['nombre']} ({_pct(t['pct'])} por ciento)" for t in top])
    if len(top) == 1:
        return f"{cuando}, el principal mercado de destino fue {partes}, que adquirió {_y([_minus(p) for p in top[0]['productos'][:2]])}, principalmente."
    return f"{cuando}, los principales mercados de destino fueron: {partes}. {_detalle(top, 'm', 'adquirió')}, principalmente."


def _texto_aduanas(top, cuando):
    if not top:
        return ""
    if len(top) == 1:
        t = top[0]
        return (f"Asimismo, las exportaciones {cuando} se realizaron a través de la aduana {t['nombre']} ({_pct(t['pct'])} por ciento), "
                f"con envíos de {_y([_minus(p) for p in t['productos'][:2]])}, principalmente.")
    partes = _y([f"{t['nombre']} ({_pct(t['pct'])} por ciento)" for t in top[:3]])
    return f"Asimismo, las exportaciones {cuando} se realizaron a través de las aduanas: {partes}. {_detalle(top[:3], 'f', 'registró mayores movimientos de')}, principalmente."


def _destacados(ok, df, sel, ant, hay_ant, a, m):
    """Datos que llaman la atención, solo los que la base permite: producto y destino principales, cuántos países,
    productos y empresas; con el año anterior en la base, nuevos mercados y lo que más creció; y si el mes es récord."""
    nombre, d, da = MESES[m - 1], sel["mes"], sel["acumulado"]
    salida = []
    tot = d["FOB_USD"].sum()
    if tot <= 0:
        return salida
    prod = d.groupby("PRODUCTO")["FOB_USD"].sum().sort_values(ascending=False)
    salida.append({"titulo": "Producto estrella del mes", "valor": prod.index[0],
                   "detalle": f"{_pct(_r1(prod.iloc[0] / tot * 100))} % de lo exportado en {nombre} (US$ {_pct(_r1(prod.iloc[0] / 1e6))} millones)"})
    pais = d.groupby(d["PAIS"].fillna("(sin dato)"))["FOB_USD"].sum().sort_values(ascending=False)
    salida.append({"titulo": "Principal destino", "valor": pais.index[0], "detalle": f"compró el {_pct(_r1(pais.iloc[0] / tot * 100))} % de las exportaciones del mes"})
    cont = d.groupby(d["CONTINENTE"].fillna("(sin dato)"))["FOB_USD"].sum().sort_values(ascending=False)
    salida.append({"titulo": "Continente principal", "valor": cont.index[0], "detalle": f"{_pct(_r1(cont.iloc[0] / tot * 100))} % del valor del mes"})
    salida.append({"titulo": "Llegamos a", "valor": f"{d['PAIS'].nunique()} países",
                   "detalle": f"en {nombre}; {da['PAIS'].nunique()} países en lo que va del año"})
    salida.append({"titulo": "Productos exportados", "valor": f"{d['PRODUCTO'].nunique()}", "detalle": f"en {nombre}; {da['PRODUCTO'].nunique()} en lo que va del año"})
    salida.append({"titulo": "Empresas exportadoras", "valor": f"{d['RUC'].nunique()}", "detalle": f"en {nombre}; {da['RUC'].nunique()} en lo que va del año"})
    if hay_ant:
        antes = set(ant["acumulado"]["PAIS"].dropna())
        nuevos = da[~da["PAIS"].isin(antes)].groupby("PAIS")["FOB_USD"].sum().sort_values(ascending=False)
        if len(nuevos):
            salida.append({"titulo": "Nuevos mercados", "valor": f"{len(nuevos)}", "detalle": "no compraron en el mismo periodo del año anterior: " + _y(list(nuevos.index[:4]))})
        crec = (da.groupby("PRODUCTO")["FOB_USD"].sum() - ant["acumulado"].groupby("PRODUCTO")["FOB_USD"].sum()).dropna().sort_values(ascending=False)
        if len(crec) and crec.iloc[0] > 0:
            salida.append({"titulo": "El que más creció", "valor": crec.index[0], "detalle": f"US$ {_pct(_r1(crec.iloc[0] / 1e6))} millones más que en enero-{nombre} del año anterior"})
    mensual = ok.groupby(["AÑO", "MES"])["FOB_USD"].sum()
    if len(mensual) >= 13 and mensual.idxmax() == (a, m):
        salida.append({"titulo": "Récord", "valor": f"{nombre.capitalize()} {a}", "detalle": f"el mayor valor mensual desde {MESES[df['MES'].min() - 1]} {df['AÑO'].min()}"})
    return salida


def resumen(depto, periodo=None):
    """Del mes y de enero-mes (filas ok): total, productos, destinos, continentes, aduanas y empresas, con su
    variación contra el mismo periodo del año anterior (si está en la base) y el borrador de los párrafos de la síntesis."""
    df = _base(depto)
    if df.empty:
        return None
    df["AÑO"], df["MES"] = df["AÑO"].astype(int), df["MES"].astype(int)
    ok = df[df["ESTADO"] == "ok"].copy()
    ok["EMPRESA_NOMBRE"] = ok["EMPRESA"].str.replace(r"^\d{11}\s*-\s*", "", regex=True).str.strip()
    meses = sorted({f"{a}-{m:02d}" for a, m in zip(df["AÑO"], df["MES"])})
    periodo = periodo if periodo in meses else meses[-1]
    a, m = int(periodo[:4]), int(periodo[5:7])
    sel = {"mes": ok[(ok["AÑO"] == a) & (ok["MES"] == m)], "acumulado": ok[(ok["AÑO"] == a) & (ok["MES"] <= m)]}
    ant = {"mes": ok[(ok["AÑO"] == a - 1) & (ok["MES"] == m)], "acumulado": ok[(ok["AÑO"] == a - 1) & (ok["MES"] <= m)]}
    hay_ant = bool(((df["AÑO"] == a - 1) & (df["MES"] == m)).any())
    salida = {"periodo": periodo, "meses": meses, "nombre_mes": MESES[m - 1], "anio": a, "hay_anterior": hay_ant}
    for k, d in sel.items():
        tot, tot_b = d["FOB_USD"].sum() / 1000, ant[k]["FOB_USD"].sum() / 1000
        salida[k] = {"total": tot, "anterior": tot_b if hay_ant else None,
                     "var": _r1((tot / tot_b - 1) * 100) if hay_ant and tot_b else None,
                     "productos": _top(d, "PRODUCTO", 8), "destinos": _top(d, "PAIS"), "continentes": _top(d, "CONTINENTE", 6),
                     "aduanas": _top(d, "ADUANA"), "empresas": _top(d, "EMPRESA_NOMBRE")}
    salida["destacados"] = _destacados(ok, df, sel, ant, hay_ant, a, m)
    cuando_m, cuando_a = f"En {MESES[m - 1]} {a}", (f"En el periodo enero-{MESES[m - 1]} {a}" if m > 1 else f"En enero {a}")
    salida["textos"] = [x for x in (_texto_destinos(salida["mes"]["destinos"], cuando_m), _texto_aduanas(salida["mes"]["aduanas"], f"en {MESES[m - 1]} {a}"),
                                    _texto_destinos(salida["acumulado"]["destinos"], cuando_a) if m > 1 else "") if x]
    return salida


def revision(depto, periodo=None):
    df = _base(depto)
    if df.empty:
        return None
    df["AÑO"], df["MES"] = df["AÑO"].astype(int), df["MES"].astype(int)
    meses = sorted({f"{a}-{m:02d}" for a, m in zip(df["AÑO"], df["MES"])})
    periodo = periodo if periodo in meses else meses[-1]
    a, m = int(periodo[:4]), int(periodo[5:7])
    mes = df[(df["AÑO"] == a) & (df["MES"] == m)]
    acum = df[(df["AÑO"] == a) & (df["MES"] <= m)]
    estados = mes.groupby("ESTADO")["FOB"].agg(["count", "sum"]).reset_index().rename(columns={"count": "filas", "sum": "fob"}).to_dict("records")
    emp = leer_tabla("empresas")
    emp = emp.assign(RUC=emp["RUC"].astype(str)).drop_duplicates("RUC", keep="last").set_index("RUC")
    ver = df[df["ESTADO"] == "verificar"]
    empresas = []
    for ruc, g in ver.groupby("RUC"):
        e = emp.loc[ruc] if ruc in emp.index else None
        empresas.append({"ruc": ruc, "razon_social": None if e is None else _txt(e["RAZON_SOCIAL"]), "borrar_para": None if e is None else _txt(e["BORRAR_PARA"]),
                         "nota": None if e is None else _txt(e["NOTA"]), "fob": float(g["FOB"].sum()), "filas": len(g),
                         "meses": sorted({f"{x}-{y:02d}" for x, y in zip(g["AÑO"], g["MES"])}),
                         "productos": g.groupby("PRODUCTO")["FOB"].sum().sort_values(ascending=False).head(3).index.tolist() or g["DESCRIPCION"].head(2).tolist(),
                         "estado_sunat": None if e is None else _txt(e["ESTADO_SUNAT"])})
    cla = df[df["ESTADO"] == "clasificar"]
    partidas = [{"partida": p, "fob": float(g["FOB"].sum()), "filas": len(g), "descripcion": g["DESCRIPCION"].iloc[0]} for p, g in cla.groupby("PARTIDA")]
    ok = df[df["ESTADO"] == "ok"]
    resto = ok[ok["RUBRO_ESTAD"].fillna("").isin(["Resto"])]
    a_resto = [{"producto": p, "categoria": g["CATEGORIA"].iloc[0], "fob": float(g["FOB"].sum())} for p, g in resto.groupby("PRODUCTO") if g["FOB"].sum() > 0]
    est = estructura_estad(depto)
    return {"periodo": periodo, "meses": meses, "estados": estados,
            "empresas": sorted(empresas, key=lambda x: -x["fob"]), "partidas": sorted(partidas, key=lambda x: -x["fob"]),
            "a_resto": sorted(a_resto, key=lambda x: -x["fob"]),
            "rubros": [f["rubro"] for f in est if f["nivel"] in ("producto", "grupo")] + [f["rubro"] for f in est if f["nivel"] == "tipo" and _clave(f["rubro"]).startswith("otras")],
            "categorias": sorted({(str(t), str(c)) for t, c in leer_tabla("partidas")[["TIPO", "CATEGORIA"]].dropna().itertuples(index=False)}),
            "propuesta": propuesta(depto, a, m, mes),
            "reglas": leer_tabla("reglas").to_dict("records"), "equivalencias": leer_tabla("equivalencias").to_dict("records")}


def actualizar_empresas(filas):
    """filas: [{RUC, BORRAR_PARA, NOTA, RAZON_SOCIAL}]: cambia solo lo que trae cada fila (el resto de la empresa se queda)."""
    emp = leer_tabla("empresas")
    emp["RUC"] = emp["RUC"].astype(str)
    for f in filas:
        ruc = re.sub(r"\D", "", str(f.get("RUC") or ""))
        if len(ruc) != 11:
            raise ValueError(f"RUC no válido: {f.get('RUC')!r}.")
        cambios = {k: (str(v).strip() or None) if v is not None else None for k, v in f.items() if k in TABLAS["empresas"] and k != "RUC"}
        if (emp["RUC"] == ruc).any():
            for k, v in cambios.items():
                emp.loc[emp["RUC"] == ruc, k] = v
        else:
            emp = pd.concat([emp, pd.DataFrame([{"RUC": ruc, **cambios}])], ignore_index=True)
        emp.loc[emp["RUC"] == ruc, "REVISADO"] = "True"
    guardar_tabla("empresas", emp)


def actualizar_partidas(filas):
    part = leer_tabla("partidas")
    for f in filas:
        k = re.sub(r"\D", "", str(f.get("PARTIDA") or "")).zfill(10)
        if k == "0" * 10:
            continue
        previa = part[part["PARTIDA"] == k]
        fila = {c: f.get(c, previa[c].iloc[0] if len(previa) else None) for c in TABLAS["partidas"]}
        fila["PARTIDA"] = k
        part = pd.concat([part[part["PARTIDA"] != k], pd.DataFrame([fila])], ignore_index=True)
    guardar_tabla("partidas", part)


def actualizar_equivalencias(filas):
    eq = leer_tabla("equivalencias")
    for f in filas:
        dep, prod = str(f.get("departamento") or "").upper(), str(f.get("producto") or "").strip()
        if not prod:
            continue
        eq = eq[~((eq["departamento"].fillna("").str.upper() == dep) & (eq["producto"].map(_clave) == _clave(prod)))]
        if not f.get("borrar") and f.get("rubro"):
            eq = pd.concat([eq, pd.DataFrame([{"departamento": dep, "producto": prod, "rubro": f["rubro"]}])], ignore_index=True)
    guardar_tabla("equivalencias", eq)


def _txt(v):
    return None if v is None or (isinstance(v, float) and pd.isna(v)) else str(v)


def pendientes_ruc():
    """RUC sin razón social (para el flujo del padrón de SUNAT): los de la base que no están en EMPRESAS o no tienen nombre."""
    base = db.leer_todo("expo-sunat")
    emp = leer_tabla("empresas")
    rucs = set(base["RUC"].astype(str)) if base is not None else set()
    con_nombre = set(emp.loc[emp["RAZON_SOCIAL"].notna() & (emp["RAZON_SOCIAL"].astype(str).str.strip() != ""), "RUC"].astype(str))
    return sorted(rucs - con_nombre)


def guardar_padron(encontrados):
    """encontrados: DataFrame RUC, RAZON_SOCIAL, ESTADO, CONDICION (del padrón reducido). Las empresas nuevas quedan con
    BORRAR PARA vacío (por verificar)."""
    emp = leer_tabla("empresas")
    emp["RUC"] = emp["RUC"].astype(str)
    for r in encontrados.itertuples(index=False):
        if (emp["RUC"] == r.RUC).any():
            i = emp.index[emp["RUC"] == r.RUC]
            emp.loc[i, "RAZON_SOCIAL"] = r.RAZON_SOCIAL
            emp.loc[i, "ESTADO_SUNAT"], emp.loc[i, "CONDICION"] = r.ESTADO, r.CONDICION
        else:
            emp = pd.concat([emp, pd.DataFrame([{"RUC": r.RUC, "RAZON_SOCIAL": r.RAZON_SOCIAL, "ESTADO_SUNAT": r.ESTADO, "CONDICION": r.CONDICION}])], ignore_index=True)
    guardar_tabla("empresas", emp)
    return len(encontrados)


def excel_propuesta(depto, anio, mes):
    from openpyxl import Workbook
    from openpyxl.styles import Font
    p = propuesta(depto, anio, mes)
    wb = Workbook()
    ws = wb.active
    ws.title = "ESTAD propuesto"
    ws.append([f"{NOMBRE[depto]}: EXPORTACIONES FOB POR GRUPO DE PRODUCTOS · {anio}-{mes:02d}"])
    ws.append(["(Miles de US dólares) · desde la base mensual (filas ok)"])
    ws.append(["Rubro", "Base mensual (propuesto)", "ESTAD cargado", "Diferencia"])
    for f in (p or {}).get("filas", []):
        ws.append([("   " * {"total": 0, "tipo": 0, "grupo": 1, "producto": 2}[f["nivel"]]) + f["rubro"], f["sunat"], f["estad"], f["dif"]])
        if f["nivel"] in ("total", "tipo"):
            ws.cell(ws.max_row, 1).font = Font(bold=True)
    ws.column_dimensions["A"].width = 36
    det = _base(depto, f"{anio}-{mes:02d}", f"{anio}-{mes:02d}")
    w2 = wb.create_sheet("Detalle")
    w2.append(COLUMNAS)
    for r in det[COLUMNAS].itertuples(index=False):
        w2.append([None if (isinstance(v, float) and pd.isna(v)) else v for v in r])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
