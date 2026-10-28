"""Ejecución del presupuesto (sección V de la síntesis): gasto corriente no financiero e inversión pública de un
departamento, con lo que pide el informe y más para mirar:

  · el mes y el acumulado Ene.-Mes de un año contra otro (por defecto, el anterior), por nivel de gobierno y por
    genérica (gasto) o función (inversión), en millones de soles, con variación % real (deflactada con el IPC de Lima)
    y nominal; en inversión, además, el PIM y su avance;
  · series mensuales de los dos años (y su acumulado), la variación real interanual y la participación de cada nivel;
  · variaciones alarmantes: los 5 mayores aumentos y caídas de la inversión en soles (mismo criterio que el Bloque 6 de
    Fiscal_Procesamiento.ipynb), por todos los niveles juntos, por nivel y función y por municipalidad.

Gasto corriente no financiero = genéricas de tipo de transacción 5 (personal, pensiones, bienes y servicios,
donaciones y transferencias corrientes y otros gastos corrientes): no incluye la deuda (tipo 7) ni el capital (6).
Inversión = genérica 2.6 (adquisición de activos no financieros), como en la base «Inversión por función».
"""
import io
import re
import unicodedata

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font

NIVELES = {"E": "Gobierno Nacional", "R": "Gobiernos Regionales", "M": "Gobiernos Locales"}
CORTO = {"E": "G.N.", "R": "G.R.", "M": "G.L."}
MESES = ["Enero", "Febrero", "Marzo", "Abril", "Mayo", "Junio", "Julio", "Agosto", "Setiembre", "Octubre", "Noviembre", "Diciembre"]
MESES_CORTOS = ["Ene.", "Feb.", "Mar.", "Abr.", "May.", "Jun.", "Jul.", "Ago.", "Set.", "Oct.", "Nov.", "Dic."]
# genéricas del gasto corriente (tipo 5), como las filas de Fiscal_Procesamiento
GENERICAS = [("Personal y obligaciones sociales", r"personal"), ("Pensiones y prestaciones sociales", r"pension"),
             ("Bienes y servicios", r"bienes y servicios"), ("Donaciones y transferencias", r"donaciones"),
             ("Otros gastos corrientes", r"otros gastos")]
TOP = 5
FUENTE = "Fuente: MEF, Consulta Amigable. Elaboración: BCRP, Sucursal Iquitos. Departamento de Estudios Económicos."


def _norm(t):
    t = unicodedata.normalize("NFKD", str(t or ""))
    return "".join(c for c in t if not unicodedata.combining(c)).lower().strip()


def nivel_de(texto):
    """«E: GOBIERNO NACIONAL», «GOBIERNOS REGIONALES», «M: GOBIERNOS LOCALES»… -> E / R / M."""
    t = _norm(texto)
    if t[:2] in ("e:", "r:", "m:"):
        return t[0].upper()
    return "E" if "nacional" in t else "R" if "regional" in t else "M" if "local" in t or "munic" in t else None


def _mes(v):
    m = re.match(r"^\s*(\d+)", str(v))
    return int(m.group(1)) if m else 0


def _funcion(f):
    """«16: COMUNICACIONES» se suma a «15: TRANSPORTE», como en los cuadros del informe."""
    f = str(f).strip()
    return "15: TRANSPORTE" if f.startswith("16:") else f


def _generica(cod, etiqueta):
    t = re.sub(r"^[\w\-\.]+\s*:\s*", "", _norm(etiqueta))
    return next((n for n, r in GENERICAS if re.search(r, t)), "Otros gastos corrientes")


def etiqueta(p, corta=False):
    return f"{(MESES_CORTOS if corta else MESES)[p % 100 - 1]} {p // 100}"


# ── Paneles: P (AAAAMM), NIVEL (E/R/M), ITEM, MONTO (soles) ──
def panel_gasto(gen):
    if gen is None or gen.empty:
        return pd.DataFrame(columns=["P", "NIVEL", "ITEM", "MONTO"])
    g = gen[gen["COD_GENERICA"].astype(str).str.strip().str.startswith("5")]
    return pd.DataFrame({"P": g["AÑO"].astype(int) * 100 + g["MES"].map(_mes), "NIVEL": g["NIVEL_GOBIERNO"].map(nivel_de),
                         "ITEM": [_generica(c, e) for c, e in zip(g["COD_GENERICA"], g["GENERICA"])],
                         "MONTO": pd.to_numeric(g["DEVENGADO"], errors="coerce").fillna(0)}).dropna(subset=["NIVEL"])


def panel_inversion(fun):
    if fun is None or fun.empty:
        return pd.DataFrame(columns=["P", "NIVEL", "ITEM", "MONTO"])
    return pd.DataFrame({"P": fun["AÑO"].astype(int) * 100 + fun["MES"].map(_mes), "NIVEL": fun["NIVEL_GOBIERNO"].map(nivel_de),
                         "ITEM": fun["FUNCION"].map(_funcion),
                         "MONTO": pd.to_numeric(fun["DEVENGADO"], errors="coerce").fillna(0)}).dropna(subset=["NIVEL"])


def panel_pim(pim):
    if pim is None or pim.empty:
        return pd.DataFrame(columns=["ANIO", "NIVEL", "ITEM", "PIM"])
    return pd.DataFrame({"ANIO": pim["AÑO"].astype(int), "NIVEL": pim["NIVEL DE GOBIERNO"].map(nivel_de),
                         "ITEM": pim["FUNCION"].map(_funcion), "PIM": pd.to_numeric(pim["PIM"], errors="coerce").fillna(0)}).dropna(subset=["NIVEL"])


def panel_municipalidad(muni):
    if muni is None or muni.empty:
        return pd.DataFrame(columns=["P", "ITEM", "MONTO"])
    return pd.DataFrame({"P": muni["AÑO"].astype(int) * 100 + muni["MES"].map(_mes),
                         "ITEM": muni["MUNICIPALIDAD"].astype(str).str.replace(r"^[\d\-]+:\s*", "", regex=True),
                         "MONTO": pd.to_numeric(muni["DEVENGADO"], errors="coerce").fillna(0)})


# ── Cálculos ──
def _var(a, b):
    return None if a is None or b is None or not a else round((b / a - 1) * 100, 1)


class Deflactor:
    """Suma en soles de un mes de referencia: monto_t × IPC_ref / IPC_t. Sin IPC para algún mes, devuelve None."""

    def __init__(self, ipc, ref):
        self.ipc, self.ref = ipc, ipc.get(f"{ref // 100}-{ref % 100:02d}")

    def real(self, df, periodos):
        if not self.ref:
            return None
        total = 0.0
        for q, monto in df[df["P"].isin(periodos)].groupby("P")["MONTO"].sum().items():
            v = self.ipc.get(f"{q // 100}-{q % 100:02d}")
            if not v:
                return None
            total += monto * self.ref / v
        return total


def _fila(nombre, df, ventanas, defl, pim=None, pim_b=None):
    nom = {w: float(df.loc[df["P"].isin(ps), "MONTO"].sum()) for w, ps in ventanas.items()}
    real = {w: defl.real(df, ps) for w, ps in ventanas.items()}
    f = {"nombre": nombre, **{w: round(v / 1e6, 3) for w, v in nom.items()},
         "var_mes": _var(real["mes_b"], real["mes_a"]) if real["mes_a"] is not None and real["mes_b"] is not None else None,
         "var_acum": _var(real["acum_b"], real["acum_a"]) if real["acum_a"] is not None and real["acum_b"] is not None else None,
         "var_mes_nominal": _var(nom["mes_b"], nom["mes_a"]), "var_acum_nominal": _var(nom["acum_b"], nom["acum_a"])}
    if pim is not None:
        f["pim"] = round(pim / 1e6, 3)
        f["avance"] = round(nom["acum_a"] / pim * 100, 1) if pim else None
        f["avance_b"] = round(nom["acum_b"] / pim_b * 100, 1) if pim_b else None
    return f


def cuadros(sector, panel, pim, ipc, periodo=None, comparar=None):
    """sector: 'gasto' o 'inversion'. panel: P/NIVEL/ITEM/MONTO de un departamento; pim: ANIO/NIVEL/ITEM/PIM (inversión)."""
    if panel.empty:
        return None
    periodos = sorted(set(panel["P"]))
    ultimo = periodos[-1]
    p = int(periodo.replace("-", "")) if periodo else ultimo
    if p not in periodos:
        raise ValueError(f"No hay datos de {etiqueta(p)} (último mes: {etiqueta(ultimo)}).")
    a, m = p // 100, p % 100
    anios = sorted({q // 100 for q in periodos})
    b = int(comparar) if comparar else a - 1
    if b >= a:
        raise ValueError(f"El año de comparación ({b}) debe ser anterior a {a}.")
    pb = b * 100 + m
    ventanas = {"mes_b": [pb], "mes_a": [p], "acum_b": list(range(b * 100 + 1, pb + 1)), "acum_a": list(range(a * 100 + 1, p + 1))}
    defl = Deflactor(ipc, p)
    con_pim = sector == "inversion" and pim is not None and not pim.empty
    pim_de = lambda anio, nivel=None, item=None: float(pim[(pim["ANIO"] == anio) & ((pim["NIVEL"] == nivel) if nivel else True)
                                                          & ((pim["ITEM"] == item) if item else True)]["PIM"].sum()) if con_pim else None

    niveles = [_fila(NIVELES[n], panel[panel["NIVEL"] == n], ventanas, defl, pim_de(a, n), pim_de(b, n)) for n in NIVELES]
    niveles.append(_fila("Total", panel, ventanas, defl, pim_de(a), pim_de(b)))
    items = sorted(set(panel["ITEM"]) | (set(pim[pim["ANIO"] == a]["ITEM"]) if con_pim else set()))
    if sector == "gasto":
        items = [n for n, _ in GENERICAS if n in items]
    detalle = [_fila(i, panel[panel["ITEM"] == i], ventanas, defl, pim_de(a, item=i), pim_de(b, item=i)) for i in items]
    detalle = [f for f in detalle if f["acum_a"] or f["acum_b"] or f.get("pim")]
    if sector == "inversion":
        detalle.sort(key=lambda f: -f["acum_a"])
    detalle.append(_fila("Total", panel, ventanas, defl, pim_de(a), pim_de(b)))

    # series de los dos años (millones de soles corrientes), mes a mes y acumuladas, en total y por nivel
    mensual = panel.groupby("P")["MONTO"].sum() / 1e6
    serie = lambda anio: [round(float(mensual.get(anio * 100 + k, 0)), 3) if anio * 100 + k in mensual.index else None for k in range(1, 13)]
    acumulada = lambda s: [None if v is None else round(sum(x or 0 for x in s[:i + 1]), 3) for i, v in enumerate(s)]
    sa, sb = serie(a), serie(b)
    por_nivel = {CORTO[n]: {"a": [round(float(panel[(panel["NIVEL"] == n) & (panel["P"] == a * 100 + k)]["MONTO"].sum()) / 1e6, 3) if a * 100 + k in mensual.index else None for k in range(1, 13)]}
                 for n in NIVELES}
    # variación real interanual del total, mes a mes (hasta 25 meses)
    interanual = []
    for q in [x for x in periodos if x <= p][-25:]:
        d = Deflactor(ipc, q)
        ra, rb = d.real(panel, [q]), d.real(panel, [q - 100]) if (q - 100) in periodos else None
        na, nb = float(mensual.get(q, 0)), float(mensual.get(q - 100, 0)) if (q - 100) in periodos else None
        interanual.append({"periodo": etiqueta(q, True), "var": _var(rb, ra) if ra is not None and rb is not None else None,
                           "var_nominal": _var(nb, na) if nb is not None else None})
    total = niveles[-1]
    participacion = [{"nivel": CORTO[n], "valor": f["acum_a"], "pct": round(f["acum_a"] / total["acum_a"] * 100, 1) if total["acum_a"] else None}
                     for n, f in zip(NIVELES, niveles[:3])]
    salida = {"sector": sector, "periodo": f"{a}-{m:02d}", "ultimo": f"{ultimo // 100}-{ultimo % 100:02d}", "anio": a, "comparar": b,
              "titulo_mes": MESES[m - 1], "anios": anios, "real": defl.real(panel, ventanas["acum_a"]) is not None and defl.real(panel, ventanas["acum_b"]) is not None,
              "niveles": niveles, "detalle": detalle, "participacion": participacion, "interanual": interanual,
              "mensual": {"meses": MESES_CORTOS, "a": sa, "b": sb, "acum_a": acumulada(sa), "acum_b": acumulada(sb), "por_nivel": por_nivel}}
    if con_pim:
        pa, pbm = pim_de(a), pim_de(b)
        # avance del PIM mes a mes (acumulado / PIM del año) de los dos años, y el ritmo «parejo» (k/12)
        salida["avance"] = {"a": [None if v is None else round(v * 1e6 / pa * 100, 1) if pa else None for v in salida["mensual"]["acum_a"]],
                            "b": [None if v is None else round(v * 1e6 / pbm * 100, 1) if pbm else None for v in salida["mensual"]["acum_b"]],
                            "parejo": [round(k / 12 * 100, 1) for k in range(1, 13)]}
    return salida


# ── Variaciones alarmantes (inversión, soles corrientes) ──
def _tops(t):
    """t: DataFrame ITEM, ANTES, AHORA. Siempre TOP aumentos (mayores diferencias) y TOP caídas (menores), sin ítems en 0
    en ambos periodos y sin repetir un ítem en las dos listas (Fiscal_Procesamiento, Bloque 6)."""
    t = t[(t["ANTES"] != 0) | (t["AHORA"] != 0)].assign(DIF=lambda x: x["AHORA"] - x["ANTES"])
    sube = t.sort_values("DIF", ascending=False).head(TOP)
    baja = t[~t["ITEM"].isin(sube["ITEM"])].sort_values("DIF").head(TOP)
    fila = lambda r: {"item": r.ITEM, "antes": round(r.ANTES / 1e6, 3), "ahora": round(r.AHORA / 1e6, 3),
                      "dif": round(r.DIF / 1e6, 3), "var": _var(r.ANTES, r.AHORA)}
    return [fila(r) for r in sube.itertuples()], [fila(r) for r in baja.itertuples()], len(t)


def variaciones(panel, muni, periodo=None, comparar=None):
    """Bloques de variaciones de la inversión de un departamento: el mes y el acumulado contra el año de comparación."""
    if panel.empty:
        return None
    periodos = sorted(set(panel["P"]))
    p = int(periodo.replace("-", "")) if periodo else periodos[-1]
    a, m = p // 100, p % 100
    b = int(comparar) if comparar else a - 1
    pb = b * 100 + m
    comparaciones = [("mes", f"{etiqueta(p)} vs {etiqueta(pb)}", [pb], [p]),
                     ("acumulado", f"Ene.-{MESES_CORTOS[m - 1]} {a} vs Ene.-{MESES_CORTOS[m - 1]} {b}",
                      list(range(b * 100 + 1, pb + 1)), list(range(a * 100 + 1, p + 1)))]
    salida = {"periodo": f"{a}-{m:02d}", "comparar": b, "titulo_mes": MESES[m - 1], "anio": a, "comparaciones": []}
    for clave, titulo, antes, ahora in comparaciones:
        bloques = []

        def bloque(nombre, df, item):
            falta = [etiqueta(q) for q in sorted(set(antes) | set(ahora)) if q not in set(df["P"])] if not df.empty else ["todos"]
            if df.empty or any(q not in set(df["P"]) for q in (antes[-1], ahora[-1])):
                bloques.append({"nombre": nombre, "sin_datos": True, "falta": falta[:6]})
                return
            t = pd.DataFrame({"ITEM": sorted(set(df[item]))})
            t["ANTES"] = t["ITEM"].map(df[df["P"].isin(antes)].groupby(item)["MONTO"].sum()).fillna(0)
            t["AHORA"] = t["ITEM"].map(df[df["P"].isin(ahora)].groupby(item)["MONTO"].sum()).fillna(0)
            sube, baja, n = _tops(t)
            ta, tb = float(t["ANTES"].sum()), float(t["AHORA"].sum())
            bloques.append({"nombre": nombre, "total": {"antes": round(ta / 1e6, 3), "ahora": round(tb / 1e6, 3), "dif": round((tb - ta) / 1e6, 3), "var": _var(ta, tb)},
                            "aumentos": sube, "caidas": baja, "pocos": n < 2 * TOP})

        todos = panel.assign(ITEM2=panel["NIVEL"].map(CORTO) + " · " + panel["ITEM"])
        bloque("Todos los niveles (nivel · función)", todos, "ITEM2")
        for n in NIVELES:
            bloque(f"{NIVELES[n]} · por función", panel[panel["NIVEL"] == n], "ITEM")
        if muni is not None and not muni.empty:
            bloque("Gobiernos Locales · por municipalidad", muni, "ITEM")
        salida["comparaciones"].append({"clave": clave, "titulo": titulo, "bloques": bloques})
    return salida


# ── Excel ──
def excel(region, c, v=None):
    """Cuadros, series y variaciones en un Excel."""
    wb = Workbook()
    neg, centro = Font(bold=True), Alignment(horizontal="center")
    a, b, mes = c["anio"], c["comparar"], c["titulo_mes"]
    titulo = "GASTO CORRIENTE NO FINANCIERO" if c["sector"] == "gasto" else "INVERSIÓN PÚBLICA"
    inv = c["sector"] == "inversion"
    ws = wb.active
    ws.title = "Cuadro"
    ws.append([f"{region.upper()}: {titulo}"]); ws.append(["(Millones de soles y variación % real)" if c["real"] else "(Millones de soles y variación % nominal: falta el IPC)"])
    cab = ["", f"{mes} {b}", f"{mes} {a}", "Var. %", f"Ene.-{mes[:3]}. {b}", f"Ene.-{mes[:3]}. {a}", "Var. %"] + (["PIM", "Avance %"] if inv else [])
    for nombre, filas in (("Por nivel de gobierno", c["niveles"]), ("Por genérica" if not inv else "Por función", c["detalle"])):
        ws.append([]); ws.append([nombre] + cab[1:])
        for x in ws[ws.max_row]:
            x.font, x.alignment = neg, centro
        for f in filas:
            vm = f["var_mes"] if c["real"] else f["var_mes_nominal"]
            va = f["var_acum"] if c["real"] else f["var_acum_nominal"]
            ws.append([f["nombre"], f["mes_b"], f["mes_a"], vm, f["acum_b"], f["acum_a"], va] + ([f.get("pim"), f.get("avance")] if inv else []))
            for i, x in enumerate(ws[ws.max_row]):
                x.number_format = "0.0" if i in (3, 6, 8) else "#,##0.0"
                if f["nombre"] == "Total":
                    x.font = neg
    ws.append([]); ws.append([FUENTE])
    ws["A1"].font = neg
    ws.column_dimensions["A"].width = 52
    for col in "BCDEFGHI":
        ws.column_dimensions[col].width = 13
    wm = wb.create_sheet("Mensual")
    wm.append([f"{region.upper()}: {titulo} POR MES"]); wm.append(["(Millones de soles corrientes)"]); wm.append([])
    wm.append(["Mes", str(b), str(a), f"Acumulado {b}", f"Acumulado {a}"] + (["Avance PIM % " + str(b), "Avance PIM % " + str(a)] if inv and c.get("avance") else []))
    mm = c["mensual"]
    for i, k in enumerate(mm["meses"]):
        wm.append([k, mm["b"][i], mm["a"][i], mm["acum_b"][i], mm["acum_a"][i]] + ([c["avance"]["b"][i], c["avance"]["a"][i]] if inv and c.get("avance") else []))
    wm["A1"].font = neg
    if v:
        for comp in v["comparaciones"]:
            wv = wb.create_sheet(f"Variaciones {comp['clave']}"[:31])
            wv.append([f"{region.upper()}: VARIACIONES DE LA INVERSIÓN · {comp['titulo']}"]); wv.append(["(Millones de soles corrientes; Var. % referencial)"])
            for bl in comp["bloques"]:
                wv.append([]); wv.append([bl["nombre"]]); wv.cell(wv.max_row, 1).font = neg
                if bl.get("sin_datos"):
                    wv.append(["Sin datos para comparar"]); continue
                wv.append(["#", "Ítem", "Antes", "Ahora", "Diferencia", "Var. %"])
                t = bl["total"]; wv.append(["", "Total", t["antes"], t["ahora"], t["dif"], t["var"]])
                for signo, lista in (("▲ Mayores aumentos", bl["aumentos"]), ("▼ Mayores caídas", bl["caidas"])):
                    wv.append([signo])
                    for i, f in enumerate(lista, 1):
                        wv.append([i, f["item"], f["antes"], f["ahora"], f["dif"], f["var"]])
            wv.column_dimensions["B"].width = 60
            wv["A1"].font = neg
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
