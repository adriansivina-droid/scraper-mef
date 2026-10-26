"""Arma los archivos de descarga. El Excel copia el formato de PIM_Funcion.xlsx: una hoja plana por base, encabezados
centrados en la fila 1, montos con separador de miles y los mismos anchos de columna."""
import io
import zipfile

from openpyxl import Workbook
from openpyxl.styles import Alignment

from .catalogo import BASES


def excel(tablas):
    """tablas: {base_id: DataFrame} en el orden del catálogo -> bytes del .xlsx."""
    wb = Workbook(write_only=False)
    wb.remove(wb.active)
    centro = Alignment(horizontal="center")
    for base_id, df in tablas.items():
        b = BASES[base_id]
        ws = wb.create_sheet(b.hoja)
        ws.append(b.columnas)
        for celda in ws[1]:
            celda.alignment = centro
        for fila in df.itertuples(index=False):
            ws.append(list(fila))
        col_valor = b.columnas.index(b.valor) + 1
        for (celda,) in ws.iter_rows(min_row=1, min_col=col_valor, max_col=col_valor):
            celda.number_format = b.formato_valor
        for i, ancho in enumerate(b.anchos):
            ws.column_dimensions[chr(65 + i)].width = ancho
    salida = io.BytesIO()
    wb.save(salida)
    return salida.getvalue()


def csv(df):
    return df.to_csv(index=False).encode("utf-8-sig")


def zip_csv(tablas, nombre_de):
    salida = io.BytesIO()
    with zipfile.ZipFile(salida, "w", zipfile.ZIP_DEFLATED) as z:
        for base_id, df in tablas.items():
            z.writestr(nombre_de(base_id), csv(df))
    return salida.getvalue()
