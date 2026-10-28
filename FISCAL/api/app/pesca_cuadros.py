"""Cuadros de pesca para la síntesis de cada región (sobre los LONG de las bases pesca-*):

  · producción pesquera: VBP (miles de soles de 2007) y volúmenes por estado, especie u origen, del mes y del
    acumulado Ene.-Mes de un año contra otro, con su variación %;
  · series del VBP de los dos años mes a mes y su variación interanual;
  · en Loreto, además, el Cuadro Nº 07 «Indicadores del sector pesquero» del mes contra el mismo mes del otro año
    (mismas filas que cuadro_07 de PESCA/Pesca_Loreto_Long.ipynb).
"""
import io

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font

from .pesca import ESPECIES_07, _clave

MESES = ["Enero", "Febrero", "Marzo", "Abril", "Mayo", "Junio", "Julio", "Agosto", "Setiembre", "Octubre", "Noviembre", "Diciembre"]
MESES_CORTOS = ["Ene.", "Feb.", "Mar.", "Abr.", "May.", "Jun.", "Jul.", "Ago.", "Set.", "Oct.", "Nov.", "Dic."]
# Cuadro Nº 07, bloque «Desembarque fresco» (como ESPECIES_FRESCO del notebook de Loreto)
ESPECIES_FRESCO = ["Llambina", "Ractacara", "Boquichico", "Palometa", "Sardina", "Carachama", "Paiche", "Fasaco", "Acarahuazú",
                   "Sabálo", "Gamitana", "Chambira", "Paña", "Bagre", "Zungaro Doncella", "Lisa", "Mota", "Maparate", "Arahuana"]
FUENTE = {"pesca-loreto": "DIREPRO Loreto", "pesca-san-martin": "DIREPRO San Martín", "pesca-ucayali": "DIREPRO Ucayali"}


def etiqueta(p, corta=False):
    return f"{(MESES_CORTOS if corta else MESES)[p % 100 - 1]} {p // 100}"


def _var(a, b):
    return None if not a else round((b / a - 1) * 100, 1)


def _filas(base_id, df):
    """[(nombre, unidad, máscara, columna, nivel)] del cuadro de producción de cada región (nivel 0 = total)."""
    if base_id == "pesca-loreto":
        pez = df["ESTADO"] != "ORNAMENTAL"
        f = [("Valor bruto pesquero", "miles de S/ de 2007", df["VBP"].notna(), "VBP", 0),
             ("Desembarque total", "t", pez, "VOLUMEN", 0),
             ("Estado fresco", "t", df["ESTADO"] == "FRESCO", "VOLUMEN", 1),
             ("Estado seco salado", "t", df["ESTADO"] == "SECO SALADO", "VOLUMEN", 1),
             ("Estado salpreso", "t", df["ESTADO"] == "SALPRESO", "VOLUMEN", 1),
             ("Boquichico", "t", pez & (df["ESPECIE_CUADRO"] == "Boquichico"), "VOLUMEN", 1),
             ("Llambina/Palometa", "t", pez & (df["ESPECIE_CUADRO"] == "Llambina/Palometa"), "VOLUMEN", 1),
             ("Peces ornamentales vivos", "miles de unidades", ~pez, "VOLUMEN", 0)]
        return f
    if base_id == "pesca-san-martin":
        sm = df["DEPARTAMENTO"].astype(str).str.upper() == "SAN MARTÍN"
        f = [("Valor bruto pesquero", "miles de S/ de 2007", sm, "VBP", 0),
             ("Producción de San Martín", "t", sm, "VOLUMEN", 0)]
        f += [(o, "t", sm & (df["ORIGEN"] == o), "VOLUMEN", 1) for o in ("Acuicultura", "Extracción de ríos")]
        f += [(f"Estado {e.lower()}", "t", sm & (df["ESTADO_OK"] == e), "VOLUMEN", 1) for e in ("FRESCO", "SECO SALADO", "SALPRESO")]
        principales = df[sm & (df["ESPECIE_OK"] != "Otras")].groupby("ESPECIE_OK")["VOLUMEN"].sum().sort_values(ascending=False).index[:6]
        f += [(e, "t", sm & (df["ESPECIE_OK"] == e), "VOLUMEN", 1) for e in principales]
        f.append(("Llega de otros departamentos", "t", ~sm, "VOLUMEN", 0))
        return f
    fresco = df["TIPO"] == "FRESCO"
    f = [("Valor bruto pesquero", "miles de S/ de 2007", df["TIPO"].isin(["FRESCO", "SECO-SALADO"]), "VBP", 0),
         ("Desembarque de pescado", "t", df["TIPO"] == "AGREGADO", "VOLUMEN", 0),
         ("Fresco", "t", fresco, "VOLUMEN", 1), ("Seco-salado", "t", df["TIPO"] == "SECO-SALADO", "VOLUMEN", 1)]
    principales = df[fresco & (df["ESPECIES_OK"] != "Otros")].groupby("ESPECIES_OK")["VOLUMEN"].sum().sort_values(ascending=False).index[:6]
    f += [(f"{e} (fresco)", "t", fresco & (df["ESPECIES_OK"] == e), "VOLUMEN", 1) for e in principales]
    f.append(("Peces ornamentales", "unidades", df["TIPO"] == "ORNAMENTAL", "VOLUMEN", 0))
    return f


def _cuadro_07(df, pb, p):
    """Filas del Cuadro Nº 07 de Loreto: (nombre, antes, ahora, var, nivel)."""
    pez, orn = df[df["ESTADO"] != "ORNAMENTAL"], df[df["ESTADO"] == "ORNAMENTAL"]
    v = lambda t, q: float(t.loc[t["P"] == q, "VOLUMEN"].sum())
    fila = lambda nombre, t, nivel=1: (nombre, round(v(t, pb), 3), round(v(t, p), 3), _var(v(t, pb), v(t, p)), nivel)
    filas = [fila("Desembarque total 1/", pez, 0)]
    filas += [fila(n, pez[pez["ESTADO"] == e]) for n, e in (("Fresco", "FRESCO"), ("Salpreso", "SALPRESO"), ("Seco-salado", "SECO SALADO"))]
    filas.append(fila("Desembarque total 2/", pez, 0))
    filas += [fila(e, pez[pez["ESPECIE_OK"] == e]) for e in ESPECIES_07]
    filas.append(fila("Otros", pez[~pez["ESPECIE_OK"].isin(ESPECIES_07)]))
    filas.append(fila("Peces ornamentales 3/ (miles de unidades)", orn, 0))
    filas += [fila(n, orn[orn["DESTINO"] == d]) for n, d in (("Mercado interno", "Mercado interno"), ("Mercado externo", "Mercado externo"),
                                                               ("Stock reman. y/o mort.", "Stock remanente y/o mortandad"))]
    fre = pez[pez["ESTADO"] == "FRESCO"]
    lista = {_clave(e): e for e in ESPECIES_FRESCO}
    grupo = fre["ESPECIE_CORREGIDA"].map(lambda e: lista.get(_clave(e), "Otras especies"))
    filas.append(fila("Desembarque fresco", fre, 0))
    filas += [fila(e, fre[grupo == e]) for e in ESPECIES_FRESCO + ["Otras especies"]]
    return [{"nombre": n, "antes": a, "ahora": b, "var": x, "nivel": nv} for n, a, b, x, nv in filas]


def cuadros(base_id, df, periodo=None, comparar=None):
    if df.empty:
        return None
    df = df.assign(P=df["AÑO"].astype(int) * 100 + df["MES"].astype(int))
    df["VOLUMEN"] = pd.to_numeric(df["VOLUMEN"], errors="coerce").fillna(0)
    df["VBP"] = pd.to_numeric(df["VBP"], errors="coerce").fillna(0)
    periodos = sorted(set(df["P"]))
    ultimo = periodos[-1]
    p = int(periodo.replace("-", "")) if periodo else ultimo
    if p not in periodos:
        raise ValueError(f"No hay datos de pesca de {etiqueta(p)} (último mes: {etiqueta(ultimo)}).")
    a, m = p // 100, p % 100
    b = int(comparar) if comparar else a - 1
    if b >= a:
        raise ValueError(f"El año de comparación ({b}) debe ser anterior a {a}.")
    pb = b * 100 + m
    ventanas = {"mes_b": [pb], "mes_a": [p], "acum_b": list(range(b * 100 + 1, pb + 1)), "acum_a": list(range(a * 100 + 1, p + 1))}
    filas = []
    for nombre, unidad, mascara, col, nivel in _filas(base_id, df):
        t = df[mascara]
        v = {w: float(t.loc[t["P"].isin(ps), col].sum()) for w, ps in ventanas.items()}
        filas.append({"nombre": nombre, "unidad": unidad, "nivel": nivel, **{w: round(x, 3) for w, x in v.items()},
                      "var_mes": _var(v["mes_b"], v["mes_a"]), "var_acum": _var(v["acum_b"], v["acum_a"])})
    vbp = df[_filas(base_id, df)[0][2]].groupby("P")["VBP"].sum()
    serie = lambda anio: [round(float(vbp[anio * 100 + k]), 3) if anio * 100 + k in vbp.index else None for k in range(1, 13)]
    interanual = [{"periodo": etiqueta(q, True), "var": _var(float(vbp[q - 100]), float(vbp[q])) if q - 100 in vbp.index else None}
                  for q in [x for x in periodos if x <= p][-25:]]
    salida = {"periodo": f"{a}-{m:02d}", "ultimo": f"{ultimo // 100}-{ultimo % 100:02d}", "anio": a, "comparar": b,
              "titulo_mes": MESES[m - 1], "anios": sorted({q // 100 for q in periodos}), "filas": filas,
              "mensual": {"meses": MESES_CORTOS, "a": serie(a), "b": serie(b)}, "interanual": interanual, "fuente": FUENTE[base_id]}
    if base_id == "pesca-loreto":
        salida["cuadro_07"] = _cuadro_07(df, pb, p)
    return salida


def excel(region, c):
    wb = Workbook()
    neg, centro = Font(bold=True), Alignment(horizontal="center")
    a, b, mes = c["anio"], c["comparar"], c["titulo_mes"]
    ws = wb.active
    ws.title = "Producción"
    ws.append([f"{region.upper()}: PRODUCCIÓN PESQUERA"]); ws.append([])
    ws.append(["", "Unidad", f"{mes} {b}", f"{mes} {a}", "Var. %", f"Ene.-{mes[:3]}. {b}", f"Ene.-{mes[:3]}. {a}", "Var. %"])
    for x in ws[3]:
        x.font, x.alignment = neg, centro
    for f in c["filas"]:
        ws.append([("   " * f["nivel"]) + f["nombre"], f["unidad"], f["mes_b"], f["mes_a"], f["var_mes"], f["acum_b"], f["acum_a"], f["var_acum"]])
        for i, x in enumerate(ws[ws.max_row]):
            x.number_format = "0.0" if i in (4, 7) else "#,##0.0"
            if f["nivel"] == 0:
                x.font = neg
    ws.append([]); ws.append([f"Fuente: {c['fuente']}. Elaboración: BCRP, Sucursal Iquitos. Departamento de Estudios Económicos."])
    ws["A1"].font = neg
    ws.column_dimensions["A"].width = 40; ws.column_dimensions["B"].width = 20
    if c.get("cuadro_07"):
        w7 = wb.create_sheet("Cuadro 07")
        w7.append([f"{region.upper()}: INDICADORES DEL SECTOR PESQUERO (t)"]); w7.append([])
        w7.append(["", f"{mes} {b}", f"{mes} {a}", "Var. %"])
        for f in c["cuadro_07"]:
            w7.append([("   " * f["nivel"]) + f["nombre"], f["antes"], f["ahora"], f["var"]])
            if f["nivel"] == 0:
                for x in w7[w7.max_row]:
                    x.font = neg
        w7.column_dimensions["A"].width = 44
    wm = wb.create_sheet("VBP mensual")
    wm.append([f"{region.upper()}: VALOR BRUTO PESQUERO (miles de S/ de 2007)"]); wm.append([]); wm.append(["Mes", str(b), str(a)])
    for i, k in enumerate(c["mensual"]["meses"]):
        wm.append([k, c["mensual"]["b"][i], c["mensual"]["a"][i]])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
