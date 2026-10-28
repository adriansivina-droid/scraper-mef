"""Cuadros, anexo y gráficos de exportaciones (e importaciones de Loreto) desde el ESTAD (EXPORT/info_export.md §5–§7).

  · síntesis (Cuadro N° 7 Loreto, N° 5 San Martín, N° 8 Ucayali): por tipo y grupo, estructura % del mes del año anterior,
    mes y enero-mes de los dos años, var. % y contribución. Millones de US$ redondeados (Loreto 2 decimales, San Martín y
    Ucayali 1); el total es la suma de los tipos ya redondeados y las variaciones salen de las cifras redondeadas (la
    regla del Cuadro_LOR). Para redactar, «var_exacta» usa los miles de US$ del ESTAD.
  · anexo: los 13 meses (mismo mes del año anterior al mes del informe) y el total Ene.-Dic. del año anterior, con todas
    las filas del ESTAD, en millones de US$.
  · gráficos (hoja ENLACE ESTAD de las SERIES): variación % anual de 12 años (el año en curso, enero-mes; Ucayali compara
    enero-mes en todos los años), los no tradicionales del año anterior y de enero-mes (6 principales + Otros) y los 5
    principales productos de enero-mes y del mes.
"""
import io
from decimal import ROUND_HALF_UP, Decimal

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font

from . import exportaciones

MESES = ["Enero", "Febrero", "Marzo", "Abril", "Mayo", "Junio", "Julio", "Agosto", "Setiembre", "Octubre", "Noviembre", "Diciembre"]
CORTOS = ["Ene.", "Feb.", "Mar.", "Abr.", "May.", "Jun.", "Jul.", "Ago.", "Set.", "Oct.", "Nov.", "Dic."]
DECIMALES = {"16": 2, "22": 1, "25": 1}
# cómo se llaman los grupos en el cuadro de la síntesis
EN_SINTESIS = {"Forestales": "Maderas y sus manufacturas"}
NOTA_OTRAS = "Incluye reparación de bienes de capital, manufacturas de plástico, partes de aeronaves, entre otros."


def _r(v, dec):
    """REDONDEAR de Excel (la mitad hacia arriba), no el redondeo bancario de Python."""
    return float(Decimal(repr(float(v))).quantize(Decimal(1).scaleb(-dec), rounding=ROUND_HALF_UP)) + 0.0   # sin «-0,0»


def _var(a, b, dec=1):
    if b in (None, 0) or a is None:
        return None
    return _r((a / b - 1) * 100, dec)


def _estructura(df):
    """Filas del ESTAD en su orden: [{orden, rubro, nivel, tipo, grupo}] y la serie {(año, mes): {orden: miles}}."""
    filas = df.drop_duplicates("ORDEN").sort_values("ORDEN")[["ORDEN", "RUBRO", "NIVEL", "TIPO", "GRUPO"]]
    filas = [dict(orden=int(r.ORDEN), rubro=r.RUBRO, nivel=r.NIVEL, tipo=r.TIPO, grupo=r.GRUPO) for r in filas.itertuples()]
    serie = df.pivot_table(index="ORDEN", columns=["AÑO", "MES"], values="FOB", aggfunc="sum").fillna(0.0)
    return filas, serie


def _suma(serie, orden, anio, m0, m1):
    cols = [(a, m) for a, m in serie.columns if a == anio and m0 <= m <= m1]
    return float(serie.loc[orden, cols].sum()) if cols and orden in serie.index else 0.0


def _productos(filas, solo_tipo=None):
    """Lo que la SERIES llama «producto»: las filas de producto y los grupos sin productos (no «Resto»)."""
    con_hijos = {(f["tipo"], f["grupo"]) for f in filas if f["nivel"] == "producto"}
    salida = []
    for f in filas:
        if solo_tipo and f["tipo"] != solo_tipo:
            continue
        if f["nivel"] == "producto" or (f["nivel"] == "grupo" and (f["tipo"], f["rubro"]) not in con_hijos and f["rubro"] != "Resto"):
            salida.append(f)
    return salida


def _ranking(filas, serie, total_orden, anio, m0, m1, n, otros=True):
    total = _suma(serie, total_orden, anio, m0, m1)
    vals = sorted(((f["rubro"], _suma(serie, f["orden"], anio, m0, m1)) for f in filas), key=lambda x: -x[1])
    top = [{"nombre": k, "valor": v} for k, v in vals if v > 0][:n]
    if otros:
        resto = total - sum(x["valor"] for x in top)
        if resto > 0.5:
            top.append({"nombre": "Otros", "valor": resto})
    return {"total": total, "partes": top}


def cuadros(depto, df, periodo=None, flujo="expo"):
    """depto: código ("16", "22", "25"); df: filas de la base expo-estad (o impo-estad) de ese departamento."""
    if df is None or df.empty:
        return None
    df = df.copy()
    df["AÑO"], df["MES"], df["ORDEN"] = df["AÑO"].astype(int), df["MES"].astype(int), df["ORDEN"].astype(int)
    df["FOB"] = pd.to_numeric(df["FOB"], errors="coerce").fillna(0.0)
    filas, serie = _estructura(df)
    total = next(f for f in filas if f["nivel"] == "total")
    con_datos = [(a, m) for (a, m) in serie.columns if serie.loc[total["orden"], (a, m)] != 0]
    if not con_datos:
        return None
    ultimo = max(con_datos)
    meses = sorted(f"{a}-{m:02d}" for a, m in con_datos)
    if periodo and periodo in meses:
        a, m = int(periodo[:4]), int(periodo[5:7])
    else:
        a, m = ultimo
    b = a - 1
    dec = DECIMALES.get(depto, 1) if flujo == "expo" else 1
    mill = lambda v: _r(v / 1000, dec)
    salida = {"periodo": f"{a}-{m:02d}", "ultimo": f"{ultimo[0]}-{ultimo[1]:02d}", "anio": a, "comparar": b, "mes": m,
              "titulo_mes": MESES[m - 1], "meses": meses, "dec": dec, "departamento": exportaciones.NOMBRE[depto], "flujo": flujo,
              "notas_estad": exportaciones.notas(f"{flujo}-estad", depto)}

    # ── síntesis: tipos y grupos ──
    tipos = [f for f in filas if f["nivel"] == "tipo"]
    lineas = []
    for t in tipos:
        hijos = [f for f in filas if f["tipo"] == t["rubro"] and f["nivel"] == ("grupo" if flujo == "expo" else "producto")]
        if flujo == "expo" and "NO TRADICIONALES" in t["rubro"]:
            hijos = sorted(hijos, key=lambda f: (f["rubro"] == "Resto", EN_SINTESIS.get(f["rubro"], f["rubro"])))
        lineas.append((t, 0))
        lineas += [(h, 1) for h in hijos]
    valores = {}
    for f, _ in lineas + [(total, 0)]:
        o = f["orden"]
        valores[o] = {"mes_b": mill(_suma(serie, o, b, m, m)), "mes_a": mill(_suma(serie, o, a, m, m)),
                      "acum_b": mill(_suma(serie, o, b, 1, m)), "acum_a": mill(_suma(serie, o, a, 1, m)),
                      "miles": [_suma(serie, o, b, m, m), _suma(serie, o, a, m, m), _suma(serie, o, b, 1, m), _suma(serie, o, a, 1, m)]}
    # el total: la suma de los tipos ya redondeados
    tot = valores[total["orden"]]
    for k in ("mes_b", "mes_a", "acum_b", "acum_a"):
        tot[k] = _r(sum(valores[t["orden"]][k] for t in tipos), dec)
    sint = []
    for f, nivel in lineas + [(total, 0)]:
        v = valores[f["orden"]]
        mi = v["miles"]
        sint.append({"nombre": EN_SINTESIS.get(f["rubro"], f["rubro"]) if nivel else f["rubro"], "nivel": nivel, "total": f is total,
                     "estructura": _r(v["mes_b"] / tot["mes_b"] * 100, 1) if tot["mes_b"] else None,
                     "mes_b": v["mes_b"], "mes_a": v["mes_a"], "var_mes": _var(v["mes_a"], v["mes_b"]),
                     "contrib_mes": _r((v["mes_a"] - v["mes_b"]) / tot["mes_b"] * 100, 1) if tot["mes_b"] else None,
                     "acum_b": v["acum_b"], "acum_a": v["acum_a"], "var_acum": _var(v["acum_a"], v["acum_b"]),
                     "contrib_acum": _r((v["acum_a"] - v["acum_b"]) / tot["acum_b"] * 100, 1) if tot["acum_b"] else None,
                     "var_exacta_mes": _var(mi[1], mi[0]), "var_exacta_acum": _var(mi[3], mi[2])})
    salida["sintesis"] = sint

    # ── anexo: 13 meses + total Ene.-Dic. del año anterior ──
    columnas = [{"etq": CORTOS[k - 1], "anio": b, "k": (b, k)} for k in range(m, 13)] + [{"etq": "Ene.-Dic.", "anio": "Total", "total": True, "k": None}] + \
               [{"etq": CORTOS[k - 1], "anio": a, "k": (a, k)} for k in range(1, m + 1)]

    def fila_anexo(f):
        vals = [round((_suma(serie, f["orden"], b, 1, 12) if c["k"] is None else _suma(serie, f["orden"], c["k"][0], c["k"][1], c["k"][1])) / 1000, 3)
                for c in columnas]
        return vals
    anexo = []
    for f in filas:
        if f["nivel"] == "total":
            continue
        anexo.append({"nombre": f["rubro"] + ("2/" if f["rubro"] == "Resto" else "3/" if f["rubro"].startswith("OTRAS") else ""),
                      "tipo": "grupo" if f["nivel"] == "tipo" else "item", "ind": {"tipo": 0, "grupo": 1, "producto": 2}[f["nivel"]] if flujo == "expo" else (0 if f["nivel"] == "tipo" else 1),
                      "valores": fila_anexo(f)})
    anexo.append({"nombre": total["rubro"], "tipo": "grupo", "ind": 0, "valores": fila_anexo(total)})
    salida["anexo"] = {"columnas": [{k: v for k, v in c.items() if k != "k"} for c in columnas], "filas": anexo}

    if flujo != "expo":
        return salida

    # ── gráficos ──
    anios = sorted({x for x, _ in serie.columns})
    todos_enemes = depto == "25"
    var_anual = []
    for y in [y for y in anios if y <= a][-12:]:
        m1 = m if (y == a or todos_enemes) else 12
        v, v0 = _suma(serie, total["orden"], y, 1, m1), _suma(serie, total["orden"], y - 1, 1, m1)
        var_anual.append({"anio": y, "etq": f"{y}*" if y == a and m < 12 else str(y), "meses": m1, "valor": v, "anterior": v0 if (y - 1) in anios else None,
                          "var": _var(v, v0) if (y - 1) in anios else None})
    salida["variacion_anual"] = var_anual
    nt = next((t for t in tipos if "NO TRADICIONALES" in t["rubro"]), None)
    if nt:
        prods = _productos(filas, nt["rubro"])
        salida["no_tradicionales"] = {"anterior": _ranking(prods, serie, nt["orden"], b, 1, 12, 6), "actual": _ranking(prods, serie, nt["orden"], a, 1, m, 6)}
    todos = _productos(filas)
    salida["principales"] = {"acumulado": _ranking(todos, serie, total["orden"], a, 1, m, 5, otros=False),
                             "mes": _ranking(todos, serie, total["orden"], a, m, m, 5, otros=False)}
    return salida


def excel(region, c):
    """El cuadro de la síntesis y el anexo en un Excel (millones de US$)."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Síntesis"
    neg = Font(bold=True)
    a, b, mes = c["anio"], c["comparar"], c["titulo_mes"]
    ws.append([f"{region}: {'EXPORTACIONES POR GRUPO DE PRODUCTOS' if c['flujo'] == 'expo' else 'IMPORTACIONES FOB SEGÚN USO O DESTINO ECONÓMICO'}"])
    ws.append(["(Valor FOB en millones de USD)"])
    ws.append(["TIPO", "Estructura %", mes, None, "Var. %", "Contribución", f"Enero-{mes}", None, "Var. %", "Contribución"])
    ws.append([None, b, b, a, None, None, b, a, None, None])
    for f in c["sintesis"]:
        ws.append([("   " * f["nivel"]) + f["nombre"], f["estructura"], f["mes_b"], f["mes_a"], f["var_mes"], f["contrib_mes"],
                   f["acum_b"], f["acum_a"], f["var_acum"], f["contrib_acum"]])
        if f["nivel"] == 0:
            ws.cell(ws.max_row, 1).font = neg
    fmt = "0." + "0" * c["dec"]
    for row in ws.iter_rows(min_row=5, min_col=2, max_col=10):
        for cell in row:
            cell.number_format = "0.0" if cell.column in (2, 5, 6, 9, 10) else fmt
            cell.alignment = Alignment(horizontal="right")
    ws.column_dimensions["A"].width = 34
    ax = wb.create_sheet("Anexo")
    ax.append([f"{region}: {'EXPORTACIONES POR GRUPO DE PRODUCTOS' if c['flujo'] == 'expo' else 'IMPORTACIONES FOB SEGÚN USO O DESTINO ECONÓMICO'}"])
    ax.append(["(Valor FOB - Millones de USD)"])
    ax.append([None] + [str(x["anio"]) for x in c["anexo"]["columnas"]])
    ax.append([None] + [x["etq"] for x in c["anexo"]["columnas"]])
    for f in c["anexo"]["filas"]:
        ax.append([("   " * f["ind"]) + f["nombre"]] + f["valores"])
        if f["tipo"] == "grupo":
            ax.cell(ax.max_row, 1).font = neg
    for row in ax.iter_rows(min_row=5, min_col=2):
        for cell in row:
            cell.number_format = "0.0"
    ax.column_dimensions["A"].width = 34
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
