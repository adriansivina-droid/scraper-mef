"""Excel de los «Cuadros y gráficos BCRP»: el mismo cuadro que dibuja static/bcrp.js, con el formato de la síntesis.

Recibe el tipo y los datos del cuadro (los mismos con que se dibujó el SVG) y arma la hoja:
  · cuadro: título centrado, filetes horizontales (cabecera y total), sangrías y negritas;
  · anexo: horizontal, cabecera gris con los años arriba, total Ene.-Dic. separado, rótulos repetidos a la derecha;
  · barras: la serie (mes y valor) y el gráfico de barras de Excel con los colores de la síntesis.
Los números van como números (no como texto), con los decimales que se ven en el cuadro.
"""
import io
import re

from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.chart.label import DataLabelList
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter as L

ELABORACION = "Elaboración: BCRP, Sucursal Iquitos. Departamento de Estudios Económicos."
F = "Arial"
NEG, NOR, TIT, NOTA = Font(name=F, size=9, bold=True), Font(name=F, size=9), Font(name=F, size=10, bold=True), Font(name=F, size=8)
GRIS = PatternFill("solid", fgColor="C0C0C0")
FINO = Side(style="thin", color="000000")
DER, IZQ, CEN = Alignment(horizontal="right"), Alignment(horizontal="left"), Alignment(horizontal="center", vertical="center")


def _numero(v):
    """«1 215» -> 1215, «-3,7» -> -3.7 (con su formato); «-», «n.d.» o texto quedan como texto."""
    if isinstance(v, (int, float)) and not isinstance(v, bool):
        return v, None
    t = str(v if v is not None else "").strip()
    if re.fullmatch(r"-?\d{1,3}(?:[  ]\d{3})*(?:,\d+)?|-?\d+(?:,\d+)?", t):
        dec = len(t.split(",")[1]) if "," in t else 0
        return float(t.replace(" ", "").replace(" ", "").replace(",", ".")), "#,##0" + ("." + "0" * dec if dec else "")
    return (t or None), None


def _encabezado(ws, spec, ancho, fila=1):
    """Número, título (con su llamada) y unidad, centrados sobre el ancho del cuadro. Devuelve la siguiente fila."""
    for texto, fuente in ((spec.get("numero"), TIT), (f"{spec.get('titulo', '')}{(' ' + spec['llamada']) if spec.get('llamada') else ''}", TIT), (spec.get("unidad"), NOR)):
        if not texto:
            continue
        ws.merge_cells(start_row=fila, start_column=1, end_row=fila, end_column=ancho)
        c = ws.cell(fila, 1, texto)
        c.font, c.alignment = fuente, CEN
        fila += 1
    return fila


def _notas(ws, notas, fila):
    for t in [*(notas or []), ELABORACION]:
        ws.cell(fila, 1, t).font = NOTA
        fila += 1


def _cuadro(ws, s):
    filas, cab = s.get("filas", []), s.get("cab", [])
    n = max([len(f.get("c", [])) for f in filas] + [sum(h.get("span", 1) for h in r) for r in cab] + [1])
    r = _encabezado(ws, s, n) + 1
    for k, fila in enumerate(cab):
        j = 1
        for h in fila:
            span = h.get("span", 1)
            if span > 1:
                ws.merge_cells(start_row=r, start_column=j, end_row=r, end_column=j + span - 1)
            c = ws.cell(r, j, h.get("t") if h.get("t") not in ("", None) else None)
            c.font, c.alignment = NEG, (IZQ if j == 1 else CEN if span > 1 else DER)
            for jj in range(j, j + span):
                b = ws.cell(r, jj)
                b.border = Border(top=FINO if k == 0 else None, bottom=FINO if (h.get("linea") or k == len(cab) - 1) else None)
            j += span
        r += 1
    for i, f in enumerate(filas):
        if f.get("sep") and i:
            r += 1
        ultima = i == len(filas) - 1
        for j, v in enumerate(f.get("c", []), start=1):
            val, fmt = (v, None) if j == 1 else _numero(v)
            c = ws.cell(r, j, val if val not in ("",) else None)
            c.font = NEG if f.get("b") else NOR
            c.alignment = Alignment(horizontal="left", indent=int(f.get("ind") or 0)) if j == 1 else DER
            if fmt:
                c.number_format = fmt
        for j in range(1, n + 1):
            ws.cell(r, j).border = Border(top=FINO if f.get("total") else None, bottom=FINO if (f.get("total") or ultima) else None)
        r += 1
    _notas(ws, s.get("notas"), r + 1)
    ws.column_dimensions["A"].width = max(18, min(48, max([len(str((f.get("c") or [""])[0])) + 2 * int(f.get("ind") or 0) for f in filas] + [10]) * 1.05))
    for j in range(2, n + 1):
        ws.column_dimensions[L(j)].width = 11


def _anexo(ws, s):
    cols, filas, dec = s.get("columnas", []), s.get("filas", []), int(s.get("dec", 1))
    n = len(cols) + 2                                   # rótulo, columnas y el rótulo repetido a la derecha
    if s.get("numero"):
        ws.cell(1, 1, s["numero"]).font = TIT
    r = _encabezado(ws, {**s, "numero": None}, n, fila=3)
    r0 = r
    # años (y «Total») arriba, meses abajo, sobre fondo gris
    grupos = []
    for j, c in enumerate(cols):
        clave = f"t{j}" if c.get("total") else str(c.get("anio"))
        if not grupos or grupos[-1][0] != clave:
            grupos.append([clave, "Total" if c.get("total") else str(c.get("anio")), j, j])
        else:
            grupos[-1][3] = j
    for _, texto, j0, j1 in grupos:
        if j1 > j0:
            ws.merge_cells(start_row=r, start_column=j0 + 2, end_row=r, end_column=j1 + 2)
        ws.cell(r, j0 + 2, texto)
    for j, c in enumerate(cols):
        ws.cell(r + 1, j + 2, c.get("etq"))
    for col in (1, n):
        ws.merge_cells(start_row=r, start_column=col, end_row=r + 1, end_column=col)
        ws.cell(r, col, s.get("rotulo") or None)
    for rr in (r, r + 1):
        for j in range(1, n + 1):
            c = ws.cell(rr, j)
            c.fill, c.font, c.alignment = GRIS, NEG, CEN
    r += 3
    formato = "#,##0" + ("." + "0" * dec if dec else "")
    for f in filas:
        tipo = f.get("tipo", "item")
        if tipo in ("grupo", "titulo"):
            r += 1 if r > r0 + 3 else 0
        negrita = tipo != "item"
        for col in (1, n):
            c = ws.cell(r, col, f.get("nombre"))
            c.font, c.alignment = (NEG if negrita else NOR), Alignment(horizontal="left", indent=0 if negrita else 1)
        for j, v in enumerate(f.get("valores") or [], start=2):
            c = ws.cell(r, j, v)
            c.font, c.alignment, c.number_format = (NEG if negrita else NOR), DER, formato
        if tipo in ("titulo", "total"):
            for j in range(1, n + 1):
                ws.cell(r, j).fill = GRIS
        r += 1
    # marco, separadores de los rótulos y del total
    separa = {1, 2, n}
    for j, c in enumerate(cols):
        if c.get("total"):
            separa |= {j + 2, j + 3}
    for rr in range(r0, r):
        for j in range(1, n + 1):
            c = ws.cell(rr, j)
            c.border = Border(left=FINO if j in separa else None, right=FINO if j == n else None,
                              top=FINO if rr in (r0, r0 + 2) else None, bottom=FINO if rr == r - 1 else None)
    _notas(ws, s.get("notas"), r + 1)
    ancho = max([len(str(f.get("nombre", ""))) for f in filas] + [12]) * 1.05
    for col in (1, n):
        ws.column_dimensions[L(col)].width = min(42, ancho)
    for j in range(2, n):
        ws.column_dimensions[L(j)].width = 9.5 if not cols[j - 2].get("total") else 11


def _barras(ws, s):
    puntos, dec = [p for p in s.get("puntos", []) if p.get("valor") is not None], int(s.get("dec", 1))
    r = _encabezado(ws, s, 8) + 1
    ws.cell(r, 1, "Mes").font = NEG
    ws.cell(r, 2, "Valor").font = NEG
    for k, p in enumerate(puntos, start=1):
        ws.cell(r + k, 1, p.get("periodo")).font = NOR
        c = ws.cell(r + k, 2, p["valor"])
        c.font, c.number_format = NOR, "#,##0" + ("." + "0" * dec if dec else "")
    ch = BarChart()
    ch.type, ch.style, ch.legend, ch.gapWidth = "col", 10, None, 50
    ch.title = None
    ch.add_data(Reference(ws, min_col=2, min_row=r, max_row=r + len(puntos)), titles_from_data=True)
    ch.set_categories(Reference(ws, min_col=1, min_row=r + 1, max_row=r + len(puntos)))
    serie = ch.series[0]
    serie.graphicalProperties.solidFill = "C6D9F1"
    serie.graphicalProperties.line.solidFill = "4F81BD"
    serie.invertIfNegative = False
    ch.dataLabels = DataLabelList()
    ch.dataLabels.showVal = True
    ch.dataLabels.showSerName = ch.dataLabels.showCatName = ch.dataLabels.showLegendKey = ch.dataLabels.showPercent = False
    ch.dataLabels.numFmt = "#,##0" + ("." + "0" * dec if dec else "")
    ch.y_axis.numFmt = "#,##0.0"
    ch.y_axis.majorGridlines = None
    ch.height, ch.width = 9, 22
    ws.add_chart(ch, f"D{r}")
    _notas(ws, s.get("notas"), r + len(puntos) + 2)
    ws.column_dimensions["A"].width = 12


def libro(tipo, spec):
    wb = Workbook()
    ws = wb.active
    ws.title = re.sub(r"[\[\]\*\?/\\:]", "", (spec.get("numero") or "Cuadro"))[:31] or "Cuadro"
    ws.sheet_view.showGridLines = False
    {"cuadro": _cuadro, "anexo": _anexo, "barras": _barras}[tipo](ws, spec)
    # al imprimir: una página de ancho (los anexos, horizontales, como en la síntesis)
    ws.page_setup.orientation = "landscape" if tipo in ("anexo", "barras") else "portrait"
    ws.page_setup.fitToWidth, ws.page_setup.fitToHeight = 1, 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
