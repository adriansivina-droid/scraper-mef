"""Ejecución del presupuesto: gasto corriente no financiero, inversión, variación real con IPC, avance del PIM,
comparación de años y variaciones alarmantes, con números hechos a mano para poder verificarlos."""
import io
from datetime import date

import pandas as pd
import pytest
from openpyxl import Workbook, load_workbook

from app import ipc
from tests.test_api import _csv, zip_de
from tests.test_hidrocarburos import cliente  # noqa: F401  (misma base vacía por prueba)

M = 1e5                                   # los montos de abajo × 100 mil soles: 1000 → S/ 100 millones
CARPETA = {"GobNacional": "Gobierno Nacional", "GobRegional": "Gobierno Regional", "Municipalidad": "Gobierno Municipal"}
NIVEL = {"GobNacional": "E: GOBIERNO NACIONAL", "GobRegional": "R: GOBIERNOS REGIONALES", "Municipalidad": "M: GOBIERNOS LOCALES"}
MESES = [(2025, m) for m in range(1, 13)] + [(2026, m) for m in range(1, 10)]


def scraper():
    s = {}
    for a, m in MESES:
        nuevo = a == 2026
        suf = f"{a} - Mes {m:02d} - LOR - SMN - UCA.csv"
        for corto, etq in NIVEL.items():
            base = dict(AÑO=a, MES=f"{m}: Mes", NIVEL_GOBIERNO=etq, DEPARTAMENTO="16: LORETO")
            personal = {"GobNacional": 100, "GobRegional": 200, "Municipalidad": 300}[corto] * (1.1 if nuevo else 1)
            gen = [dict(base, COD_GENERICA="5-21", GENERICA="5-21: PERSONAL Y OBLIGACIONES SOCIALES", DEVENGADO=personal * M),
                   dict(base, COD_GENERICA="5-23", GENERICA="5-23: BIENES Y SERVICIOS", DEVENGADO=(60 if nuevo else 50) * M),
                   dict(base, COD_GENERICA="6-26", GENERICA="6-26: ADQUISICION DE ACTIVOS NO FINANCIEROS", DEVENGADO=777 * M),
                   dict(base, COD_GENERICA="7-28", GENERICA="7-28: SERVICIO DE LA DEUDA PUBLICA", DEVENGADO=999 * M)]
            s[f"output/DEVENGADO/{CARPETA[corto]}/DEV Generica.{corto} {suf}"] = _csv(pd.DataFrame(gen))
            if corto == "Municipalidad":
                fun = [dict(AÑO=a, MES=f"{m}: Mes", DEPARTAMENTO="16: LORETO", FUNCION=f, MUNICIPALIDAD=mu, DEVENGADO=v * M)
                       for f, mu, v in [("15: TRANSPORTE", "160101-301263: MUNICIPALIDAD PROVINCIAL DE MAYNAS", 300 if nuevo else 200),
                                        ("16: COMUNICACIONES", "160101-301263: MUNICIPALIDAD PROVINCIAL DE MAYNAS", 10),
                                        ("20: SALUD", "160108-301270: MUNICIPALIDAD DISTRITAL DE PUNCHANA", 50 if nuevo else 100)]]
            else:
                f, v = ("15: TRANSPORTE", 1500 if nuevo else 1000) if corto == "GobNacional" else ("22: EDUCACION", 400 if nuevo else 500)
                fun = [dict(base, ENTIDAD="X", COD_FUNCION=f[:2], FUNCION=f, DEVENGADO=v * M)]
            s[f"output/DEVENGADO/{CARPETA[corto]}/DEV Funcion.{corto} {suf}"] = _csv(pd.DataFrame(fun))
    pim = [dict(AÑO=a, DEPARTAMENTO="16: LORETO", **{"NIVEL DE GOBIERNO": n}, FUNCION=f, PIM=v * M)
           for a in (2025, 2026) for n, f, v in [("GOBIERNO NACIONAL", "15: TRANSPORTE", 20000), ("GOBIERNOS REGIONALES", "22: EDUCACION", 10000),
                                                ("GOBIERNOS LOCALES", "15: TRANSPORTE", 5000), ("GOBIERNOS LOCALES", "20: SALUD", 2000)]]
    s["output/PIM/PIM X FUNCIÓN/DATA/PIM Funcion 2026 - LOR - SMN - UCA.csv"] = _csv(pd.DataFrame(pim))
    return s


def ipc_xlsx(valor_2026=105.0):
    wb = Workbook()
    ws = wb.active
    for _ in range(3):
        ws.append(["IPC Lima Metropolitana"])
    for a, m in MESES:
        ws.append([None, date(a, m, 1), valor_2026 if a == 2026 else 100.0])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


@pytest.fixture()
def datos(cliente):  # noqa: F811
    r = cliente.post("/v1/cargas", files=[("archivos", ("output.zip", zip_de(scraper()), "application/zip")),
                                          ("archivos", ("IPC.xlsx", ipc_xlsx(), "application/octet-stream"))])
    assert r.status_code == 200, r.text
    assert any("IPC de Lima" in a for a in r.json()["avisos"])
    return cliente


def test_gasto_corriente_real_y_nominal(datos):
    c = datos.get("/v1/cuadros/ejecucion?departamento=16&sector=gasto").json()["cuadros"]
    assert c["periodo"] == "2026-09" and c["comparar"] == 2025 and c["real"] and c["anios"] == [2025, 2026]
    total = c["niveles"][-1]
    # sin capital ni deuda: (100+200+300) + 3×50 = 750 en 2025; (660) + 3×60 = 840 en 2026 (× 100 mil = millones / 10)
    assert (total["mes_b"], total["mes_a"]) == (75.0, 84.0) and total["acum_a"] == 756.0
    assert total["var_mes_nominal"] == 12.0 and total["var_mes"] == round((840 / 105 / 7.5 - 1) * 100, 1) == 6.7
    assert [f["nombre"] for f in c["detalle"]] == ["Personal y obligaciones sociales", "Bienes y servicios", "Total"]
    assert c["detalle"][0]["var_acum"] == 4.8                         # 1,10 / 1,05
    assert [p["nivel"] for p in c["participacion"]] == ["G.N.", "G.R.", "G.L."]
    assert c["mensual"]["a"][8] == 84.0 and c["mensual"]["a"][9] is None and c["mensual"]["acum_b"][11] == 900.0
    assert c["interanual"][-1] == {"periodo": "Set. 2026", "var": 6.7, "var_nominal": 12.0}


def test_inversion_avance_pim_y_comparar(datos):
    c = datos.get("/v1/cuadros/ejecucion?departamento=16&sector=inversion&periodo=2026-06").json()["cuadros"]
    n = {f["nombre"]: f for f in c["niveles"]}
    assert n["Gobierno Nacional"]["acum_a"] == 900.0 and n["Gobierno Nacional"]["pim"] == 2000.0 and n["Gobierno Nacional"]["avance"] == 45.0
    assert n["Gobierno Nacional"]["avance_b"] == 30.0
    d = {f["nombre"]: f for f in c["detalle"]}
    assert d["15: TRANSPORTE"]["acum_a"] == 900.0 + 186.0               # G.N. + G.L. (300 + 10 de comunicaciones) × 6
    assert "16: COMUNICACIONES" not in d
    assert c["avance"]["parejo"][5] == 50.0 and c["avance"]["a"][5] == round(c["mensual"]["acum_a"][5] / 3700 * 100, 1)
    assert datos.get("/v1/cuadros/ejecucion?departamento=16&sector=inversion&comparar=2026").status_code == 422
    assert datos.get("/v1/cuadros/ejecucion?departamento=16&sector=inversion&periodo=2027-01").status_code == 422


def test_variaciones_alarmantes(datos):
    v = datos.get("/v1/variaciones?departamento=16").json()["variaciones"]
    mes, acum = v["comparaciones"]
    assert mes["titulo"] == "Setiembre 2026 vs Setiembre 2025" and acum["clave"] == "acumulado"
    todos = mes["bloques"][0]
    assert todos["total"] == {"antes": 181.0, "ahora": 226.0, "dif": 45.0, "var": 24.9}
    assert todos["aumentos"][0]["item"] == "G.N. · 15: TRANSPORTE" and todos["aumentos"][0]["dif"] == 50.0
    assert todos["pocos"] and len(todos["aumentos"]) + len(todos["caidas"]) == 4      # 4 ítems, ninguno repetido
    muni = mes["bloques"][-1]
    assert muni["nombre"].endswith("por municipalidad")
    assert [f["item"] for f in muni["aumentos"]][0] == "MUNICIPALIDAD PROVINCIAL DE MAYNAS"


def test_sin_ipc_queda_nominal(cliente):  # noqa: F811
    cliente.post("/v1/cargas", files=[("archivos", ("output.zip", zip_de(scraper()), "application/zip"))])
    c = cliente.get("/v1/cuadros/ejecucion?departamento=16&sector=gasto").json()["cuadros"]
    assert not c["real"] and c["niveles"][-1]["var_mes"] is None and c["niveles"][-1]["var_mes_nominal"] == 12.0


def test_ipc_del_bcrp(cliente, monkeypatch):  # noqa: F811
    monkeypatch.setattr(ipc, "de_bcrp", lambda: {"2026-08": 113.1, "2026-09": 113.4, "2026-10": 0.2})   # 0,2 % no es un índice
    r = cliente.post("/v1/ipc/actualizar").json()
    assert r["meses"] == 2 and r["hasta"] == "2026-09" and r["fuente"].startswith("BCRP")


def test_excel(datos):
    r = datos.get("/v1/cuadros/ejecucion?departamento=16&sector=inversion&formato=xlsx")
    wb = load_workbook(io.BytesIO(r.content))
    assert wb.sheetnames == ["Cuadro", "Mensual", "Variaciones mes", "Variaciones acumulado"]
    assert "2026-09 vs 2025" in r.headers["content-disposition"]


def test_lee_la_respuesta_del_bcrp(monkeypatch):
    class Resp:
        def raise_for_status(self):
            pass

        def json(self):
            return {"config": {}, "periods": [{"name": "Ene.2026", "values": ["112.51"]}, {"name": "Feb.2026", "values": ["n.d."]},
                                              {"name": "Set.2026", "values": ["113.40"]}]}

    class Sesion:
        def get(self, url, **k):
            assert "/PN01270PM/json/2010-1/" in url
            return Resp()

    monkeypatch.undo()                                    # la función real (conftest la reemplaza para no salir a la red)
    assert ipc.de_bcrp(sesion=Sesion()) == {"2026-01": 112.51, "2026-09": 113.4}


def test_proyectos_del_notebook(datos, tmp_path, monkeypatch):
    """Bloques 1 a 3 de Fiscal_Proyectos.ipynb de verdad (con los CSV armados desde la base); la consulta al MEF
    (bloques 4 y 5) se simula. Luego el endpoint los muestra."""
    import sys
    from pathlib import Path
    sys.path.insert(0, str(Path(__file__).resolve().parents[2] / "automatizacion"))
    import proyectos_servidor as ps
    raiz = tmp_path / "Fiscal"
    ps.preparar_csv(raiz, "2026-09")
    assert sorted(p.name for p in (raiz / "Scrapeo/output/DEVENGADO/Gobierno Nacional").iterdir()) == [
        "DEV Funcion.GobNacional 2025 - Mes 09 - LOR - SMN - UCA.csv", "DEV Funcion.GobNacional 2026 - Mes 09 - LOR - SMN - UCA.csv"]
    ps.sustituir_colab()
    g = {"__name__": "__main__", "display": lambda *x, **k: None}
    celdas = ps.celdas("2026-09", raiz)
    assert len(celdas) == 5 and "MES_ANALISIS = '2026-09'" in celdas[1]
    for codigo in celdas[:3]:                                  # librerías, configuración y funciones a abrir
        exec(compile(codigo, "nb", "exec"), g)
    plan = {(p["cod_depto"], p["nivel"]["codigo"]): p for p in g["plan"]}
    assert plan[("16", "E")]["positivo"] and plan[("16", "E")]["funciones"][0]["cod"] == "15"
    assert not plan[("16", "R")]["positivo"] and plan[("16", "R")]["fecha"].year == 2025      # cayó: se consulta 2025
    for p in g["plan"]:                                        # bloque 5 simulado
        for f in p["funciones"]:
            f["proyectos"], f["devengado_funcion"] = [("2001621", f"PROYECTO DE {f['nombre']}", 80e6), ("2195", "OTRO", 20e6)], 100e6
    df = ps.filas(g)
    from app import db
    db.guardar("inv-proyectos", df)
    r = datos.get("/v1/proyectos?departamento=16").json()
    assert r["periodo"] == "2026-09" and [n["nivel"] for n in r["proyectos"]] == ["GOBIERNO NACIONAL", "GOBIERNO REGIONAL", "GOBIERNOS LOCALES"]
    gn = r["proyectos"][0]
    assert gn["consultado"] == "Set-26" and gn["funciones"][0]["proyectos"][0] == {"codigo": "2001621", "proyecto": "PROYECTO DE 15: TRANSPORTE (incl. 16: Comunicaciones)",
                                                                                 "devengado": 80e6, "pct": 80.0}
    gr = r["proyectos"][1]
    assert gr["funciones"] == [] and gr["consultado"] == "Set-25" and gr["nota"] == "Ninguna función en la lista de caídas"
