"""Pruebas de pesca con Excel que imitan los crudos de las DIREPRO y los LONG de los notebooks de PESCA/."""
import io
import os

import pandas as pd
import pytest
from openpyxl import Workbook, load_workbook
from sqlalchemy import text

from app import db
from tests.test_api import cliente_http


def _libro(hojas):
    """hojas: {nombre: [filas]} -> bytes de un .xlsx con esas celdas tal cual."""
    wb = Workbook()
    wb.remove(wb.active)
    for nombre, filas in hojas.items():
        ws = wb.create_sheet(nombre)
        for f in filas:
            ws.append(f)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def loreto_consumo(mes="FEBRERO", anio=2026, boquichico=10.0):
    def cuadro(titulo, especies):
        total = sum(v for _, v in especies)
        return [[], [titulo], [f"PERIODO  :  {mes} - {anio}"], ["N°", "Especie", "Iquitos", "Total"],
                [len(especies), "Total", total, total]] + [[i, e, v, v] for i, (e, v) in enumerate(especies, 1)] + [[None, "NOTA :"]]
    return _libro({"Desemb_Fres": cuadro("DESEMBARQUE FRESCO", [("Boquichico", boquichico), ("Manitoa", 2.0), ("Bagre", 1.0)]),
                   "Desemb_salpr": cuadro("DESEMBARQUE SALPRESO", [("Palometa", 0.5)]),
                   "Desemb_secosalad": cuadro("DESEMBARQUE SECO SALADO", [("Paiche", 4.0)])})


def loreto_ornamentales():
    return _libro({"POR ESPECIES": [
        [], ["INGRESO, EGRESO Y VALOR"], ["PERIODO : SETIEMBRE - 2025."],
        ["N°", "NombreComunBD", None, "INGRESO", None, "EGRESOS"],
        [None, None, None, None, None, "EXPORTACION", None, "NACIONAL", None, "TOTAL DE EGRESOS", None, "STOCK"],
        [None, None, "Nombre cientifico BD", "UNIDADES", "VALOR S/.", "UNIDADES", "VALOR S/.", "UNIDADES", "VALOR S/.", "UNIDADES", "%", "UNIDADES"],
        [1, "otocinclos comun", "Otocinclus vestitus", 3000, 10, 2500, 8, 500, 2, 3000, 0.6, 0],
        [2, "sturiosoma", "Sturisoma sp.", 1000, 5, 1000, 5, None, None, 1000, 0.2, 0],
        [3, "sturiosoma", "Sturisoma nigrirostrum", 1000, 5, 800, 4, None, None, 800, 0.2, 200],
        [3, "TOTAL", None, 5000, 20, 4300, 17, 500, 2, 4800, 1, 200],
        [None, "Principales especies"], [1, "otocinclos comun", "Otocinclus vestitus", 3000]]})


def smn_crudo():
    col = ["PROCEDENCIA", "DEPARTAMENTO", "ENCABEZADO", "ESPECIE", "FRESCO", "FRESCO REFRIGERADO", "SALPRESO", "SECO SALADO"]
    return _libro({"SMN_PESCA": [col,
        ["CONTINENTAL", None, 1, "CONTINENTAL", 5000, 9000, 0, 2000],
        ["ACUICULTURA CONTINENTAL", "SAN MARTÍN", 1, "ACUICULTURA", 5000, 1000, 0, 0],
        ["ACUICULTURA CONTINENTAL", "SAN MARTÍN", 0, "TILAPIA", 4000, 1000, None, None],
        ["ACUICULTURA CONTINENTAL", "SAN MARTÍN", 0, "TRUCHA", 1000, None, None, None],
        ["CONTINENTAL", "SAN MARTÍN", 0, "BOQUICHICO", None, 500, None, None],
        ["CONTINENTAL", "PUCALLPA", 1, "PUCALLPA", 0, 7500, 0, 2000],
        ["CONTINENTAL", "PUCALLPA", 0, "BOQUICHICO", None, 7500, None, 2000],
        ["TOTAL", None, 1, "TOTAL", 10000, 18000, 0, 4000]]})


def uca_crudo(anio=2026):
    meses = [f"{anio}-{m:02d}-01" for m in range(1, 13)]
    def cuadro(titulo, filas):
        return [[titulo], [None, "( t )"], ["ESPECIES", *meses, "TOTAL"]] + filas + [["TOTAL (tm)"]]
    datos = lambda *v: list(v) + [0] * (12 - len(v))
    return _libro({"TOTAL DE PESCADO ": cuadro("CUADRO 01", [["BAGRE", *datos(3, 4)], ["NUEVA ESPECIE", *datos(1)]]),
                   "TOTAL FRESCO": cuadro("CUADRO 02", [["BAGRE", *datos(2, 3)], ["NUEVA ESPECIE", *datos(1)]]),
                   "TOTAL SECOSALADO": cuadro("CUADRO 03", [["BAGRE", *datos(1, 1)], ["NUEVA ESPECIE", *datos(0)]]),
                   "ORNAMENTALES": [["CUADRO 37"], [None, None, "(Unidades)"], ["N/O", "ESPECIES", *meses], [1, "ABRAMITES", None, 50], [None, "TOTAL"]]})


def uca_long_historico():
    filas = [dict(AÑO=2025, MES=m, TIPO=t, ESPECIES="BAGRE", ESPECIES_CORREGIDO="Bagre", ESPECIES_OK="Bagre", VOLUMEN=1.0, VBP=0.0)
             for m in (1, 2) for t in ("FRESCO", "SECO-SALADO", "AGREGADO")]
    buf = io.BytesIO()
    pd.DataFrame(filas).to_excel(buf, index=False, sheet_name="UCA_PESCA")
    return buf.getvalue()


@pytest.fixture()
def cliente(tmp_path, monkeypatch):
    monkeypatch.setenv("DATABASE_URL", os.environ.get("PRUEBA_DATABASE_URL") or f"sqlite:///{tmp_path / 'pesca.db'}")
    monkeypatch.delenv("APP_CLAVE", raising=False)
    db.reiniciar_motor()
    if db.motor().dialect.name == "postgresql":
        with db.motor().begin() as cx:
            cx.execute(text("DROP SCHEMA public CASCADE; CREATE SCHEMA public"))
        db.reiniciar_motor()
    yield cliente_http()
    db.reiniciar_motor()


def subir(c, **archivos):
    r = c.post("/v1/cargas", files=[("archivos", (n, b, "application/octet-stream")) for n, b in archivos.items()])
    assert r.status_code == 200, r.text
    return r.json()


def datos(c, base, depto):
    j = c.get(f"/v1/descargas?base={base}&departamento={depto}&desde=2020-01&hasta=2026-12&formato=json").json()
    return pd.DataFrame(j[base]["datos"])


def test_loreto_consumo_y_ornamentales(cliente):
    r = subir(cliente, **{"CONSUMO_FEBRERO_-_2026-OPP.xlsx": loreto_consumo(), "P._ORNAMENTALES_SETIEMBRE_-_OPP_2025.xlsx": loreto_ornamentales()})
    assert r["bases"]["pesca-loreto"]["periodos"] == ["2025-09", "2026-02"] and r["avisos"] == []
    d = datos(cliente, "pesca-loreto", "16")
    fre = d[(d.MES == 2) & (d.ESTADO == "FRESCO")].set_index("ESPECIE")
    assert fre.loc["Boquichico", "VBP"] == pytest.approx(10 * 3.27504)
    assert fre.loc["Manitoa", "ESPECIE_OK"] == "Zúngaro Manitoa" and fre.loc["Bagre", "ESPECIE_OK"] == "Otros"
    assert fre.loc["Boquichico", "ESPECIE_CUADRO"] == "Boquichico"
    orn = d[d.ESTADO == "ORNAMENTAL"]
    assert orn.VOLUMEN.sum() == pytest.approx(5.0)                       # miles de unidades, como el TOTAL
    assert len(orn) == 5                                                  # nombre común repetido: el científico las separa
    assert orn.groupby("DESTINO").VOLUMEN.sum().to_dict() == pytest.approx(
        {"Mercado externo": 4.3, "Mercado interno": 0.5, "Stock remanente y/o mortandad": 0.2})
    assert orn.VBP.sum() == pytest.approx(5.0 * 1.0329)


def test_loreto_recargar_consumo_no_borra_ornamentales(cliente):
    subir(cliente, **{"orn.xlsx": loreto_ornamentales(), "consumo.xlsx": loreto_consumo(mes="SETIEMBRE", anio=2025)})
    subir(cliente, **{"consumo.xlsx": loreto_consumo(mes="SETIEMBRE", anio=2025, boquichico=20.0)})
    d = datos(cliente, "pesca-loreto", "16")
    assert set(d.ESTADO) == {"FRESCO", "SALPRESO", "SECO SALADO", "ORNAMENTAL"}
    assert d[d.ESPECIE == "Boquichico"].VOLUMEN.tolist() == [20.0]


def test_san_martin(cliente):
    r = subir(cliente, **{"SMN_PESCA_2026-08.xlsx": smn_crudo()})
    assert r["bases"]["pesca-san-martin"]["periodos"] == ["2026-08"]
    d = datos(cliente, "pesca-san-martin", "22")
    sm = d[d.DEPARTAMENTO == "SAN MARTÍN"]
    assert sorted(zip(sm.ESPECIE, sm.ESTADO, sm.VOLUMEN)) == [("BOQUICHICO", "FRESCO REFRIGERADO", 0.5), ("TILAPIA", "FRESCO", 4.0),
                                                             ("TILAPIA", "FRESCO REFRIGERADO", 1.0), ("TRUCHA", "FRESCO", 1.0)]
    assert set(sm.ESTADO_OK) == {"FRESCO"} and sm.VBP.sum() == pytest.approx(6.5 * 3.27504)
    assert dict(zip(sm.ESPECIE, sm.ESPECIE_OK)) == {"TILAPIA": "Tilapia", "TRUCHA": "Otras", "BOQUICHICO": "Boquichico"}
    otros = d[d.DEPARTAMENTO == "Otros"]                                  # Pucallpa: su total por estado, sin VBP
    assert len(otros) == 4 and otros.VOLUMEN.sum() == pytest.approx(9.5) and otros.VBP.sum() == 0


def test_san_martin_sin_mes_en_el_nombre(cliente):
    r = subir(cliente, **{"crudo_smn.xlsx": smn_crudo()})
    assert "pesca-san-martin" not in r["bases"] and any("año y el mes" in a for a in r["avisos"])


def test_ucayali_usa_el_historico_y_recorta_meses_vacios(cliente):
    subir(cliente, **{"UCA_PESCA_LONG.xlsx": uca_long_historico()})
    r = subir(cliente, **{"ESTADISTICA_PESQUERA_DIREPRO_UCA_202601-202602.xlsx": uca_crudo()})
    assert r["bases"]["pesca-ucayali"]["periodos"] == ["2026-01", "2026-02"]          # marzo-diciembre en cero: no se guardan
    assert any("NUEVA ESPECIE" in a for a in r["avisos"])
    d = datos(cliente, "pesca-ucayali", "25")
    bagre = d[(d.AÑO == 2026) & (d.ESPECIES == "BAGRE")].set_index(["TIPO", "MES"])
    assert bagre.loc[("FRESCO", 1), "ESPECIES_OK"] == "Bagre" and bagre.loc[("FRESCO", 2), "VBP"] == pytest.approx(3 * 3.27504)
    assert bagre.loc[("SECO-SALADO", 1), "VBP"] == pytest.approx(2.98778) and bagre.loc[("AGREGADO", 1), "VBP"] == 0
    nueva = d[(d.ESPECIES == "NUEVA ESPECIE") & (d.TIPO == "FRESCO")]
    assert set(nueva.ESPECIES_OK) == {"Otros"} and set(nueva.ESPECIES_CORREGIDO) == {"Nueva Especie"}
    assert d[d.TIPO == "ORNAMENTAL"][["MES", "VOLUMEN"]].values.tolist() == [[2, 50]]
    assert len(d[d.AÑO == 2025]) == 6                                                   # el histórico queda


def test_excel_de_pesca(cliente):
    subir(cliente, **{"consumo.xlsx": loreto_consumo(), "SMN_PESCA_2026-08.xlsx": smn_crudo()})
    v = cliente.get("/v1/descargas/vista-previa?base=pesca-loreto,pesca-san-martin&departamento=16,25&desde=2026-01&hasta=2026-12").json()
    assert v["avisos"] == ["Pesca San Martín: solo tiene datos de San Martín."]       # Ucayali no es de Loreto: sin aviso
    assert v["descarga"] == "bases_pesca_2026-01_2026-12.xlsx"
    r = cliente.get("/v1/descargas?base=pesca-loreto&departamento=16&desde=2026-01&hasta=2026-12&formato=xlsx")
    ws = load_workbook(io.BytesIO(r.content))["PESCA Loreto"]
    assert ws["M1"].value == "VBP" and ws["M2"].number_format == "#,##0.000"


def test_excel_que_no_es_de_pesca(cliente):
    r = subir(cliente, **{"otro.xlsx": _libro({"Hoja1": [["A", "B"], [1, 2]]})})
    assert r["bases"] == {} and any("ni un Excel de pesca" in a for a in r["avisos"])


class DriveFalsa:
    """Imita LectorDrive: {carpeta: {file_id: (nombre, modificado, bytes)}}."""
    def __init__(self, carpetas):
        self.carpetas, self.bajados = carpetas, []

    def listar(self, carpeta):
        return [{"id": i, "name": n, "modifiedTime": m} for i, (n, m, _) in self.carpetas.get(carpeta, {}).items()]

    def bajar(self, file_id):
        self.bajados.append(file_id)
        return next(c[file_id][2] for c in self.carpetas.values() if file_id in c)


@pytest.fixture()
def drive_falsa(cliente, monkeypatch):
    from app import drive
    falsa = DriveFalsa({drive.CARPETAS["pesca-loreto"]: {"L1": ("CONSUMO_FEBRERO_-_2026-OPP.xlsx", "2026-03-10T10:00:00Z", loreto_consumo())},
                        drive.CARPETAS["pesca-san-martin"]: {"S1": ("SMN_PESCA_2026-08.xlsx", "2026-09-05T10:00:00Z", smn_crudo())},
                        drive.CARPETAS["pesca-ucayali"]: {}})
    monkeypatch.setenv("GOOGLE_DRIVE_CREDENCIAL", "{}")
    monkeypatch.setattr(drive, "LectorDrive", lambda credencial: falsa)
    return falsa


def test_drive_sin_credencial(cliente, monkeypatch):
    monkeypatch.delenv("GOOGLE_DRIVE_CREDENCIAL", raising=False)
    assert cliente.get("/v1/drive").json()["configurado"] is False
    assert cliente.post("/v1/drive/sincronizar").status_code == 503


def test_drive_sincroniza_solo_lo_nuevo(cliente, drive_falsa):
    from app import drive
    r = cliente.post("/v1/drive/sincronizar").json()
    assert sorted(r["bases"]) == ["pesca-loreto", "pesca-san-martin"] and len(r["archivos"]) == 2
    assert cliente.get("/v1/estado").json()["bases"]["pesca-san-martin"]["hasta"] == "2026-08"
    r = cliente.post("/v1/drive/sincronizar").json()                         # nada cambió: no baja nada
    assert r["bases"] == {} and r["archivos"] == [] and drive_falsa.bajados == ["L1", "S1"]
    # DIREPRO corrige el archivo de Loreto y llega uno nuevo de Ucayali
    drive_falsa.carpetas[drive.CARPETAS["pesca-loreto"]]["L1"] = ("CONSUMO_FEBRERO_-_2026-OPP.xlsx", "2026-03-20T10:00:00Z", loreto_consumo(boquichico=30.0))
    drive_falsa.carpetas[drive.CARPETAS["pesca-ucayali"]]["U1"] = ("ESTADISTICA_UCA_2026.xlsx", "2026-10-01T10:00:00Z", uca_crudo())
    r = cliente.post("/v1/drive/sincronizar?origen=GitHub Actions").json()
    assert sorted(a["nombre"] for a in r["archivos"]) == ["CONSUMO_FEBRERO_-_2026-OPP.xlsx", "ESTADISTICA_UCA_2026.xlsx"]
    assert datos(cliente, "pesca-loreto", "16").query("ESPECIE == 'Boquichico'").VOLUMEN.tolist() == [30.0]
    e = cliente.get("/v1/drive").json()
    assert e["configurado"] and e["ultima"]["origen"] == "GitHub Actions" and len(e["carpetas"]) == 3
    r = cliente.post("/v1/drive/sincronizar?todo=true").json()                # volver a procesar todo
    assert len(r["archivos"]) == 3


def test_drive_mismo_nombre_en_dos_carpetas(cliente, drive_falsa):
    from app import drive
    drive_falsa.carpetas[drive.CARPETAS["pesca-ucayali"]]["X"] = ("SMN_PESCA_2026-08.xlsx", "2026-10-01T10:00:00Z", uca_crudo())
    r = cliente.post("/v1/drive/sincronizar").json()
    assert sorted(r["bases"]) == ["pesca-loreto", "pesca-san-martin", "pesca-ucayali"]


def test_drive_cuenta_sin_acceso_a_la_carpeta(cliente, monkeypatch):
    """Drive responde 404 a una carpeta que la cuenta no ve: debe ser un error claro, no «no hay archivos nuevos»."""
    from app import drive

    class Respuesta:
        def __init__(self, codigo, datos=None):
            self.status_code, self._datos = codigo, datos or {}
        def raise_for_status(self):
            if self.status_code >= 400:
                raise RuntimeError(self.status_code)
        def json(self):
            return self._datos

    class Sesion:
        def get(self, url, params=None, timeout=None):
            if url.endswith(drive.CARPETAS["pesca-san-martin"]):
                return Respuesta(404)
            if "alt" not in (params or {}) and "q" not in (params or {}):
                return Respuesta(200, {"mimeType": "application/vnd.google-apps.folder"})
            return Respuesta(200, {"files": []})

    lector = drive.LectorDrive.__new__(drive.LectorDrive)
    lector.sesion = Sesion()
    assert lector.listar(drive.CARPETAS["pesca-loreto"]) == []
    monkeypatch.setenv("GOOGLE_DRIVE_CREDENCIAL", "{}")
    monkeypatch.setattr(drive, "LectorDrive", lambda credencial: lector)
    r = cliente.post("/v1/drive/sincronizar")
    assert r.status_code == 502 and "no puede ver la carpeta" in r.json()["detail"]


def test_drive_dice_cuantos_excel_hay_en_cada_carpeta(cliente, drive_falsa):
    cliente.post("/v1/drive/sincronizar")
    r = cliente.post("/v1/drive/sincronizar").json()
    assert r["carpetas"] == {"pesca-loreto": 1, "pesca-san-martin": 1, "pesca-ucayali": 0}
    assert r["avisos"] == ["No hay archivos nuevos ni modificados en las carpetas de Drive "
                           "(Excel en cada carpeta · Loreto: 1, San Martín: 1, Ucayali: 0)."]
