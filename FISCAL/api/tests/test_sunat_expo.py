"""Pruebas de la base mensual de SUNAT: diccionarios de ANUAL_PRO, estados de las filas, revisión, ESTAD propuesto."""
import importlib.util
import io
import zipfile
from pathlib import Path

from openpyxl import Workbook, load_workbook

from app import sunat_expo
from tests.test_exportaciones import estadlo
from tests.test_pesca import cliente, datos, subir  # noqa: F401

CAB = ["aduana", "ubigeo", "ruc11", "ano_mes", "fecembar", "pais", "nandi07", "partorig", "dcom", "fob", "kilos", "qunifis",
       "tunifis", "ruc", "secx_3", "producto", "ubigeo_dep"]


def fila(ruc, partida, dcom, fob, dep="16", pais="105", ano_mes="2608"):
    return ["226", "160113", ruc, ano_mes, ano_mes + "10", pais, partida, partida, dcom, fob, 100, 10, "U", ruc[2:10], "500", "R_500", dep]


def anual_pro():
    """Un ANUAL_PRO mínimo: UNICOS con sus tablas lado a lado, PARTIDA y ORIGEN (encabezado en la fila 7)."""
    wb = Workbook()
    o = wb.active
    o.title = "ORIGEN"
    for _ in range(6):
        o.append([])
    o.append(CAB + ["FOB (MILES)"])
    o.append(fila("20513842377", "2709000000", "PETROLEO CRUDO", 17033000, ano_mes="2607"))
    u = wb.create_sheet("UNICOS")
    cab = ["PRODUCTO_ORIGINAL", "PRODUCTO_CORREGIDO", None, None, None, "COD", "COD_TEXTO", "PAIS", "CONTINENTE", None, None, None,
           "RUC", "RUC_TEXTO", "RAZON_SOCIAL", "DOMICILIO_FISCAL", "DPTO_EXPORTACION", "RUC_NOMBRE", "BORRAR_PARA", "NOTA", "REVISADO",
           None, None, None, "COD", "ADUANA_JURISDICCION"]
    u.append(cab)
    empresas = [("20513842377", "PETROTAL PERU S.R.L.", "OK"), ("20100000001", "CACAO LORETO SAC", "OK"),
                ("20601551951", "REYCAL COMPANY S.A.C.", "LOR SMN UCA"), ("20100000003", "CAFETALERA SAC", "OK")]
    for i in range(4):
        r = [None] * len(cab)
        r[0], r[1] = ("CACAO EN GRANO SECO", "Cacao en grano") if i == 0 else (None, None)
        r[5], r[6], r[7], r[8] = (105, "105", "Brasil", "América") if i == 0 else (None, None, None, None)
        r[12], r[13], r[14], r[18] = empresas[i][0], empresas[i][0], empresas[i][1], empresas[i][2]
        r[24], r[25] = (226, "Iquitos") if i == 0 else (None, None)
        u.append(r)
    p = wb.create_sheet("PARTIDA")
    p.append(["PARTIDA", "PARTIDA_TEXTO", "|", "CATEGORIA", "PRODUCTO ✔ ", "PRODUCTO", "NOTA"])
    p.append([2709000000, "2709000000", "1. PRODUCTOS TRADICIONALES", "1.2. PETRÓLEO Y SUS DERIVADOS", "Barriles de petróleo crudo", "- Aceites crudos", None])
    p.append([1801001900, "1801001900", "2. PRODUCTOS NO TRADICIONALES", "2.4. AGROPECUARIOS", "Cacao", "- - Los demás", None])
    p.append([901119000, "0901119000", "1. PRODUCTOS TRADICIONALES", "1.1. CAFÉ", "Café en grano", "- - Los demás", None])
    p.append([2523290000, "2523290000", "3. OTROS PRODUCTOS", "3.5. OTROS", "Cemento", "- Los demás", None])
    p.append([4407999000, "4407999000", None, None, None, "- - Las demás maderas", None])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def sunat_agosto():
    wb = Workbook()
    ws = wb.active
    ws.append(CAB)
    for f in [fila("20513842377", "2709000000", "PETROLEO CRUDO", 17033000),             # ok -> petróleo
              fila("20100000001", "1801001900", "CACAO EN GRANO SECO", 290400),          # ok -> cacao (PROD_UNICOS)
              fila("20601551951", "1801001900", "CACAO TOSTADO", 5000),                   # BORRAR PARA: LOR SMN UCA
              fila("20100000003", "0901119000", "CAFE VERDE", 70000),                     # regla: café en Loreto
              fila("20999999999", "1801001900", "CACAO", 8000),                           # empresa nueva
              fila("20100000001", "4407999000", "MADERA TORNILLO", 9000),                 # partida sin clasificar
              fila("20100000001", "2523290000", "CEMENTO", 1000),                         # otras exportaciones
              fila("20100000001", "1801001900", "CACAO EN GRANO SECO", 7000, dep="15"),   # Lima: no se guarda
              [None] * 9 + [999999999]]:                                                  # fila de totales
        ws.append(f)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def test_flujo_sunat_completo(cliente):
    r = subir(cliente, **{"ESTADLO.xlsx": estadlo(), "ANUAL_PRO_2026.xlsx": anual_pro()})
    assert r["bases"]["expo-sunat"]["periodos"] == ["2026-07"]                 # la hoja ORIGEN como base
    assert any("Partidas (PARTIDA): 5, 4 clasificadas" in a for a in r["avisos"])
    r = subir(cliente, **{"X_202608.xlsx": sunat_agosto()})
    assert any("1 filas de otros departamentos" in a for a in r["avisos"])
    df = datos(cliente, "expo-sunat", "16")
    ago = df[df["MES"] == 8].set_index("DESCRIPCION")
    assert ago.loc["PETROLEO CRUDO", "ESTADO"] == "ok" and ago.loc["PETROLEO CRUDO", "RUBRO_ESTAD"] == "Petróleo crudo y derivados"
    assert ago.loc["CACAO EN GRANO SECO", "PRODUCTO"] == "Cacao en grano" and ago.loc["CACAO EN GRANO SECO", "RUBRO_ESTAD"] == "Cacao en grano"
    assert ago.loc["CACAO EN GRANO SECO", "FOB"] == 290 and ago.loc["CACAO EN GRANO SECO", "PAIS"] == "Brasil"
    assert ago.loc["CACAO TOSTADO", "ESTADO"] == "borrar"
    assert ago.loc["CAFE VERDE", "ESTADO"] == "borrar" and "café" in ago.loc["CAFE VERDE", "MOTIVO"]
    assert ago.loc["CACAO", "ESTADO"] == "verificar"
    assert ago.loc["MADERA TORNILLO", "ESTADO"] == "clasificar"
    assert ago.loc["CEMENTO", "RUBRO_ESTAD"] == "OTRAS EXPORTACIONES"

    x = cliente.get("/v1/exportaciones/sunat?departamento=16&periodo=2026-08").json()["revision"]
    assert [e["ruc"] for e in x["empresas"]] == ["20999999999"]
    assert [p["partida"] for p in x["partidas"]] == ["4407999000"]
    prop = {f["rubro"]: f for f in x["propuesta"]["filas"]}
    assert prop["Petróleo crudo y derivados"]["sunat"] == 17033 and prop["Petróleo crudo y derivados"]["dif"] == 0
    assert prop["TOTAL EXPORTACIONES"]["sunat"] == 17033 + 290 + 1
    assert x["destinos"]["mes"][0]["nombre"] == "Brasil"

    # revisión: la empresa nueva pasa a OK, la partida se clasifica como madera y va a «Otros acabados»
    assert cliente.post("/v1/exportaciones/sunat/empresas", json={"filas": [{"RUC": "20999999999", "BORRAR_PARA": "OK"}]}).status_code == 200
    assert cliente.post("/v1/exportaciones/sunat/partidas", json={"filas": [{"PARTIDA": "4407999000", "TIPO": "2. PRODUCTOS NO TRADICIONALES",
                                                                            "CATEGORIA": "2.1. MADERAS Y SUS MANUFACTURAS", "PRODUCTO": "Madera tornillo"}]}).status_code == 200
    ago = datos(cliente, "expo-sunat", "16").query("MES == 8").set_index("DESCRIPCION")
    assert ago.loc["CACAO", "ESTADO"] == "ok"
    assert ago.loc["MADERA TORNILLO", "RUBRO_ESTAD"] == "Otros acabados"
    # una equivalencia manda el cemento a «Resto»; una regla nueva borra el cacao de la empresa nueva
    cliente.post("/v1/exportaciones/sunat/equivalencias", json={"filas": [{"departamento": "LOR", "producto": "Cemento", "rubro": "Resto"}]})
    reglas = x["reglas"] + [{"departamento": "LOR", "campo": "ruc", "valor": "20999999999", "motivo": "prueba"}]
    cliente.post("/v1/exportaciones/sunat/reglas", json={"filas": reglas})
    ago = datos(cliente, "expo-sunat", "16").query("MES == 8").set_index("DESCRIPCION")
    assert ago.loc["CEMENTO", "RUBRO_ESTAD"] == "Resto" and ago.loc["CACAO", "ESTADO"] == "borrar"

    # pasar al ESTAD: los cuadros usan lo de SUNAT
    assert cliente.post("/v1/exportaciones/sunat-estad?departamento=16&periodo=2026-08").status_code == 200
    c = cliente.get("/v1/cuadros/exportaciones?departamento=16&periodo=2026-08").json()["cuadros"]
    s = {f["nombre"]: f for f in c["sintesis"]}
    assert s["Petróleo crudo y derivados"]["mes_a"] == 17.03 and s["Maderas y sus manufacturas"]["mes_a"] == 0.01
    r = cliente.get("/v1/exportaciones/sunat-estad.xlsx?departamento=16&periodo=2026-08")
    assert load_workbook(io.BytesIO(r.content)).sheetnames == ["ESTAD propuesto", "Detalle SUNAT"]


def test_padron_de_ruc(cliente):
    subir(cliente, **{"ESTADLO.xlsx": estadlo(), "ANUAL_PRO_2026.xlsx": anual_pro()})
    subir(cliente, **{"X_202608.xlsx": sunat_agosto()})
    assert "20999999999" in sunat_expo.pendientes_ruc()
    spec = importlib.util.spec_from_file_location("padron_ruc", Path(__file__).resolve().parents[3] / "EXPORT" / "automatizacion" / "padron_ruc.py")
    padron = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(padron)
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w") as z:
        z.writestr("padron_reducido_ruc.txt", "RUC|NOMBRE O RAZÓN SOCIAL|ESTADO|CONDICIÓN|\n20999999999|NUEVA EXPORTADORA SAC|ACTIVO|HABIDO|\n"
                                              "20111111111|OTRA SAC|ACTIVO|HABIDO|\n".encode("latin-1"))
    df = padron.buscar(["20999999999"], buf.getvalue())
    assert df.to_dict("records") == [{"RUC": "20999999999", "RAZON_SOCIAL": "NUEVA EXPORTADORA SAC", "ESTADO": "ACTIVO", "CONDICION": "HABIDO"}]
    sunat_expo.guardar_padron(df)
    sunat_expo.reprocesar()
    assert "20999999999" not in sunat_expo.pendientes_ruc()
    x = cliente.get("/v1/exportaciones/sunat?departamento=16&periodo=2026-08").json()["revision"]
    assert x["empresas"][0]["razon_social"] == "NUEVA EXPORTADORA SAC"          # sigue por verificar (BORRAR PARA vacío)
