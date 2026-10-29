"""Pruebas de la base mensual de importaciones (aduana de Iquitos): clasificación por CUODE, ESTAD propuesto y resumen."""
import io

from openpyxl import Workbook, load_workbook

from tests.test_exportaciones import estadlo
from tests.test_pesca import cliente, datos, subir  # noqa: F401

CAB = ["aduana", "ruc", "partida", "fob", "flete", "seguro", "kilos", "pais_adqui", "cuode", "periodo", "cuod1", "daduana"]


def anual_impo():
    wb = Workbook()
    o = wb.active
    o.title = "ORIGEN"
    for _ in range(4):
        o.append([])
    o.append(["ADUANA", "RUC", "PARTIDA", "FOB", "FLETE", "SEGURO", "KILOS", "COD_PAIS", "CUODE", "FENUM", "CUOD"])
    o.append(["226", "20393529076", "0303890000", 25050, 4500, 375.75, 15000, 105, 113, 202607, 1])
    p = wb.create_sheet("PARTIDAS")
    p.append(["PARTIDA", "PARTIDA_TEXTO", "TIPO", "CATEGORIA", "DESC. CUODE", "PRODUCTO", "NOTAS", "DESCRIPCION", "FUENTE", None,
              "CUODE PRIMARIO", "CUODE SECUNDARIO", "Cuode", "TIPO", "CATEGORÍA", "DESC. CUODE"])
    p.append(["0303890000", "0303890000", "BIENES DE CONSUMO", "Bienes de consumo no duradero", "x", "Pescado congelado", None, None, "ESTIMADO", None,
              "01.", 100, 113, "BIENES DE CONSUMO", "Bienes de consumo no duradero", "Productos alimenticios elaborados"])
    p.append(["1701140000", "1701140000", "BIENES DE CONSUMO", "Bienes de consumo no duradero", "x", "Azúcar de caña", None, None, "ESTIMADO", None,
              "02.", 400, 412, "MATERIAS PRIMAS", "Materias primas para la agricultura", "Alimentos para animales: semielaborados"])
    p.append(["8430410000", "8430410000", "BIENES DE CAPITAL", "Para la industria", "x", "Perforadoras", None, None, "MANUAL", None,
              "03.", 800, 840, "BIENES DE CAPITAL", "Para la industria", "Maquinaria industrial"])
    u = wb.create_sheet("UNICOS")
    u.append(["CÓDIGO", "TEXTO", "PAÍS", "CONTINENTE", None, None, None, "RUC", "RAZÓN SOCIAL", None, "PRODUCTO_ORIGINAL", "PRODUCTO_CORREGIDO"])
    u.append([105, "105", "Brasil", "AMÉRICA", None, None, None, "20393529076", "PESCADOS SAC", None, "Perforadoras", "Perforadoras de pozos"])
    u.append([169, "169", "Colombia", "AMÉRICA"])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def base_agosto():
    wb = Workbook()
    ws = wb.active
    ws.append(CAB)
    for f in [["226", "20393529076", "0303890000", 25050, 4500, 375.75, 15000, 105, 113, 202608, 1, "iquitos"],   # CUODE 113 -> no duraderos
              ["226", "20999999999", "1701140000", 1000400, 0, 0, 1, 105, 512, 202608, 2, "iquitos"],          # CUODE 512 no está -> partida
              ["226", "20393529076", "8430410000", 6772684, 0, 0, 1, 169, 412, 202608, 3, "iquitos"],          # MANUAL gana al CUODE
              ["226", "20393529076", "9999999999", 500, 0, 0, 1, 105, None, 202608, 3, "iquitos"],             # sin CUODE ni partida
              ["235", "20393529076", "0303890000", 9000, 0, 0, 1, 105, 113, 202608, 1, "aerea"]]:               # otra aduana
        ws.append(f)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def test_importaciones_base_mensual(cliente):
    r = subir(cliente, **{"ESTADLO.xlsx": estadlo(), "ANUAL_PRO_-_IMPORTACIONES.xlsm": anual_impo()})
    assert r["bases"]["impo-base"]["periodos"] == ["2026-07"]
    subir(cliente, **{"M_202608.xlsx": base_agosto()})
    df = datos(cliente, "impo-base", "16").query("MES == 8").set_index("PARTIDA")
    assert df.loc["0303890000"].query("ADUANA == '226'")["RUBRO_ESTAD"].iloc[0] == "Bienes no duraderos (100)"
    assert df.loc["0303890000"].query("ADUANA == '226'")["CIF"].iloc[0] == 25050 + 4500 + 375.75
    assert df.loc["1701140000", "CATEGORIA"] == "Bienes de consumo no duradero" and df.loc["1701140000", "EMPRESA"] == "SIN RAZÓN SOCIAL"
    assert df.loc["8430410000", "CATEGORIA"] == "Para la industria" and df.loc["8430410000", "PRODUCTO"] == "Perforadoras de pozos"
    assert df.loc["8430410000", "PAIS"] == "Colombia"
    assert df.loc["9999999999", "ESTADO"] == "clasificar"
    assert df.loc["0303890000"].query("ADUANA == '235'")["ESTADO"].iloc[0] == "otra aduana"
    x = cliente.get("/v1/importaciones/base?periodo=2026-08").json()["revision"]
    assert [e["ruc"] for e in x["empresas"]] == ["20999999999"] and [p["partida"] for p in x["partidas"]] == ["9999999999"]
    prop = {f["rubro"]: f["sunat"] for f in x["propuesta"]["filas"]}
    assert prop["Bienes no duraderos (100)"] == 25 + 1000 and prop["Bienes para la industria (800)"] == 6773 and prop["TOTAL IMPORTACIONES"] == 7798
    # correcciones: razón social y clasificación manual de la partida que faltaba
    cliente.post("/v1/importaciones/base/empresas", json={"filas": [{"RUC": "20999999999", "RAZON_SOCIAL": "AZUCARERA SAC"}]})
    cliente.post("/v1/importaciones/base/partidas", json={"filas": [{"PARTIDA": "9999999999", "CATEGORIA": "Para la industria", "PRODUCTO": "Repuestos"}]})
    df = datos(cliente, "impo-base", "16").query("MES == 8").set_index("PARTIDA")
    assert df.loc["1701140000", "EMPRESA"] == "AZUCARERA SAC" and df.loc["9999999999", "ESTADO"] == "ok"
    assert df.loc["9999999999", "RUBRO_ESTAD"] == "Bienes para la industria (800)"
    assert cliente.post("/v1/importaciones/estad?periodo=2026-08").status_code == 200
    c = cliente.get("/v1/cuadros/exportaciones?departamento=16&flujo=impo&periodo=2026-08").json()["cuadros"]
    assert c["sintesis"][-1]["mes_a"] == 7.8
    r = cliente.get("/v1/importaciones/estad.xlsx?periodo=2026-08")
    assert load_workbook(io.BytesIO(r.content)).sheetnames == ["ESTAD propuesto", "Detalle"]
    res = cliente.get("/v1/importaciones/resumen?periodo=2026-08").json()["resumen"]
    assert res["textos"][0].startswith("Las importaciones por la aduana de Iquitos fueron de USD 7,8 millones en agosto de 2026")
    assert "bienes de capital para la industria (perforadoras de pozos" in res["textos"][0]
    assert res["destacados"][0]["valor"] == "Perforadoras de pozos"


def test_importaciones_nombres_y_diccionarios(cliente):
    subir(cliente, **{"ESTADLO.xlsx": estadlo(), "ANUAL_PRO_-_IMPORTACIONES.xlsm": anual_impo()})
    subir(cliente, **{"M_202608.xlsx": base_agosto()})
    x = cliente.get("/v1/importaciones/base?periodo=2026-08").json()["revision"]
    assert {n["original"] for n in x["nombres"]} == {"Pescado congelado", "Azúcar de caña"}     # Perforadoras sí está en PROD_UNICOS
    cliente.post("/v1/importaciones/base/nombres", json={"filas": [{"original": "Azúcar de caña", "nombre": "Azúcar"}]})
    df = datos(cliente, "impo-base", "16").query("MES == 8").set_index("PARTIDA")
    assert df.loc["1701140000", "PRODUCTO"] == "Azúcar"
    r = cliente.get("/v1/importaciones/diccionarios.xlsx")
    assert {"LEEME", "PARTIDAS", "CUODE", "PROD_UNICOS"} <= set(load_workbook(io.BytesIO(r.content)).sheetnames)
    r = subir(cliente, **{"dic_impo.xlsx": r.content})
    assert any("Importaciones: base mensual reprocesada" in a for a in r["avisos"])
