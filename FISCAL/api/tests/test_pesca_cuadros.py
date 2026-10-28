"""Cuadros de pesca: producción del mes y acumulada contra otro año, series del VBP y Cuadro Nº 07 de Loreto."""
import io

import pandas as pd
from openpyxl import load_workbook

from tests.test_hidrocarburos import cliente  # noqa: F401
from app.pesca import LOR_COLUMNAS, SMN_COLUMNAS, UCA_COLUMNAS

MESES = [(2025, m) for m in range(1, 13)] + [(2026, m) for m in range(1, 4)]


def long(hoja, columnas, filas):
    buf = io.BytesIO()
    with pd.ExcelWriter(buf) as w:
        pd.DataFrame(filas, columns=columnas).to_excel(w, sheet_name=hoja, index=False)
    return buf.getvalue()


def lor():
    f = []
    for a, m in MESES:
        k = 1.2 if a == 2026 else 1.0
        f += [dict(AÑO=a, MES=m, DEPARTAMENTO="LORETO", FUENTE="DESEMBARQUE", ESTADO="FRESCO", DESTINO="", ESPECIE="Boquichico",
                   ESPECIE_CORREGIDA="Boquichico", ESPECIE_OK="Boquichico", ESPECIE_CUADRO="Boquichico", UNIDAD="t", VOLUMEN=100 * k, VBP=327.504 * k),
              dict(AÑO=a, MES=m, DEPARTAMENTO="LORETO", FUENTE="DESEMBARQUE", ESTADO="SECO SALADO", DESTINO="", ESPECIE="Paiche",
                   ESPECIE_CORREGIDA="Paiche", ESPECIE_OK="Paiche", ESPECIE_CUADRO="Otros", UNIDAD="t", VOLUMEN=10, VBP=29.8778),
              dict(AÑO=a, MES=m, DEPARTAMENTO="LORETO", FUENTE="POR ESPECIES", ESTADO="ORNAMENTAL", DESTINO="Mercado externo", ESPECIE="Otocinclo",
                   ESPECIE_CORREGIDA="Otocinclo", ESPECIE_OK="Otocinclo", ESPECIE_CUADRO="Otros", UNIDAD="miles de unidades", VOLUMEN=50, VBP=51.645)]
    return long("LOR_PESCA", LOR_COLUMNAS, f)


def smn():
    f = []
    for a, m in MESES:
        f += [dict(AÑO=a, MES=m, DEPARTAMENTO="SAN MARTÍN", ORIGEN="Acuicultura", ESTADO="FRESCO", ESTADO_OK="FRESCO", ESPECIE="TILAPIA",
                   ESPECIE_OK="Tilapia", VOLUMEN=20 if a == 2026 else 10, VBP=0),
              dict(AÑO=a, MES=m, DEPARTAMENTO="Otros", ORIGEN="Extracción de ríos", ESTADO="FRESCO", ESTADO_OK="FRESCO", ESPECIE="UCAYALI",
                   ESPECIE_OK="Otras", VOLUMEN=5, VBP=0)]
    return long("SMN_PESCA", SMN_COLUMNAS, f)


def uca():
    f = []
    for a, m in MESES:
        f += [dict(AÑO=a, MES=m, TIPO="FRESCO", ESPECIES="BOQUICHICO", ESPECIES_CORREGIDO="Boquichico", ESPECIES_OK="Boquichico", VOLUMEN=30, VBP=98.25),
              dict(AÑO=a, MES=m, TIPO="AGREGADO", ESPECIES="TOTAL", ESPECIES_CORREGIDO="Total", ESPECIES_OK="Otros", VOLUMEN=30, VBP=0),
              dict(AÑO=a, MES=m, TIPO="ORNAMENTAL", ESPECIES="NEON", ESPECIES_CORREGIDO="Neon", ESPECIES_OK="Neon", VOLUMEN=1000, VBP=0)]
    return long("UCA_PESCA", UCA_COLUMNAS, f)


def test_cuadros_de_las_tres_regiones(cliente):  # noqa: F811
    r = cliente.post("/v1/cargas", files=[("archivos", (n, b, "application/octet-stream")) for n, b in
                                          (("LOR_PESCA_LONG.xlsx", lor()), ("SMN_PESCA_LONG.xlsx", smn()), ("UCA_PESCA_LONG.xlsx", uca()))])
    assert r.status_code == 200, r.text
    c = cliente.get("/v1/cuadros/pesca?departamento=16").json()["cuadros"]
    assert c["periodo"] == "2026-03" and c["comparar"] == 2025
    f = {x["nombre"]: x for x in c["filas"]}
    assert f["Estado fresco"]["var_mes"] == 20.0 and f["Estado fresco"]["acum_a"] == 360.0
    assert f["Peces ornamentales vivos"]["mes_a"] == 50.0 and f["Desembarque total"]["mes_a"] == 130.0
    assert c["mensual"]["b"][11] is not None and c["mensual"]["a"][3] is None
    sl = {x["nombre"]: x for x in c["sintesis"]["filas"]}
    assert c["sintesis"]["titulo"] == "LORETO: PRODUCCIÓN PESQUERA 1/" and sl["Consumo Humano (toneladas)"]["mes_a"] == 130.0
    assert sl["Boquichico"]["mes_a"] == 120.0 and sl["Mercado externo"]["mes_a"] == 50.0 and sl["SECTOR PESCA 2/"]["solo_var"]
    c7 = {}
    for x in c["cuadro_07"]["filas"]:                   # el bloque fresco repite nombres: vale la primera aparición
        c7.setdefault(x["nombre"], x)
    assert c["cuadro_07"]["columnas"][0] == "Ene. 2025" and c["cuadro_07"]["columnas"][-1] == "Mar. 2026"
    assert c7["Paiche"]["totales"]["2025"] == 120.0 and c7["PECES ORNAMENTALES 1/"]["totales"]["2026"] == 150.0
    s = cliente.get("/v1/cuadros/pesca?departamento=22").json()["cuadros"]
    fs = {x["nombre"]: x for x in s["filas"]}
    assert fs["Producción de San Martín"]["var_mes"] == 100.0 and fs["Llega de otros departamentos"]["mes_a"] == 5.0 and "Tilapia" in fs
    ss = {x["nombre"]: x for x in s["sintesis"]["filas"]}
    assert ss["Producción Interna"]["mes_a"] == 20.0 and ss["Tilapia"]["var_acum"] == 100.0 and ss["Otras especies"]["mes_a"] == 0.0
    c7s = {x["nombre"]: x for x in s["cuadro_07"]["filas"]}
    assert c7s["TOTAL"]["totales"]["2026"] == 75.0 and c7s["INGRESO OTROS DPTOS. (CONTINENTAL)"]["totales"]["2026"] == 15.0
    u = cliente.get("/v1/cuadros/pesca?departamento=25&periodo=2026-02&formato=xlsx")
    assert load_workbook(io.BytesIO(u.content)).sheetnames == ["Producción pesquera", "Cuadro 07", "Producción", "VBP mensual"]
    assert cliente.get("/v1/cuadros/pesca?departamento=16&comparar=2026").status_code == 422


def test_ucayali_formato_de_la_sucursal(cliente):  # noqa: F811
    """«UCAYALI: PRODUCCIÓN PESQUERA» de junio 2026 (la imagen que mandó el usuario, con cifras redondeadas) y Cuadro Nº 07."""
    fresco = {"Bagre": (14, 90), "Ractacara": (0, 186), "Boquichico": (19, 2), "Palometa": (3, 48), "Chiu Chiu": (0, 0),
              "Sardina": (12, 30), "Doncella": (8, 139)}                            # Doncella va a «Otras especies»
    f = []
    for (a, i) in ((2025, 0), (2026, 1)):
        for m in range(1, 7):
            for e, v in fresco.items():
                vol = v[i] if m == 6 else 1
                f += [dict(AÑO=a, MES=m, TIPO="FRESCO", ESPECIES=e.upper(), ESPECIES_CORREGIDO=e, ESPECIES_OK=e if e != "Doncella" else "Otros",
                           VOLUMEN=vol, VBP=vol * 3.27504),
                      dict(AÑO=a, MES=m, TIPO="AGREGADO", ESPECIES=e.upper(), ESPECIES_CORREGIDO=e, ESPECIES_OK="Otros", VOLUMEN=vol, VBP=0)]
            seco = (16, 28)[i] if m == 6 else 1
            f += [dict(AÑO=a, MES=m, TIPO="SECO-SALADO", ESPECIES="BAGRE", ESPECIES_CORREGIDO="Bagre", ESPECIES_OK="Bagre", VOLUMEN=seco, VBP=seco * 2.98778),
                  dict(AÑO=a, MES=m, TIPO="AGREGADO", ESPECIES="BAGRE ", ESPECIES_CORREGIDO="Bagre", ESPECIES_OK="Otros", VOLUMEN=seco, VBP=0),
                  dict(AÑO=a, MES=m, TIPO="ORNAMENTAL", ESPECIES="NEON", ESPECIES_CORREGIDO="Neon", ESPECIES_OK="Neon", VOLUMEN=1000, VBP=0)]
    cliente.post("/v1/cargas", files=[("archivos", ("UCA_PESCA_LONG.xlsx", long("UCA_PESCA", UCA_COLUMNAS, f), "application/octet-stream"))])
    c = cliente.get("/v1/cuadros/pesca?departamento=25&periodo=2026-06").json()["cuadros"]
    s = {x["nombre"]: x for x in c["sintesis"]["filas"]}
    assert [x["nombre"] for x in c["sintesis"]["filas"]] == ["Consumo Humano", "Fresco", "Bagre", "Ractacara", "Boquichico", "Palometa", "Chiu Chiu",
                                                              "Sardina", "Otras especies", "Seco salado", "SECTOR PESCA 2/"]
    assert (s["Consumo Humano"]["mes_b"], s["Consumo Humano"]["mes_a"]) == (72.0, 523.0)
    assert s["Fresco"]["mes_a"] == 495.0 and s["Otras especies"]["mes_a"] == 139.0 and s["Boquichico"]["var_mes"] == -89.5
    assert s["Ractacara"]["var_mes"] is None and s["SECTOR PESCA 2/"]["solo_var"] and s["Seco salado"]["var_mes"] == 75.0
    c7 = c["cuadro_07"]
    assert c7["columnas"][0] == "Ene. 2025" and len(c7["columnas"]) == 12
    filas = {}
    for x in c7["filas"]:
        if not x.get("titulo"):
            filas.setdefault(x["nombre"], x)
    assert filas["TOTAL 1/"]["totales"]["2026"] == filas["TOTAL"]["totales"]["2026"] == 523 + 5 * 8      # 1/ = fresco + seco
    assert filas["PECES ORNAMENTALES (unidades)"]["totales"]["2025"] == 6000
    wb = load_workbook(io.BytesIO(cliente.get("/v1/cuadros/pesca?departamento=25&periodo=2026-06&formato=xlsx").content))
    ws = wb["Producción pesquera"]
    assert ws["B7"].value == "UCAYALI: PRODUCCIÓN PESQUERA 1/" and [ws.cell(13, j).value for j in range(2, 5)] == ["Consumo Humano", 72, 523]
    w7 = wb["Cuadro 07"]
    assert w7["A2"].value == "UCAYALI: DESEMBARQUE DE PESCADO PARA CONSUMO HUMANO" and w7["B4"].value == 2025 and w7["H5"].value == 2025      # 2025 trae ene.-jun. en la prueba
