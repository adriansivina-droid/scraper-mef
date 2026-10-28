"""Empleo: carga del Excel del scraper y cuadro de la síntesis, comparado con lo publicado (Loreto, Cuadro N° 6, junio 2026)."""
import io
import os

import pandas as pd
import pytest
from openpyxl import load_workbook
from sqlalchemy import text

from app import db
from tests.test_api import cliente_http

# Cuadro N° 6 de la síntesis de Loreto (junio 2026): abril 2025, abril 2026 y diciembre 2025
LORETO = {"Agrícola": (3199, 3373, 3829), "Comercio": (10295, 10469, 10338), "Construcción": (3085, 2947, 2996),
          "Manufactura": (2495, 2669, 2696), "Minería": (233, 304, 307), "No precisa": (1461, 1339, 1284),
          "Servicios": (17506, 18577, 18187)}
PUBLICADO = {"Agrícola": (5.4, -11.9), "Comercio": (1.7, 1.3), "Construcción": (-4.5, -1.6), "Manufactura": (7.0, -1.0),
             "Minería": (30.5, -1.0), "Otros": (-8.4, 4.3), "Servicios": (6.1, 2.1), "Total": (3.7, 0.1)}


def excel_scraper():
    filas = []
    for act, (abr25, abr26, dic25) in LORETO.items():
        for (a, m), v in {(2025, 4): abr25, (2026, 4): abr26, (2025, 12): dic25}.items():
            filas.append(dict(AÑO=a, MES=m, REGIÓN="Loreto", ACTIVIDAD_ECONOMICA=act, EMPLEO=v))
        filas.append(dict(AÑO=2026, MES=4, REGIÓN="Perú", ACTIVIDAD_ECONOMICA=act, EMPLEO=1))
    for region in ("Lima Metropolitana", "Lima Provincias"):
        filas.append(dict(AÑO=2026, MES=4, REGIÓN=region, ACTIVIDAD_ECONOMICA="Comercio", EMPLEO=100))
    buf = io.BytesIO()
    pd.DataFrame(filas).to_excel(buf, index=False)
    return buf.getvalue()


@pytest.fixture()
def cliente(tmp_path, monkeypatch):
    monkeypatch.setenv("DATABASE_URL", os.environ.get("PRUEBA_DATABASE_URL") or f"sqlite:///{tmp_path / 'empleo.db'}")
    monkeypatch.delenv("APP_CLAVE", raising=False)
    db.reiniciar_motor()
    if db.motor().dialect.name == "postgresql":
        with db.motor().begin() as cx:
            cx.execute(text("DROP SCHEMA public CASCADE; CREATE SCHEMA public"))
        db.reiniciar_motor()
    c = cliente_http()
    r = c.post("/v1/cargas", files=[("archivos", ("Loreto_empleo.xlsx", excel_scraper(), "application/octet-stream"))])
    assert r.status_code == 200, r.text
    yield c
    db.reiniciar_motor()


def test_carga(cliente):
    e = cliente.get("/v1/estado").json()["bases"]["empleo"]
    assert e["departamentos"] == ["15", "16"] and e["desde"] == "2025-04" and e["hasta"] == "2026-04"
    avisos = cliente.get("/v1/estado").json()["ultima_carga"]["resumen"]["avisos"]
    assert any("Perú" in a for a in avisos)                      # el total nacional no se guarda


def test_lima_metropolitana_y_provincias_no_se_pisan(cliente):
    d = cliente.get("/v1/descargas?base=empleo&departamento=15&desde=2026-01&hasta=2026-12&formato=json").json()["empleo"]["datos"]
    assert sorted(f["REGIÓN"] for f in d) == ["Lima Metropolitana", "Lima Provincias"]


def test_cuadro_igual_al_publicado(cliente):
    c = cliente.get("/v1/cuadros/empleo?departamento=16&periodo=2026-04").json()["cuadros"]
    assert c["periodo"] == "2026-04" and [f["rama"] for f in c["cuadro"]] == list(PUBLICADO)
    for f in c["cuadro"]:
        assert (f["var_mes"], f["var_dic"]) == PUBLICADO[f["rama"]], f["rama"]
    total = c["cuadro"][-1]
    assert (total["mes_anterior"], total["mes"], total["diciembre"]) == (38274, 39678, 39637)
    assert c["grafico"][-1] == {"periodo": "Abr. 2026", "var": 3.7}
    assert c["anexo"][-1]["Agricultura"] == 5.4 and c["anexo"][-1]["Total"] == 3.7


def test_cuadro_por_defecto_y_errores(cliente):
    assert cliente.get("/v1/cuadros/empleo?departamento=16").json()["cuadros"]["periodo"] == "2026-04"
    assert cliente.get("/v1/cuadros/empleo?departamento=16&periodo=2026-09").status_code == 422
    assert cliente.get("/v1/cuadros/empleo?departamento=25").json()["cuadros"] is None


def test_cuadro_en_excel(cliente):
    r = cliente.get("/v1/cuadros/empleo?departamento=16&periodo=2026-04&formato=xlsx")
    wb = load_workbook(io.BytesIO(r.content))
    assert wb.sheetnames == ["Cuadro", "Gráfico", "Anexo"]
    ws = wb["Cuadro"]
    assert ws["A1"].value == "LORETO: TRABAJADORES ASALARIADOS PRIVADOS FORMALES"
    assert [c.value for c in ws[4]] == ["Rama de actividad", "Abril 2025", "Abril 2026", "Var. %", "Dic. 2025", "Abr. 2026", "Var. %"]
    assert [c.value for c in ws[12]][:4] == ["Total", 38274, 39678, 3.7]
