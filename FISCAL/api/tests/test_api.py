"""Pruebas de la fase 1 con archivos que imitan los nombres y columnas de Fiscal_Scraper (sin MEF ni Railway)."""
import io
import os
import zipfile

import pandas as pd
import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook
from sqlalchemy import text

from app import db

DEPTOS = {"16": "16: LORETO", "22": "22: SAN MARTÍN"}
NIVELES = {"GobNacional": "E: GOBIERNO NACIONAL", "GobRegional": "R: GOBIERNOS REGIONALES", "Municipalidad": "M: GOBIERNOS LOCALES"}
MESES = [(2025, 9), (2026, 8), (2026, 9)]


def _csv(df):
    return df.to_csv(index=False).encode("utf-8-sig")


def archivos_scraper():
    """{ruta dentro del zip: bytes}, con la estructura de Fiscal/Scrapeo/output."""
    salida = {}
    for corto, etq in NIVELES.items():
        carpeta = {"GobNacional": "Gobierno Nacional", "GobRegional": "Gobierno Regional", "Municipalidad": "Gobierno Municipal"}[corto]
        for a, m in MESES:
            base = dict(AÑO=a, MES=f"{m}: Setiembre", NIVEL_GOBIERNO=etq)
            fun = []
            for d in DEPTOS.values():
                for cod, f in [("15", "15: TRANSPORTE"), ("22", "22: EDUCACION")]:
                    if corto == "Municipalidad":
                        for mu in ["160101-301263: MUNICIPALIDAD PROVINCIAL DE MAYNAS", "160108-301270: MUNICIPALIDAD DISTRITAL DE PUNCHANA"]:
                            fun.append(dict(AÑO=a, MES=base["MES"], DEPARTAMENTO=d, FUNCION=f, MUNICIPALIDAD=mu, DEVENGADO=100.0 * m))
                    else:
                        fun.append(dict(base, DEPARTAMENTO=d, ENTIDAD="X", COD_FUNCION=cod, FUNCION=f, DEVENGADO=1000.5 * m))
            sufijo = f"{a} - Mes {m:02d} - LOR - SMN - UCA.csv"
            salida[f"output/DEVENGADO/{carpeta}/DEV Funcion.{corto} {sufijo}"] = _csv(pd.DataFrame(fun))
            gen, deu, fue = [], [], []
            for d in DEPTOS.values():
                gen += [dict(base, DEPARTAMENTO=d, COD_GENERICA="5-21", GENERICA="5-21: PERSONAL Y OBLIGACIONES SOCIALES", DEVENGADO=500.0),
                        dict(base, DEPARTAMENTO=d, COD_GENERICA="7-28", GENERICA="7-28: SERVICIO DE LA DEUDA PUBLICA", DEVENGADO=90.0)]
                deu += [dict(base, DEPARTAMENTO=d, GENERICA="7-2-8", RUTA="TOTAL GENÉRICA 2.8", CLASIFICACION="TOTAL", DEVENGADO=90.0),
                        dict(base, DEPARTAMENTO=d, GENERICA="7-2-8", RUTA="2: AMORTIZACION DE LA DEUDA", CLASIFICACION="AMORTIZACION", DEVENGADO=60.0),
                        dict(base, DEPARTAMENTO=d, GENERICA="7-2-8", RUTA="1: INTERESES", CLASIFICACION="OTRO", DEVENGADO=25.0)]
                fue += [dict(base, DEPARTAMENTO=d, GENERICA="6-2-6", COD_FUENTE="1", FUENTE="1: RECURSOS ORDINARIOS", DEVENGADO=70.0)]
            salida[f"output/DEVENGADO/{carpeta}/DEV Generica.{corto} {sufijo}"] = _csv(pd.DataFrame(gen))
            if corto != "GobNacional":             # GN sin desglose de deuda: la 2.8 queda en SIN DESGLOSE
                salida[f"output/DEVENGADO/{carpeta}/DEV Deuda.{corto} {sufijo}"] = _csv(pd.DataFrame(deu))
            salida[f"output/DEVENGADO/{carpeta}/DEV InvFuente.{corto} {sufijo}"] = _csv(pd.DataFrame(fue))
            salida[f"output/DEVENGADO/{carpeta}/DEV Fuente.{corto} {sufijo}"] = _csv(pd.DataFrame(fue).assign(GENERICA="TODAS"))
    pim = [dict(AÑO=a, DEPARTAMENTO=d, **{"NIVEL DE GOBIERNO": n}, FUNCION=f, PIM=1_000_000)
           for a in (2025, 2026) for d in DEPTOS.values() for n in ["GOBIERNO NACIONAL", "GOBIERNOS LOCALES", "GOBIERNOS REGIONALES"]
           for f in ["15: TRANSPORTE", "22: EDUCACION"]]
    salida["output/PIM/PIM X FUNCIÓN/DATA/PIM Funcion 2026 - LOR - SMN - UCA.csv"] = _csv(pd.DataFrame(pim))
    salida["output/PIM/PIM X FUNCIÓN X MUNICIPALIDAD/DATA/PIM Funcion.Municipalidad 2026 - LOR - SMN - UCA.csv"] = _csv(pd.DataFrame(
        [dict(AÑO=2026, DEPARTAMENTO="16: LORETO", FUNCION="15: TRANSPORTE", MUNICIPALIDAD="160101-301263: MUNICIPALIDAD PROVINCIAL DE MAYNAS", PIM=5)]))
    salida["output/DEVENGADO/CARPETA DE CONTROL/CONTROL GobRegional 2026.csv"] = b"ESTADO\nOK\n"
    return salida


def zip_de(archivos):
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w") as z:
        for ruta, datos in archivos.items():
            z.writestr(ruta, datos)
    return buf.getvalue()


def cliente_http(auth=None):
    from app.main import app
    c = TestClient(app)
    c.auth = auth
    return c


@pytest.fixture()
def cliente(tmp_path, monkeypatch):
    # PRUEBA_DATABASE_URL=postgresql://… corre las mismas pruebas contra Postgres (como en Railway)
    monkeypatch.setenv("DATABASE_URL", os.environ.get("PRUEBA_DATABASE_URL") or f"sqlite:///{tmp_path / 'prueba.db'}")
    monkeypatch.setenv("APP_USUARIO", "equipo")
    monkeypatch.setenv("APP_CLAVE", "secreta")
    db.reiniciar_motor()
    if db.motor().dialect.name == "postgresql":
        with db.motor().begin() as cx:
            cx.execute(text("DROP SCHEMA public CASCADE; CREATE SCHEMA public"))
        db.reiniciar_motor()
    c = cliente_http(("equipo", "secreta"))
    r = c.post("/v1/cargas", files=[("archivos", ("output.zip", zip_de(archivos_scraper()), "application/zip"))])
    assert r.status_code == 200, r.text
    yield c
    db.reiniciar_motor()


Q = "departamento=16,22&nivel=nacional,regional,locales&desde=2025-01&hasta=2026-09"


def test_pide_clave(cliente):
    assert cliente_http().get("/v1/estado").status_code == 401
    assert cliente_http(("equipo", "mala")).get("/").status_code == 401
    assert cliente_http().get("/salud").status_code == 200


def test_carga_y_estado(cliente):
    e = cliente.get("/v1/estado").json()
    assert e["bases"]["inv-dev-funcion"] == {"filas": 3 * 2 * 2 * 3, "departamentos": ["16", "22"], "desde": "2025-09", "hasta": "2026-09"}
    assert e["bases"]["inv-pim-funcion"]["filas"] == 24
    assert e["bases"]["inv-dev-municipalidad"]["filas"] == 3 * 2 * 2 * 2
    assert "CONTROL" not in str(e["ultima_carga"]["resumen"]["avisos"])


def test_recargar_no_duplica(cliente):
    antes = cliente.get("/v1/estado").json()["bases"]
    cliente.post("/v1/cargas", files=[("archivos", ("output.zip", zip_de(archivos_scraper()), "application/zip"))])
    assert cliente.get("/v1/estado").json()["bases"] == antes


def test_locales_es_suma_de_municipalidades(cliente):
    v = cliente.get(f"/v1/descargas?base=inv-dev-funcion&formato=json&{Q}").json()["inv-dev-funcion"]["datos"]
    loc = [f for f in v if f["NIVEL_GOBIERNO"] == "M: GOBIERNOS LOCALES" and f["AÑO"] == 2026 and f["MES"].startswith("9:")
           and f["DEPARTAMENTO"] == "16: LORETO" and f["FUNCION"] == "15: TRANSPORTE"]
    assert [f["DEVENGADO"] for f in loc] == [1800.0] and loc[0]["ENTIDAD"] == "MUNICIPALIDADES (SUMA)"


def test_deuda_desglosada_y_total_intacto(cliente):
    v = cliente.get(f"/v1/descargas?base=gasto-dev-generica&formato=json&{Q}").json()["gasto-dev-generica"]["datos"]
    clave = lambda f: (f["AÑO"], f["MES"], f["DEPARTAMENTO"], f["NIVEL_GOBIERNO"])
    reg = [f for f in v if clave(f) == (2026, "9: Setiembre", "16: LORETO", "R: GOBIERNOS REGIONALES")]
    assert sorted((f["DESGLOSE"], f["DEVENGADO"]) for f in reg if f["COD_GENERICA"] == "7-28") == [("AMORTIZACIÓN", 60.0), ("INTERESES Y COMISIONES", 25.0), ("OTROS (sin clasificar)", 5.0)]
    assert sum(f["DEVENGADO"] for f in reg) == 590.0
    nac = [f for f in v if clave(f) == (2026, "9: Setiembre", "16: LORETO", "E: GOBIERNO NACIONAL") and f["COD_GENERICA"] == "7-28"]
    assert [(f["DESGLOSE"], f["DEVENGADO"]) for f in nac] == [("SIN DESGLOSE", 90.0)]


def test_filtros(cliente):
    v = cliente.get("/v1/descargas/vista-previa?base=inv-dev-funcion,inv-pim-funcion&departamento=16&nivel=regional&desde=2026-09&hasta=2026-09").json()
    filas = {a["base"]: a["filas"] for a in v["archivos"]}
    assert filas == {"inv-dev-funcion": 2, "inv-pim-funcion": 2}
    v = cliente.get("/v1/descargas/vista-previa?base=inv-dev-funcion&departamento=16,15&nivel=regional&desde=2026-01&hasta=2026-09").json()
    assert v["avisos"] == ["Inversión por función: sin datos para Lima."]
    assert cliente.get("/v1/descargas/vista-previa?base=inv-dev-funcion&departamento=&nivel=regional&desde=2026-01&hasta=2026-09").status_code == 422
    assert cliente.get("/v1/descargas/vista-previa?base=inv-dev-funcion&departamento=16&nivel=regional&desde=2026-09&hasta=2026-01").status_code == 422
    assert cliente.get("/v1/descargas/vista-previa?base=otra&departamento=16&nivel=regional&desde=2026-01&hasta=2026-09").status_code == 422


def test_excel_con_formato_de_pim_funcion(cliente):
    r = cliente.get(f"/v1/descargas?base=gasto-dev-fuente,inv-pim-funcion,inv-dev-funcion&formato=xlsx&{Q}")
    assert r.status_code == 200 and "attachment" in r.headers["content-disposition"]
    wb = load_workbook(io.BytesIO(r.content))
    assert wb.sheetnames == ["INV Funciones", "PIM Funciones", "GASTO Fuentes"]       # orden del catálogo
    ws = wb["PIM Funciones"]
    assert [c.value for c in ws[1]] == ["AÑO", "DEPARTAMENTO", "NIVEL DE GOBIERNO", "FUNCION", "PIM"]
    assert ws["A1"].alignment.horizontal == "center" and ws["E2"].number_format == "#,##0"
    assert ws.column_dimensions["B"].width == 41.6
    filas = [tuple(c.value for c in r) for r in ws.iter_rows(min_row=2)]
    assert filas[0][:3] == (2026, "16: LORETO", "GOBIERNO NACIONAL") and filas[-1][:3] == (2025, "22: SAN MARTÍN", "GOBIERNOS REGIONALES")


def test_csv_y_zip(cliente):
    r = cliente.get(f"/v1/descargas?base=inv-dev-fuente&formato=csv&{Q}")
    assert r.headers["content-type"].startswith("text/csv") and r.content.decode("utf-8-sig").startswith("AÑO,MES,NIVEL_GOBIERNO")
    r = cliente.get(f"/v1/descargas?base=inv-dev-fuente,gasto-dev-fuente&formato=csv&{Q}")
    assert sorted(zipfile.ZipFile(io.BytesIO(r.content)).namelist()) == [
        "GASTO Fuentes 2025-01_2026-09 - 16-22.csv", "INV Fuentes 2025-01_2026-09 - 16-22.csv"]


def test_pim_funcion_xlsx(cliente, tmp_path):
    df = pd.DataFrame([dict(AÑO=2020, DEPARTAMENTO="01: AMAZONAS", **{"NIVEL DE GOBIERNO": "GOBIERNO NACIONAL"}, FUNCION="15: TRANSPORTE", PIM=7)])
    buf = io.BytesIO(); df.to_excel(buf, index=False, sheet_name="PIM Funciones")
    r = cliente.post("/v1/cargas", files=[("archivos", ("PIM_Funcion.xlsx", buf.getvalue(), "application/octet-stream"))])
    assert r.json()["bases"]["inv-pim-funcion"]["departamentos"] == ["01"]
    e = cliente.get("/v1/estado").json()["bases"]["inv-pim-funcion"]
    assert e["desde"] == "2020" and "01" in e["departamentos"]


def test_zip_con_rutas_peligrosas_se_ignora(cliente):
    r = cliente.post("/v1/cargas", files=[("archivos", ("x.zip", zip_de({"../../fuera.csv": b"a,b\n1,2\n"}), "application/zip"))])
    assert r.status_code == 200 and r.json()["bases"] == {}


def test_mismo_periodo_en_dos_archivos_no_duplica(cliente):
    """El zip del scraper trae el PIM 2026 de Loreto y San Martín; PIM_Funcion.xlsx también: debe quedar una sola versión."""
    df = pd.DataFrame([dict(AÑO=2026, DEPARTAMENTO="16: LORETO", **{"NIVEL DE GOBIERNO": "GOBIERNO NACIONAL"}, FUNCION="15: TRANSPORTE", PIM=7)])
    buf = io.BytesIO(); df.to_excel(buf, index=False)
    arch = archivos_scraper() | {"PIM_Funcion.xlsx": buf.getvalue()}
    r = cliente.post("/v1/cargas", files=[("archivos", ("todo.zip", zip_de(arch), "application/zip"))]).json()
    assert any("más de un archivo" in a for a in r["avisos"])
    v = cliente.get("/v1/descargas?base=inv-pim-funcion&formato=json&departamento=16&nivel=nacional&desde=2026-01&hasta=2026-12").json()
    assert sorted((f["FUNCION"], f["PIM"]) for f in v["inv-pim-funcion"]["datos"]) == [("15: TRANSPORTE", 1000000), ("22: EDUCACION", 1000000)]
