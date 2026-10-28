"""Hidrocarburos: carga de los CSV de PERUPETRO (como los deja el flujo diario en Drive) y cuadros de la síntesis.

Datos reales de 2024 y 2025 (HIDROCARBUROS/analisis/Loreto_Ucayali_datos_mensuales.xlsx); las cifras se comparan con
el texto de Loreto 2025 (analisis/Loreto_petroleo_2025_texto.md): 7,9 millones de barriles, +22,9 %; Lote 95, +5,9 %.
"""
import io
import os
from pathlib import Path

import pandas as pd
import pytest
from openpyxl import load_workbook
from sqlalchemy import text

from app import db
from tests.test_api import cliente_http

ANALISIS = Path(__file__).resolve().parents[3] / "HIDROCARBUROS" / "analisis" / "Loreto_Ucayali_datos_mensuales.xlsx"
DEPTO = {"95": "LORETO", "8": "LORETO", "192": "LORETO", "67": "LORETO", "39": "LORETO", "131": "UCAYALI / HUANUCO", "31-C": ""}


def csvs():
    d = pd.read_excel(ANALISIS)
    d = d[(d["TEMA"] == "Producción") & d["LOTE"].notna() & (d["LOTE"] != "TODOS (PAÍS)")].copy()
    d["LOTE"] = d["LOTE"].astype(str)
    comunes = lambda x: pd.DataFrame({"AÑO": x["AÑO"], "MES": x["MES"], "FECHA": x["FECHA"], "ZONA": "SELVA",
                                      "DEPARTAMENTO": x["LOTE"].map(DEPTO), "OPERADOR": x["OPERADOR"], "LOTE": x["LOTE"],
                                      "LOTE_ORIGINAL": x["LOTE"], "LOTE_HISTORICO": x["LOTE"], "PRODUCTO": x["DETALLE"]})
    liq = d[d["DETALLE"] != "GAS NATURAL"]
    liq = comunes(liq).assign(BARRILES=liq["VALOR"], BPD=None, ESTADO_FILA="OK", FUENTE="prueba")
    zocalo = liq.iloc[:1].assign(LOTE="Z-69", LOTE_ORIGINAL="Z-69", LOTE_HISTORICO="Z-2B/Z-69", DEPARTAMENTO="", BARRILES=500)
    gas = d[d["DETALLE"] == "GAS NATURAL"]
    gas = comunes(gas).assign(MPC=gas["VALOR"], MMPCD=None, ESTADO_FILA="OK", FUENTE="prueba")
    c = pd.read_excel(ANALISIS)
    c = c[c["SERIE"].str.startswith("Canon transferido por departamento", na=False)]
    canon = pd.DataFrame({"AÑO": c["AÑO"], "MES": c["MES"], "FECHA": c["FECHA"], "DEPARTAMENTO": c["DEPARTAMENTO"],
                          "CANON_MM_SOLES": c["VALOR"], "FUENTE": "prueba"})
    canon = pd.concat([canon, canon.assign(DEPARTAMENTO="TOTAL")])
    a_bytes = lambda df: df.to_csv(index=False).encode("utf-8-sig")
    return {"PP_produccion_lote_mes_1994-2026.csv": a_bytes(pd.concat([liq, zocalo])),
            "PP_produccion_gas_lote_mes_1994-2026.csv": a_bytes(gas), "PP_canon_departamento_mes.csv": a_bytes(canon)}


@pytest.fixture()
def cliente(tmp_path, monkeypatch):
    monkeypatch.setenv("DATABASE_URL", os.environ.get("PRUEBA_DATABASE_URL") or f"sqlite:///{tmp_path / 'hidro.db'}")
    monkeypatch.delenv("APP_CLAVE", raising=False)
    db.reiniciar_motor()
    if db.motor().dialect.name == "postgresql":
        with db.motor().begin() as cx:
            cx.execute(text("DROP SCHEMA public CASCADE; CREATE SCHEMA public"))
        db.reiniciar_motor()
    yield cliente_http()
    db.reiniciar_motor()


def subir(c, archivos):
    r = c.post("/v1/cargas", files=[("archivos", (n, b, "text/csv")) for n, b in archivos.items()])
    assert r.status_code == 200, r.text
    return r.json()


def test_carga_y_departamentos(cliente):
    r = subir(cliente, csvs())
    assert set(r["bases"]) == {"hidro-produccion", "hidro-canon"}
    assert r["bases"]["hidro-produccion"]["departamentos"] == ["16", "25", "99"]   # 31-C sin departamento → Ucayali por su lote
    assert any("Z-69" in a for a in r["avisos"])
    d = cliente.get("/v1/descargas?base=hidro-produccion&departamento=25&desde=2025-01&hasta=2025-12&formato=json").json()["hidro-produccion"]["datos"]
    assert {(f["LOTE"], f["PRODUCTO"], f["UNIDAD"], f["DEPARTAMENTO"]) for f in d} == {
        ("131", "PETROLEO", "BARRILES", "UCAYALI"), ("31-C", "LGN", "BARRILES", "UCAYALI"), ("31-C", "GAS NATURAL", "MPC", "UCAYALI")}
    canon = cliente.get("/v1/descargas?base=hidro-canon&departamento=16&desde=2025-01&hasta=2025-12&formato=json").json()["hidro-canon"]
    assert canon["filas"] == 12 and abs(sum(f["CANON_MM_SOLES"] for f in canon["datos"]) - 213.85) < 0.005


def test_recargar_no_duplica_y_mueve_de_departamento(cliente):
    subir(cliente, csvs())
    a = csvs()
    liq = pd.read_csv(io.BytesIO(a["PP_produccion_lote_mes_1994-2026.csv"]), dtype=str)
    liq.loc[liq["LOTE"] == "Z-69", "DEPARTAMENTO"] = "TUMBES"          # el mapa le asigna departamento
    a["PP_produccion_lote_mes_1994-2026.csv"] = liq.to_csv(index=False).encode("utf-8-sig")
    subir(cliente, a)
    e = cliente.get("/v1/estado").json()["bases"]["hidro-produccion"]
    assert e["departamentos"] == ["16", "24", "25"]
    total = sum(len(pd.read_csv(io.BytesIO(b))) for n, b in a.items() if "produccion" in n)
    assert e["filas"] == total


def test_cuadro_loreto_igual_al_texto(cliente):
    subir(cliente, csvs())
    c = cliente.get("/v1/cuadros/hidrocarburos?departamento=16").json()["cuadros"]
    assert c["periodo"] == "2025-12" and [b["producto"] for b in c["bloques"]] == ["Petróleo"]
    filas = {f["lote"]: f for f in c["bloques"][0]["filas"]}
    assert list(filas) == ["Lote 95", "Lote 8", "Total"]                 # sin los lotes que no produjeron
    assert round(filas["Total"]["acum"]) == 7905 and filas["Total"]["var_acum"] == 22.9
    assert filas["Lote 95"]["var_acum"] == 5.9 and round(filas["Lote 8"]["acum"]) == 1130
    assert c["grafico"]["producto"] == "Petróleo" and len(c["grafico"]["serie"]) == 24
    assert c["anexo"][0]["lotes"] == ["Lote 95", "Lote 8"] and len(c["anexo"][0]["filas"]) == 13
    assert c["canon"]["acum"] == 213.85 and c["canon"]["var_acum"] is None   # la muestra no trae el canon de 2024


def test_cuadro_ucayali_y_excel(cliente):
    subir(cliente, csvs())
    c = cliente.get("/v1/cuadros/hidrocarburos?departamento=25&periodo=2025-06").json()["cuadros"]
    assert [b["producto"] for b in c["bloques"]] == ["Petróleo", "Líquidos de gas natural", "Gas natural"]
    assert c["bloques"][2]["unidad"] == "millones de pies cúbicos"
    r = cliente.get("/v1/cuadros/hidrocarburos?departamento=25&periodo=2025-06&formato=xlsx")
    assert "Hidrocarburos%20Ucayali%202025-06.xlsx" in r.headers["content-disposition"]
    wb = load_workbook(io.BytesIO(r.content))
    assert wb.sheetnames == ["Cuadro", "Gráfico", "Anexo", "Canon"]
    assert [x.value for x in wb["Cuadro"][3]] == ["Producto / lote", "Junio 2024", "Junio 2025", "Var. %", "Ene.-Jun. 2024", "Ene.-Jun. 2025", "Var. %"]
    assert cliente.get("/v1/cuadros/hidrocarburos?departamento=25&periodo=2026-03").status_code == 422
    assert cliente.get("/v1/cuadros/hidrocarburos?departamento=22").json()["cuadros"] is None


def _drive():
    """La API de Drive que ven los notebooks (en GitHub Actions apunta al almacén de la página)."""
    from app import almacen
    return almacen, almacen.build("drive", "v3", cache_discovery=False)


def _subir_como_notebook(drive, carpeta, nombre, datos, tmp_path, mime="text/csv"):
    from app import almacen
    ruta = tmp_path / nombre
    ruta.write_bytes(datos)
    return drive.files().create(body={"name": nombre, "parents": [carpeta]}, media_body=almacen.MediaFileUpload(str(ruta), mimetype=mime),
                                fields="id", supportsAllDrives=True).execute()["id"]


def test_almacen_imita_a_drive(cliente, tmp_path):
    almacen, drive = _drive()
    assert almacen.vacio()
    fid = _subir_como_notebook(drive, almacen.CARPETA_DATA, "PP_LOTES_GEO.csv", b"LOTE\n95\n", tmp_path)
    sub = drive.files().create(body={"name": "PDF", "parents": [almacen.CARPETA_DATA], "mimeType": "application/vnd.google-apps.folder"},
                               fields="id", supportsAllDrives=True).execute()["id"]
    r = drive.files().list(q=f"'{almacen.CARPETA_DATA}' in parents and trashed = false", fields="nextPageToken, files(id,name,mimeType)",
                           supportsAllDrives=True, includeItemsFromAllDrives=True).execute()
    assert {f["name"]: f["id"] for f in r["files"]} == {"PP_LOTES_GEO.csv": fid, "PDF": sub} and "nextPageToken" not in r
    (tmp_path / "PP_LOTES_GEO.csv").write_bytes(b"LOTE\n95\n8\n")
    drive.files().update(fileId=fid, media_body=almacen.MediaFileUpload(str(tmp_path / "PP_LOTES_GEO.csv")), supportsAllDrives=True).execute()
    fh = io.BytesIO()
    dl = almacen.MediaIoBaseDownload(fh, drive.files().get_media(fileId=fid, supportsAllDrives=True))
    hecho = False
    while not hecho:
        _, hecho = dl.next_chunk()
    assert fh.getvalue() == b"LOTE\n95\n8\n" == drive.files().get_media(fileId=fid).execute()
    assert not almacen.vacio() and almacen.resumen()["DATA"]["archivos"] == 1


def test_carga_solo_los_csv_de_hidrocarburos(cliente, tmp_path):
    almacen, drive = _drive()
    a = csvs()
    _subir_como_notebook(drive, almacen.CARPETA_DATA, "PP_produccion_lote_mes_1994-2025.csv", b"x", tmp_path)   # de una corrida anterior
    for n in ["PP_produccion_lote_mes_1994-2026.csv", "PP_produccion_gas_lote_mes_1994-2026.csv", "PP_canon_departamento_mes.csv"]:
        _subir_como_notebook(drive, almacen.CARPETA_DATA, n, a[n], tmp_path)
    _subir_como_notebook(drive, almacen.CARPETA_DATA, "PP_regalias_lote_anio.csv", b"x", tmp_path)
    r = cliente.post("/v1/hidrocarburos/cargar").json()
    assert set(r["bases"]) == {"hidro-produccion", "hidro-canon"} and len(r["archivos"]) == 3
    assert cliente.post("/v1/hidrocarburos/cargar").json()["bases"] == {}            # nada nuevo
    assert len(cliente.post("/v1/hidrocarburos/cargar?todo=true").json()["archivos"]) == 3
    e = cliente.get("/v1/hidrocarburos/almacen").json()
    assert e["resumen"]["DATA"]["archivos"] == 5 and e["ultima"]["origen"] == "botón de la página"
    lista = cliente.get("/v1/hidrocarburos/archivos").json()["archivos"]
    canon = next(x for x in lista if x["nombre"] == "PP_canon_departamento_mes.csv")
    d = cliente.get("/v1/hidrocarburos/archivos/descargar", params={"id": canon["id"]})
    assert d.status_code == 200 and d.content == a["PP_canon_departamento_mes.csv"]


def diaria():
    """Producción diaria: diciembre 2025 (ya está en la mensual: se ignora), enero-febrero 2026 completos y marzo al día 10."""
    filas = []
    for (a, m, dias) in ((2025, 12, 31), (2026, 1, 31), (2026, 2, 28), (2026, 3, 10)):
        for d in range(1, dias + 1):
            filas += [dict(FECHA=f"{d:02d}/{m:02d}/{a}", AÑO=a, MES=m, DIA=d, FLUIDO="LIQUIDOS", PRODUCTO="PETROLEO Y CONDENSADO",
                           UNIDAD="BBL", OPERADOR="PETROTAL", LOTE_ORIGINAL="95", LOTE="95", VALOR=15000, FUENTE="x", DEPARTAMENTO="LORETO"),
                      dict(FECHA="", AÑO=a, MES=m, DIA=d, FLUIDO="LIQUIDOS", PRODUCTO="PETROLEO Y CONDENSADO", UNIDAD="BBL",
                           OPERADOR="UPLAND", LOTE_ORIGINAL="8", LOTE="8", VALOR=3000, FUENTE="x", DEPARTAMENTO="LORETO"),
                      dict(FECHA="", AÑO=a, MES=m, DIA=d, FLUIDO="GAS", PRODUCTO="GAS NATURAL", UNIDAD="M3", OPERADOR="X",
                           LOTE_ORIGINAL="95", LOTE="95", VALOR=999, FUENTE="x", DEPARTAMENTO="LORETO")]
    return pd.DataFrame(filas).to_csv(index=False).encode("utf-8-sig")


def test_meses_de_la_produccion_diaria(cliente):
    subir(cliente, csvs())
    subir(cliente, {"PP_produccion_diaria.csv": diaria()})
    c = cliente.get("/v1/cuadros/hidrocarburos?departamento=16").json()["cuadros"]
    assert c["periodo"] == "2026-02" and c["de_diaria"] == ["Ene. 2026", "Feb. 2026"]
    assert c["parcial"] == {"periodo": "Marzo 2026", "dias": 10, "de": 31}
    filas = {f["lote"]: f for f in c["bloques"][0]["filas"]}
    assert filas["Lote 95"]["mes"] == 420.0 and filas["Lote 8"]["mes"] == 84.0      # 28 días × 15 y × 3 mil barriles
    assert [b["producto"] for b in c["bloques"]] == ["Petróleo"]                      # los m3 de gas no se suman
    ultimo = c["grafico"]["serie"][-1]
    assert ultimo == {"periodo": "Mar. 2026", "valor": 18.0, "var": None, "parcial": True}
    # diciembre 2025 sigue saliendo de la estadística mensual
    dic = cliente.get("/v1/cuadros/hidrocarburos?departamento=16&periodo=2025-12").json()["cuadros"]
    assert round([f for f in dic["bloques"][0]["filas"] if f["lote"] == "Total"][0]["acum"]) == 7905
    d = cliente.get("/v1/descargas?base=hidro-produccion&departamento=16&desde=2026-03&hasta=2026-03&formato=json").json()["hidro-produccion"]["datos"]
    assert {f["FUENTE"] for f in d} == {"PRODUCCIÓN DIARIA · 10 de 31 días"}


def test_mapa_desde_el_almacen(cliente, tmp_path):
    assert cliente.get("/v1/hidrocarburos/mapa").status_code == 404
    almacen, drive = _drive()
    _subir_como_notebook(drive, almacen.CARPETA_DATA, "PP_produccion_diaria.csv", diaria(), tmp_path)
    mapa = drive.files().create(body={"name": "MAPA", "parents": [almacen.CARPETA_DATA], "mimeType": "application/vnd.google-apps.folder"},
                                fields="id").execute()["id"]
    _subir_como_notebook(drive, mapa, "PP_mapa_lotes.html", b"<html>simple</html>", tmp_path, "text/html")
    _subir_como_notebook(drive, mapa, "PP_mapa_lotes_plus.html", "<html>plus · Perú</html>".encode(), tmp_path, "text/html")
    _subir_como_notebook(drive, mapa, "PP_lotes_indicadores.geojson", b"{}", tmp_path)
    r = almacen.sincronizar(origen="GitHub Actions")
    assert "hidro-produccion" in r["bases"] and any("Mapa de lotes" in a for a in r["avisos"])
    m = cliente.get("/v1/hidrocarburos/mapa")
    assert m.status_code == 200 and m.text == "<html>plus · Perú</html>"
    assert "attachment" in cliente.get("/v1/hidrocarburos/mapa?descargar=true").headers["content-disposition"]


def test_tabla_vieja_sin_columna_fuente(cliente):
    with db.motor().begin() as cx:                      # como estaba en Neon antes de agregar FUENTE
        cx.execute(text('CREATE TABLE "base_hidro_produccion" ("AÑO" BIGINT, "MES" BIGINT, "DEPARTAMENTO" TEXT, "ZONA" TEXT, '
                        '"OPERADOR" TEXT, "LOTE" TEXT, "LOTE_HISTORICO" TEXT, "PRODUCTO" TEXT, "UNIDAD" TEXT, "VOLUMEN" FLOAT, '
                        '"PROMEDIO_DIARIO" FLOAT, _anio BIGINT, _mes BIGINT, _depto TEXT, _nivel TEXT)'))
        cx.execute(text("INSERT INTO base_hidro_produccion VALUES (2025,12,'LORETO','','X','95','95','PETROLEO','BARRILES',100,3,2025,12,'16','LIQUIDOS')"))
    assert cliente.get("/v1/cuadros/hidrocarburos?departamento=16").status_code == 200       # antes de la primera carga: era un 500
    assert cliente.get("/v1/descargas?base=hidro-produccion&departamento=16&desde=2025-01&hasta=2025-12&formato=json").status_code == 200
    subir(cliente, csvs())
    assert cliente.get("/v1/cuadros/hidrocarburos?departamento=16").json()["cuadros"]["periodo"] == "2025-12"


def test_flujo_usa_el_almacen_en_vez_de_drive(cliente, tmp_path, monkeypatch):
    """ejecutar_notebooks.py con PP_ALMACEN=pagina: «googleapiclient» de los notebooks queda apuntando al almacén, y con
    el almacén vacío no corre (no se pierden los históricos que estaban en Drive)."""
    import sys
    sys.path.insert(0, str(Path(__file__).resolve().parents[3] / "HIDROCARBUROS" / "automatizacion"))
    import ejecutar_notebooks as en
    for m in [k for k in sys.modules if k.startswith("googleapiclient")]:
        monkeypatch.delitem(sys.modules, m)
    with pytest.raises(SystemExit, match="migrar desde Drive"):
        en.usar_almacen_de_la_pagina()
    almacen, drive = _drive()
    _subir_como_notebook(drive, almacen.CARPETA_CONTROL, "CONTROL PP produccion.csv", b"x", tmp_path)
    en.usar_almacen_de_la_pagina()
    from googleapiclient.discovery import build
    from googleapiclient.http import MediaFileUpload, MediaIoBaseDownload  # noqa: F401
    try:
        assert build is almacen.build and MediaFileUpload is almacen.MediaFileUpload
    finally:
        for m in [k for k in sys.modules if k.startswith("googleapiclient")]:
            sys.modules.pop(m)


def test_pdf_de_perupetro_van_a_disco(cliente, tmp_path, monkeypatch):
    """DATA/PDF (~450 MB de PDF públicos) no llena la base: con ALMACEN_LOCAL va a disco (la caché del flujo)."""
    monkeypatch.setenv("ALMACEN_LOCAL", str(tmp_path / "local"))
    almacen, drive = _drive()
    pdf = drive.files().create(body={"name": "PDF", "parents": [almacen.CARPETA_DATA], "mimeType": almacen.CARPETA}, fields="id").execute()["id"]
    fid = _subir_como_notebook(drive, pdf, "PP_ANUAL_2024.pdf", b"%PDF-1.4 ...", tmp_path, "application/pdf")
    assert fid == f"disco:{pdf}/PP_ANUAL_2024.pdf" and (tmp_path / "local" / pdf / "PP_ANUAL_2024.pdf").read_bytes() == b"%PDF-1.4 ..."
    r = drive.files().list(q=f"'{pdf}' in parents and trashed = false").execute()["files"]
    assert [f["name"] for f in r] == ["PP_ANUAL_2024.pdf"] and drive.files().get_media(fileId=fid).execute() == b"%PDF-1.4 ..."
    assert [a["nombre"] for a in almacen.arbol()] == ["PDF"]                           # en la base solo queda la carpeta
    with pytest.raises(FileNotFoundError):
        almacen.leer(f"disco:{pdf}/../../secreto")


def test_se_carga_el_csv_mas_reciente_segun_su_fecha_original(cliente):
    """La migración conserva la fecha de Drive: de dos CSV de producción se carga el modificado más recientemente,
    aunque se haya copiado antes."""
    from app import almacen
    a = csvs()
    almacen.crear("PP_produccion_lote_mes_1994-2026.csv", almacen.CARPETA_DATA, a["PP_produccion_lote_mes_1994-2026.csv"],
                  modificado="2026-10-28T12:40:00Z")
    almacen.crear("PP_produccion_lote_mes_2026-2026.csv", almacen.CARPETA_DATA, b"x", modificado="2026-03-01T00:00:00Z")
    r = almacen.sincronizar(origen="migración desde Drive", todo=True)
    assert [x["nombre"] for x in r["archivos"]] == ["PP_produccion_lote_mes_1994-2026.csv"] and "hidro-produccion" in r["bases"]
